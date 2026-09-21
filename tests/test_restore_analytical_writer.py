"""
Тесты Analytical Restore Writer (Stage 9C) — app.restore.analytical_writer.

Проверяют исключительно metadata-семантику и atomic-write поведение —
без resolution/lookup pipeline (это забота
tests/test_restore_analytical.py / tests/test_restore_analytical_orchestration.py).
Workbook строится напрямую через openpyxl.Workbook().
"""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.packaging.custom import IntProperty, StringProperty
from openpyxl.styles import Font

from app.restore.analytical_writer import write_analytical_restored_workbook

MARKER_PROPERTY = "DataAnonymizer.AnalyticallyRestored"
JOB_ID_PROPERTY = "DataAnonymizer.JobId"
RESTORED_FROM_JOB_ID_PROPERTY = "DataAnonymizer.RestoredFromJobId"


def _simple_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "restored value"
    return wb


def _property_names(workbook: openpyxl.Workbook) -> list[str]:
    return [p.name for p in workbook.custom_doc_props]


# ---------------------------------------------------------------------------
# Metadata
# ---------------------------------------------------------------------------


def test_successful_destination_created(tmp_path: Path) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    write_analytical_restored_workbook(wb, destination)

    assert destination.exists()
    output = openpyxl.load_workbook(destination)
    assert output.active["A1"].value == "restored value"


def test_exact_marker_value(tmp_path: Path) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    write_analytical_restored_workbook(wb, destination)

    output = openpyxl.load_workbook(destination)
    names = _property_names(output)
    assert names.count(MARKER_PROPERTY) == 1
    assert output.custom_doc_props[MARKER_PROPERTY].value == "true"


def test_stale_duplicate_marker_removed(tmp_path: Path) -> None:
    wb = _simple_workbook()
    # Эмулируем повреждённый/дважды восстановленный файл: прямая
    # манипуляция .props, обходящая уникальность CustomPropertyList.append().
    wb.custom_doc_props.props.append(StringProperty(name=MARKER_PROPERTY, value="true"))
    wb.custom_doc_props.props.append(StringProperty(name=MARKER_PROPERTY, value="true"))
    destination = tmp_path / "restored.xlsx"

    write_analytical_restored_workbook(wb, destination)

    output = openpyxl.load_workbook(destination)
    names = _property_names(output)
    assert names.count(MARKER_PROPERTY) == 1
    assert output.custom_doc_props[MARKER_PROPERTY].value == "true"


def test_unrelated_custom_properties_preserved(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb.custom_doc_props.append(StringProperty(name="Company.Department", value="Finance"))
    wb.custom_doc_props.append(IntProperty(name="Company.Revision", value=7))
    destination = tmp_path / "restored.xlsx"

    write_analytical_restored_workbook(wb, destination)

    output = openpyxl.load_workbook(destination)
    assert output.custom_doc_props["Company.Department"].value == "Finance"
    assert output.custom_doc_props["Company.Revision"].value == 7


def test_existing_job_id_preserved(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb.custom_doc_props.append(StringProperty(name=JOB_ID_PROPERTY, value="original-job-123"))
    destination = tmp_path / "restored.xlsx"

    write_analytical_restored_workbook(wb, destination)

    output = openpyxl.load_workbook(destination)
    assert output.custom_doc_props[JOB_ID_PROPERTY].value == "original-job-123"
    assert output.custom_doc_props[MARKER_PROPERTY].value == "true"


def test_existing_restored_from_job_id_preserved(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb.custom_doc_props.append(StringProperty(name=RESTORED_FROM_JOB_ID_PROPERTY, value="source-job-456"))
    destination = tmp_path / "restored.xlsx"

    write_analytical_restored_workbook(wb, destination)

    output = openpyxl.load_workbook(destination)
    assert output.custom_doc_props[RESTORED_FROM_JOB_ID_PROPERTY].value == "source-job-456"
    assert output.custom_doc_props[MARKER_PROPERTY].value == "true"


# ---------------------------------------------------------------------------
# Preservation of formatting/structure
# ---------------------------------------------------------------------------


def test_styles_number_format_merged_ranges_sheet_visibility_preserved(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "0012345678"
    ws1["A1"].number_format = "@"
    ws1["B1"] = "bold text"
    ws1["B1"].font = Font(bold=True)
    ws1.merge_cells("C1:D1")
    ws1["C1"] = "merged"

    ws2 = wb.create_sheet("Hidden")
    ws2.sheet_state = "hidden"

    destination = tmp_path / "restored.xlsx"
    write_analytical_restored_workbook(wb, destination)

    output = openpyxl.load_workbook(destination)
    assert output["First"]["A1"].value == "0012345678"
    assert output["First"]["A1"].number_format == "@"
    assert output["First"]["B1"].font.bold is True
    assert "C1:D1" in [str(r) for r in output["First"].merged_cells.ranges]
    assert output.sheetnames == ["First", "Hidden"]
    assert output["Hidden"].sheet_state == "hidden"


# ---------------------------------------------------------------------------
# Atomic output
# ---------------------------------------------------------------------------


def test_same_directory_temp_and_fsync_and_rename(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    fsync_calls = {"n": 0}
    real_fsync = os.fsync

    def counting_fsync(fd: int) -> None:
        fsync_calls["n"] += 1
        real_fsync(fd)

    rename_calls = {"n": 0}
    real_rename = os.rename

    def counting_rename(src, dst):
        rename_calls["n"] += 1
        return real_rename(src, dst)

    monkeypatch.setattr(os, "fsync", counting_fsync)
    monkeypatch.setattr(os, "rename", counting_rename)

    write_analytical_restored_workbook(wb, destination)

    monkeypatch.undo()
    assert fsync_calls["n"] == 1
    assert rename_calls["n"] == 1
    assert destination.exists()
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []


def test_os_replace_never_used(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    def failing_replace(*args, **kwargs):
        raise AssertionError("os.replace должен НЕ вызываться analytical writer'ом")

    monkeypatch.setattr(os, "replace", failing_replace)

    write_analytical_restored_workbook(wb, destination)

    monkeypatch.undo()
    assert destination.exists()


def test_save_failure_no_destination_and_temp_cleaned_up(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    def failing_save(self, *args, **kwargs):
        raise OSError("simulated save failure")

    monkeypatch.setattr(openpyxl.Workbook, "save", failing_save)

    with pytest.raises(OSError):
        write_analytical_restored_workbook(wb, destination)

    monkeypatch.undo()
    assert not destination.exists()
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []


def test_fsync_failure_no_destination_and_temp_cleaned_up(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    def failing_fsync(fd: int) -> None:
        raise OSError("simulated fsync failure")

    monkeypatch.setattr(os, "fsync", failing_fsync)

    with pytest.raises(OSError):
        write_analytical_restored_workbook(wb, destination)

    monkeypatch.undo()
    assert not destination.exists()
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []


def test_existing_destination_race_raises_file_exists_and_bytes_unchanged(tmp_path: Path) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"
    original_bytes = b"PRE-EXISTING CONFIDENTIAL CONTENT"
    destination.write_bytes(original_bytes)

    with pytest.raises(FileExistsError):
        write_analytical_restored_workbook(wb, destination)

    assert destination.read_bytes() == original_bytes
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []
