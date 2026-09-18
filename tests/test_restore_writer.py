"""
Тесты Restore-specific atomic writer (app.restore.writer).

Проверяют исключительно metadata-семантику и atomic-write поведение —
без anonymization/provenance pipeline (это забота
tests/test_restore_single_file.py). Workbook строится напрямую через
openpyxl.Workbook().
"""

from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.packaging.custom import IntProperty, StringProperty

from app.restore.writer import write_restored_workbook

JOB_ID_PROPERTY = "DataAnonymizer.JobId"
RESTORED_FROM_JOB_ID_PROPERTY = "DataAnonymizer.RestoredFromJobId"


def _simple_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "restored value"
    return wb


def _property_names(workbook: openpyxl.Workbook) -> list[str]:
    return [p.name for p in workbook.custom_doc_props]


# ---------------------------------------------------------------------------
# Metadata
# ---------------------------------------------------------------------------


def test_job_id_absent_after_write(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb.custom_doc_props.append(StringProperty(name=JOB_ID_PROPERTY, value="source-job-1"))
    destination = tmp_path / "restored.xlsx"

    write_restored_workbook(wb, destination, source_job_id="source-job-1")

    output = openpyxl.load_workbook(destination)
    assert JOB_ID_PROPERTY not in _property_names(output)


def test_exactly_one_restored_from_job_id_with_exact_value(tmp_path: Path) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    write_restored_workbook(wb, destination, source_job_id="source-job-42")

    output = openpyxl.load_workbook(destination)
    names = _property_names(output)
    assert names.count(RESTORED_FROM_JOB_ID_PROPERTY) == 1
    assert output.custom_doc_props[RESTORED_FROM_JOB_ID_PROPERTY].value == "source-job-42"


def test_stale_duplicate_restored_metadata_removed(tmp_path: Path) -> None:
    wb = _simple_workbook()
    # Эмулируем повреждённый/дважды восстановленный файл: несколько stale
    # occurrences ОБОИХ имён (прямая манипуляция .props, обходящая
    # уникальность CustomPropertyList.append()).
    wb.custom_doc_props.props.append(StringProperty(name=JOB_ID_PROPERTY, value="old-1"))
    wb.custom_doc_props.props.append(StringProperty(name=JOB_ID_PROPERTY, value="old-2"))
    wb.custom_doc_props.props.append(StringProperty(name=RESTORED_FROM_JOB_ID_PROPERTY, value="stale-1"))
    wb.custom_doc_props.props.append(StringProperty(name=RESTORED_FROM_JOB_ID_PROPERTY, value="stale-2"))
    destination = tmp_path / "restored.xlsx"

    write_restored_workbook(wb, destination, source_job_id="fresh-job-id")

    output = openpyxl.load_workbook(destination)
    names = _property_names(output)
    assert names.count(JOB_ID_PROPERTY) == 0
    assert names.count(RESTORED_FROM_JOB_ID_PROPERTY) == 1
    assert output.custom_doc_props[RESTORED_FROM_JOB_ID_PROPERTY].value == "fresh-job-id"


def test_unrelated_custom_properties_preserved(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb.custom_doc_props.append(StringProperty(name="Company.Department", value="Finance"))
    wb.custom_doc_props.append(IntProperty(name="Company.Revision", value=7))
    destination = tmp_path / "restored.xlsx"

    write_restored_workbook(wb, destination, source_job_id="job-1")

    output = openpyxl.load_workbook(destination)
    assert output.custom_doc_props["Company.Department"].value == "Finance"
    assert output.custom_doc_props["Company.Revision"].value == 7


# ---------------------------------------------------------------------------
# Atomic output
# ---------------------------------------------------------------------------


def test_atomic_save_success(tmp_path: Path) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    write_restored_workbook(wb, destination, source_job_id="job-1")

    assert destination.exists()
    output = openpyxl.load_workbook(destination)
    assert output.active["A1"].value == "restored value"


def test_save_failure_no_destination_and_temp_cleaned_up(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    def failing_save(self, *args, **kwargs):
        raise OSError("simulated save failure")

    monkeypatch.setattr(openpyxl.Workbook, "save", failing_save)

    with pytest.raises(OSError):
        write_restored_workbook(wb, destination, source_job_id="job-1")

    monkeypatch.undo()
    assert not destination.exists()
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []


def test_fsync_failure_no_destination_and_temp_cleaned_up(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"

    def failing_fsync(fd: int) -> None:
        raise OSError("simulated fsync failure")

    monkeypatch.setattr(os, "fsync", failing_fsync)

    with pytest.raises(OSError):
        write_restored_workbook(wb, destination, source_job_id="job-1")

    monkeypatch.undo()
    assert not destination.exists()
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []


def test_existing_destination_race_raises_file_exists_and_bytes_unchanged(tmp_path: Path) -> None:
    """
    Экспериментально подтверждённая на Contract Freeze semantics: если
    destination появился конкурентно (уже существует к моменту
    финального os.rename), запись должна отказаться, не перезаписав его.
    """
    wb = _simple_workbook()
    destination = tmp_path / "restored.xlsx"
    original_bytes = b"PRE-EXISTING CONFIDENTIAL CONTENT"
    destination.write_bytes(original_bytes)

    with pytest.raises(FileExistsError):
        write_restored_workbook(wb, destination, source_job_id="job-1")

    assert destination.read_bytes() == original_bytes
    leftover = [p for p in tmp_path.iterdir() if p.name.startswith(f".{destination.name}.")]
    assert leftover == []
