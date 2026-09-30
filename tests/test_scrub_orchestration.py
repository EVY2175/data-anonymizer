"""
Тесты Stage 10C.2.5 — Safe Scrub Orchestration.

Методология: реальные full-pipeline positive-тесты используют настоящие
workbook + настоящий InMemoryMappingStore/InMemoryIdentifierMappingStore —
2.3/2.4/Stage10A НЕ мокаются в них. Negative pre-mutation-invariant тесты
мокают ТОЛЬКО app.scrub.orchestrate.scrub_workbook_object_model, чтобы
независимо доказать "NOT CALLED" без выполнения реальной мутации.

MAJOR-1 correction (ownership guard) тесты дополнительно мокают
scrub_workbook_object_model на конкретный (в т.ч. source-эквивалентный)
Path, а в одном тесте — os.path.samefile, чтобы независимо доказать
comparison-failure fail-closed без построения hostile symlink-окружения.
"""

from __future__ import annotations

import hashlib
import os
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles.colors import Color
from openpyxl.styles.fills import GradientFill, PatternFill

from app.mapping.identifier_memory import InMemoryIdentifierMappingStore
from app.mapping.memory import InMemoryMappingStore
from app.models.entities import EntityType, MappingEntry
from app.models.identifiers import IdentifierMappingEntry, IdentifierType
from app.scrub.errors import PackageScrubError, PackageScrubReason
from app.scrub.orchestrate import (
    ScrubOrchestrationError,
    ScrubOrchestrationReason,
    ScrubResult,
    _interpret_safety_report,
    _temp_ownership_confirmed,
    scrub_workbook_for_external_candidate,
)

SECRET = "SECRET_ORCHESTRATION_VALUE_94731"

_RESTORED_MARKER_A = "DataAnonymizer.AnalyticallyRestored"
_RESTORED_MARKER_B = "DataAnonymizer.RestoredFromJobId"
_JOB_ID_PROPERTY = "DataAnonymizer.JobId"


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------


def _save(tmp_path: Path, wb, name: str = "src.xlsx") -> Path:
    src = tmp_path / name
    wb.save(src)
    wb.close()
    return src


def _sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _empty_store() -> InMemoryMappingStore:
    return InMemoryMappingStore()


# ----------------------------------------------------------------------
# §41: реальные full-pipeline positive-тесты (2.3/2.4/Stage10A НЕ мокаются)
# ----------------------------------------------------------------------


def test_plain_workbook_full_pipeline_pass(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hello"
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        assert isinstance(result, ScrubResult)
        assert result.path.exists()
        assert result.validation.worksheet_count == 1
        assert len(result.sha256) == 64
        assert all(ch in "0123456789abcdef" for ch in result.sha256)
    assert not result.path.exists()


def test_patternfill_workbook_full_pipeline_pass(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = PatternFill(patternType="solid", fgColor=Color(rgb="FF00FF00"))
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        assert result.path.exists()
        assert result.validation.worksheet_count == 1


def test_gradientfill_workbook_full_pipeline_pass(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = GradientFill(type="linear", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        assert result.path.exists()
        assert result.validation.worksheet_count == 1


# ----------------------------------------------------------------------
# §37: SafetyReport findings — restored markers / known values / precedence
# ----------------------------------------------------------------------


def test_restored_marker_analytically_restored_rejected_before_mutation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.custom_doc_props.append(StringProperty(name=_RESTORED_MARKER_A, value="true"))
    source = _save(tmp_path, wb)

    with patch("app.scrub.orchestrate.scrub_workbook_object_model") as mock_scrub:
        with pytest.raises(PackageScrubError) as excinfo:
            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                pass
    assert excinfo.value.reason == PackageScrubReason.RESTORED_MARKER_PRESENT
    mock_scrub.assert_not_called()


def test_restored_marker_restored_from_job_id_rejected_before_mutation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.custom_doc_props.append(StringProperty(name=_RESTORED_MARKER_B, value="job-123"))
    source = _save(tmp_path, wb)

    with patch("app.scrub.orchestrate.scrub_workbook_object_model") as mock_scrub:
        with pytest.raises(PackageScrubError) as excinfo:
            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                pass
    assert excinfo.value.reason == PackageScrubReason.RESTORED_MARKER_PRESENT
    mock_scrub.assert_not_called()


def test_job_id_alone_is_not_blocking(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.custom_doc_props.append(StringProperty(name=_JOB_ID_PROPERTY, value="job-123"))
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        assert result.path.exists()


def test_known_entity_value_rejected_before_mutation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "ACME Corp"
    source = _save(tmp_path, wb)
    store = InMemoryMappingStore()
    store.add(MappingEntry(alias="C_0001", real_value="ACME Corp", entity_type=EntityType.COMPANY))

    with patch("app.scrub.orchestrate.scrub_workbook_object_model") as mock_scrub:
        with pytest.raises(ScrubOrchestrationError) as excinfo:
            with scrub_workbook_for_external_candidate(source, store, None):
                pass
    assert excinfo.value.reason == ScrubOrchestrationReason.KNOWN_ENTITY_VALUE_PRESENT
    mock_scrub.assert_not_called()


def test_known_identifier_value_rejected_before_mutation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "7701234567"
    source = _save(tmp_path, wb)
    id_store = InMemoryIdentifierMappingStore()
    id_store.add(
        IdentifierMappingEntry(token="ID_0001", identifier_value="7701234567", identifier_type=IdentifierType.INN)
    )

    with patch("app.scrub.orchestrate.scrub_workbook_object_model") as mock_scrub:
        with pytest.raises(ScrubOrchestrationError) as excinfo:
            with scrub_workbook_for_external_candidate(source, _empty_store(), id_store):
                pass
    assert excinfo.value.reason == ScrubOrchestrationReason.KNOWN_IDENTIFIER_VALUE_PRESENT
    mock_scrub.assert_not_called()


def test_entity_and_identifier_precedence_entity_wins(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ACME Corp"
    ws["A2"] = "7701234567"
    source = _save(tmp_path, wb)
    store = InMemoryMappingStore()
    store.add(MappingEntry(alias="C_0001", real_value="ACME Corp", entity_type=EntityType.COMPANY))
    id_store = InMemoryIdentifierMappingStore()
    id_store.add(
        IdentifierMappingEntry(token="ID_0001", identifier_value="7701234567", identifier_type=IdentifierType.INN)
    )

    with pytest.raises(ScrubOrchestrationError) as excinfo:
        with scrub_workbook_for_external_candidate(source, store, id_store):
            pass
    assert excinfo.value.reason == ScrubOrchestrationReason.KNOWN_ENTITY_VALUE_PRESENT


def test_restored_marker_precedence_over_entity_and_identifier(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ACME Corp"
    ws["A2"] = "7701234567"
    wb.custom_doc_props.append(StringProperty(name=_RESTORED_MARKER_A, value="true"))
    source = _save(tmp_path, wb)
    store = InMemoryMappingStore()
    store.add(MappingEntry(alias="C_0001", real_value="ACME Corp", entity_type=EntityType.COMPANY))
    id_store = InMemoryIdentifierMappingStore()
    id_store.add(
        IdentifierMappingEntry(token="ID_0001", identifier_value="7701234567", identifier_type=IdentifierType.INN)
    )

    with pytest.raises(PackageScrubError) as excinfo:
        with scrub_workbook_for_external_candidate(source, store, id_store):
            pass
    assert excinfo.value.reason == PackageScrubReason.RESTORED_MARKER_PRESENT


def test_formula_finding_alone_does_not_block(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = 2
    ws["C1"] = "=A1+B1"
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        assert result.path.exists()
        assert result.validation.formula_cell_count == 1


# ----------------------------------------------------------------------
# §38: defensive fail-closed branch — unit-test приватного helper-а с
# controlled fake (production SafetyReport НЕ трогается)
# ----------------------------------------------------------------------


class _FakeBlockingReport:
    """Duck-typed минимальный fake — интерфейс count()/has_blocking_findings."""

    def count(self, kind) -> int:
        return 0

    @property
    def has_blocking_findings(self) -> bool:
        return True


class _FakeCleanReport:
    def count(self, kind) -> int:
        return 0

    @property
    def has_blocking_findings(self) -> bool:
        return False


def test_defensive_unknown_blocking_kind_fails_closed():
    with pytest.raises(ScrubOrchestrationError) as excinfo:
        _interpret_safety_report(_FakeBlockingReport())
    assert excinfo.value.reason == ScrubOrchestrationReason.INTERNAL_FAILURE


def test_interpret_safety_report_clean_does_not_raise():
    _interpret_safety_report(_FakeCleanReport())  # не должно поднять исключение


# ----------------------------------------------------------------------
# §39-40: preflight/inventory before Stage10A (order invariant)
# ----------------------------------------------------------------------


def test_preflight_failure_prevents_stage10a_and_mutation(tmp_path):
    fake = tmp_path / "not_xlsx.txt"
    fake.write_bytes(b"hello")

    with patch("app.scrub.orchestrate.inspect_workbook_for_external_ai") as mock_stage10a:
        with patch("app.scrub.orchestrate.scrub_workbook_object_model") as mock_scrub:
            with pytest.raises(PackageScrubError):
                with scrub_workbook_for_external_candidate(fake, _empty_store(), None):
                    pass
    mock_stage10a.assert_not_called()
    mock_scrub.assert_not_called()


def test_inventory_failure_prevents_stage10a_and_mutation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    with patch("app.scrub.orchestrate.inspect_package_policy", side_effect=PackageScrubError(PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT)):
        with patch("app.scrub.orchestrate.inspect_workbook_for_external_ai") as mock_stage10a:
            with patch("app.scrub.orchestrate.scrub_workbook_object_model") as mock_scrub:
                with pytest.raises(PackageScrubError) as excinfo:
                    with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                        pass
    assert excinfo.value.reason == PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT
    mock_stage10a.assert_not_called()
    mock_scrub.assert_not_called()


# ----------------------------------------------------------------------
# §29-31: normal cleanup / caller exception / BaseException
# ----------------------------------------------------------------------


def test_normal_cleanup_removes_temp(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        assert result.path.exists()
    assert not result.path.exists()


def test_caller_exception_cleans_temp_and_propagates_original(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    captured_path = None
    with pytest.raises(RuntimeError, match="caller sentinel"):
        with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
            captured_path = result.path
            raise RuntimeError("caller sentinel")
    assert captured_path is not None
    assert not captured_path.exists()


class _CustomBaseException(BaseException):
    pass


def test_caller_base_exception_cleans_temp_and_propagates(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    captured_path = None
    with pytest.raises(_CustomBaseException):
        with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
            captured_path = result.path
            raise _CustomBaseException()
    assert captured_path is not None
    assert not captured_path.exists()


# ----------------------------------------------------------------------
# §28: nested contexts
# ----------------------------------------------------------------------


def test_nested_contexts_use_distinct_paths(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as a:
        with scrub_workbook_for_external_candidate(source, _empty_store(), None) as b:
            assert a.path != b.path
            assert a.path.exists()
            assert b.path.exists()
        assert not b.path.exists()
        assert a.path.exists()
    assert not a.path.exists()


# ----------------------------------------------------------------------
# §18: STEP 6 (postvalidate) failure — best-effort unlink, original exception
# ----------------------------------------------------------------------


def test_postvalidation_failure_cleans_temp_and_preserves_original(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    sentinel_error = PackageScrubError(PackageScrubReason.POST_VALIDATION_FAILED)
    created_paths = []
    from app.scrub.mutate import scrub_workbook_object_model as real_scrub

    def _spy_scrub(path):
        result_path = real_scrub(path)
        created_paths.append(result_path)
        return result_path

    with patch("app.scrub.orchestrate.scrub_workbook_object_model", side_effect=_spy_scrub):
        with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package", side_effect=sentinel_error):
            with pytest.raises(PackageScrubError) as excinfo:
                with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                    pass
    assert excinfo.value is sentinel_error
    assert excinfo.value.reason == PackageScrubReason.POST_VALIDATION_FAILED
    assert len(created_paths) == 1
    assert not created_paths[0].exists()


# ----------------------------------------------------------------------
# §34-36: SHA tests
# ----------------------------------------------------------------------


def test_sha256_matches_independent_hash_inside_context(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        independent = _sha256_of(result.path)
        assert independent == result.sha256


def test_sha256_computed_after_validation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    call_order = []
    from app.scrub.postvalidate import validate_scrubbed_workbook_package as real_validate
    from app.scrub.orchestrate import _compute_sha256 as real_sha

    def _spy_validate(path):
        call_order.append("validate")
        return real_validate(path)

    def _spy_sha(path):
        call_order.append("sha")
        return real_sha(path)

    with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package", side_effect=_spy_validate):
        with patch("app.scrub.orchestrate._compute_sha256", side_effect=_spy_sha):
            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                pass
    assert call_order == ["validate", "sha"]


def test_caller_mutation_invalidates_sha_correspondence(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        original_sha = result.sha256
        with open(result.path, "ab") as f:
            f.write(b"\x00\x01\x02mutated-by-caller")
        new_sha = _sha256_of(result.path)
        assert new_sha != original_sha
    assert not result.path.exists()


def test_sha_failure_cleans_temp_raises_internal_failure_privacy_safe(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    created_paths = []
    from app.scrub.mutate import scrub_workbook_object_model as real_scrub

    def _spy_scrub(path):
        result_path = real_scrub(path)
        created_paths.append(result_path)
        return result_path

    def _failing_sha(path):
        raise OSError(f"simulated I/O failure: {SECRET}")

    with patch("app.scrub.orchestrate.scrub_workbook_object_model", side_effect=_spy_scrub):
        with patch("app.scrub.orchestrate._compute_sha256", side_effect=_failing_sha):
            with pytest.raises(ScrubOrchestrationError) as excinfo:
                with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                    pass
    assert excinfo.value.reason == ScrubOrchestrationReason.INTERNAL_FAILURE
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None
    assert len(created_paths) == 1
    assert not created_paths[0].exists()


# ----------------------------------------------------------------------
# §32-33: cleanup-failure success-path / error-path
# ----------------------------------------------------------------------


def test_cleanup_failure_on_success_path_raises_cleanup_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    leftover_path = None
    original_unlink = Path.unlink

    def _failing_unlink(self, *args, **kwargs):
        nonlocal leftover_path
        leftover_path = self
        raise OSError(f"simulated unlink failure: {SECRET}")

    with patch.object(Path, "unlink", _failing_unlink):
        with pytest.raises(ScrubOrchestrationError) as excinfo:
            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                pass  # штатный выход, unlink monkeypatched на сбой

    assert excinfo.value.reason == ScrubOrchestrationReason.CLEANUP_FAILED
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None

    # Test-cleanup: monkeypatch снят (with-блок patch.object уже вышел),
    # реальный unlink теперь снова доступен — убираем orphan вручную.
    if leftover_path is not None and leftover_path.exists():
        original_unlink(leftover_path, missing_ok=True)


def test_cleanup_failure_on_error_path_preserves_primary_exception(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    leftover_path = None
    original_unlink = Path.unlink

    def _failing_unlink(self, *args, **kwargs):
        nonlocal leftover_path
        leftover_path = self
        raise OSError(f"simulated unlink failure: {SECRET}")

    with patch.object(Path, "unlink", _failing_unlink):
        with pytest.raises(RuntimeError, match="primary sentinel"):
            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                raise RuntimeError("primary sentinel")

    if leftover_path is not None and leftover_path.exists():
        original_unlink(leftover_path, missing_ok=True)


# ----------------------------------------------------------------------
# §27: source immutability
# ----------------------------------------------------------------------


def test_source_unchanged_after_successful_pipeline(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)
    before = _sha256_of(source)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None):
        pass

    after = _sha256_of(source)
    assert before == after


def test_source_unchanged_after_restored_marker_failure(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.custom_doc_props.append(StringProperty(name=_RESTORED_MARKER_A, value="true"))
    source = _save(tmp_path, wb)
    before = _sha256_of(source)

    with pytest.raises(PackageScrubError):
        with scrub_workbook_for_external_candidate(source, _empty_store(), None):
            pass

    after = _sha256_of(source)
    assert before == after


# ----------------------------------------------------------------------
# §43: ScrubResult repr не содержит source filename/known values
# ----------------------------------------------------------------------


def test_result_repr_does_not_contain_source_filename_or_known_values(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "not-a-secret-cell-value"
    source_name = "very_unique_source_filename_marker.xlsx"
    source = _save(tmp_path, wb, source_name)

    with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
        rendered = repr(result)
        assert "very_unique_source_filename_marker" not in rendered
        assert source_name not in str(result.path)


# ----------------------------------------------------------------------
# §50: error privacy sweep — known entity / known identifier
# ----------------------------------------------------------------------


def test_error_privacy_known_entity(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = SECRET
    source = _save(tmp_path, wb)
    store = InMemoryMappingStore()
    store.add(MappingEntry(alias="C_0001", real_value=SECRET, entity_type=EntityType.COMPANY))

    with pytest.raises(ScrubOrchestrationError) as excinfo:
        with scrub_workbook_for_external_candidate(source, store, None):
            pass
    assert excinfo.value.reason == ScrubOrchestrationReason.KNOWN_ENTITY_VALUE_PRESENT
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None


def test_error_privacy_known_identifier(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = SECRET
    source = _save(tmp_path, wb)
    id_store = InMemoryIdentifierMappingStore()
    id_store.add(IdentifierMappingEntry(token="ID_0001", identifier_value=SECRET, identifier_type=IdentifierType.INN))

    with pytest.raises(ScrubOrchestrationError) as excinfo:
        with scrub_workbook_for_external_candidate(source, _empty_store(), id_store):
            pass
    assert excinfo.value.reason == ScrubOrchestrationReason.KNOWN_IDENTIFIER_VALUE_PRESENT
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None


def test_error_privacy_restored_marker(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.custom_doc_props.append(StringProperty(name=_RESTORED_MARKER_A, value=SECRET))
    source = _save(tmp_path, wb)

    with pytest.raises(PackageScrubError) as excinfo:
        with scrub_workbook_for_external_candidate(source, _empty_store(), None):
            pass
    assert excinfo.value.reason == PackageScrubReason.RESTORED_MARKER_PRESENT
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None


# ----------------------------------------------------------------------
# MAJOR-1 correction — ownership guard (cleanup target integrity)
# ----------------------------------------------------------------------


def test_temp_ownership_confirmed_true_for_distinct_files(tmp_path):
    a = tmp_path / "a.txt"
    b = tmp_path / "b.txt"
    a.write_bytes(b"a")
    b.write_bytes(b"b")
    assert _temp_ownership_confirmed(a, b) is True


def test_temp_ownership_confirmed_false_for_same_file(tmp_path):
    a = tmp_path / "a.txt"
    a.write_bytes(b"a")
    assert _temp_ownership_confirmed(a, a) is False


def test_ownership_guard_rejects_exact_source_path_returned(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)
    before = _sha256_of(source)

    with patch("app.scrub.orchestrate.scrub_workbook_object_model", return_value=source):
        with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package") as mock_validate:
            with patch("app.scrub.orchestrate._compute_sha256") as mock_sha:
                with patch("app.scrub.orchestrate._best_effort_unlink") as mock_best_effort:
                    with patch("app.scrub.orchestrate._unlink_or_raise_cleanup_error") as mock_strict:
                        with pytest.raises(ScrubOrchestrationError) as excinfo:
                            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                                pass

    assert excinfo.value.reason == ScrubOrchestrationReason.INTERNAL_FAILURE
    mock_validate.assert_not_called()
    mock_sha.assert_not_called()
    mock_best_effort.assert_not_called()
    mock_strict.assert_not_called()
    assert source.exists()
    assert _sha256_of(source) == before


def test_ownership_guard_rejects_dot_alias_source_path(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)
    before = _sha256_of(source)

    # pathlib сворачивает одиночный "." при построении Path, поэтому alias
    # строится как строка (os.path.samefile принимает str не хуже Path) —
    # это то самое лексическое различие, которое обязан пережить guard.
    aliased_source = str(tmp_path) + os.sep + "." + os.sep + source.name
    assert aliased_source != str(source)  # разное лексическое представление
    assert os.path.samefile(aliased_source, source)  # но тот же файл на диске

    with patch("app.scrub.orchestrate.scrub_workbook_object_model", return_value=aliased_source):
        with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package") as mock_validate:
            with patch("app.scrub.orchestrate._compute_sha256") as mock_sha:
                with patch("app.scrub.orchestrate._best_effort_unlink") as mock_best_effort:
                    with patch("app.scrub.orchestrate._unlink_or_raise_cleanup_error") as mock_strict:
                        with pytest.raises(ScrubOrchestrationError) as excinfo:
                            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                                pass

    assert excinfo.value.reason == ScrubOrchestrationReason.INTERNAL_FAILURE
    mock_validate.assert_not_called()
    mock_sha.assert_not_called()
    mock_best_effort.assert_not_called()
    mock_strict.assert_not_called()
    assert source.exists()
    assert _sha256_of(source) == before


def test_ownership_guard_rejects_relative_vs_absolute_alias_source_path(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)
    before = _sha256_of(source)

    monkeypatch.chdir(tmp_path)
    relative_alias = Path(source.name)  # относительный путь к тому же файлу
    assert relative_alias != source
    assert os.path.samefile(relative_alias, source)

    with patch("app.scrub.orchestrate.scrub_workbook_object_model", return_value=relative_alias):
        with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package") as mock_validate:
            with patch("app.scrub.orchestrate._compute_sha256") as mock_sha:
                with patch("app.scrub.orchestrate._best_effort_unlink") as mock_best_effort:
                    with patch("app.scrub.orchestrate._unlink_or_raise_cleanup_error") as mock_strict:
                        with pytest.raises(ScrubOrchestrationError) as excinfo:
                            with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                                pass

    assert excinfo.value.reason == ScrubOrchestrationReason.INTERNAL_FAILURE
    mock_validate.assert_not_called()
    mock_sha.assert_not_called()
    mock_best_effort.assert_not_called()
    mock_strict.assert_not_called()
    assert source.exists()
    assert _sha256_of(source) == before


def test_ownership_guard_allows_real_distinct_temp_full_pipeline(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)

    call_order = []
    from app.scrub.postvalidate import validate_scrubbed_workbook_package as real_validate

    def _spy_validate(path):
        call_order.append("validate")
        return real_validate(path)

    with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package", side_effect=_spy_validate):
        with scrub_workbook_for_external_candidate(source, _empty_store(), None) as result:
            assert result.path.exists()
            assert result.path != source
    assert call_order == ["validate"]


def test_ownership_guard_comparison_failure_privacy_safe(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    source = _save(tmp_path, wb)
    before = _sha256_of(source)

    created_paths = []
    from app.scrub.mutate import scrub_workbook_object_model as real_scrub

    def _spy_scrub(path):
        result_path = real_scrub(path)
        created_paths.append(result_path)
        return result_path

    def _failing_samefile(a, b):
        raise OSError(f"simulated comparison failure: {SECRET}")

    with patch("app.scrub.orchestrate.scrub_workbook_object_model", side_effect=_spy_scrub):
        with patch("app.scrub.orchestrate.os.path.samefile", side_effect=_failing_samefile):
            with patch("app.scrub.orchestrate.validate_scrubbed_workbook_package") as mock_validate:
                with patch("app.scrub.orchestrate._compute_sha256") as mock_sha:
                    with patch("app.scrub.orchestrate._best_effort_unlink") as mock_best_effort:
                        with patch("app.scrub.orchestrate._unlink_or_raise_cleanup_error") as mock_strict:
                            with pytest.raises(ScrubOrchestrationError) as excinfo:
                                with scrub_workbook_for_external_candidate(source, _empty_store(), None):
                                    pass

    assert excinfo.value.reason == ScrubOrchestrationReason.INTERNAL_FAILURE
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None
    mock_validate.assert_not_called()
    mock_sha.assert_not_called()
    mock_best_effort.assert_not_called()
    mock_strict.assert_not_called()
    assert len(created_paths) == 1
    assert source.exists()
    assert _sha256_of(source) == before


# ----------------------------------------------------------------------
# Reason enum / error class — базовая корректность
# ----------------------------------------------------------------------


def test_scrub_orchestration_reason_has_exactly_four_values():
    assert {r.value for r in ScrubOrchestrationReason} == {
        "known_entity_value_present",
        "known_identifier_value_present",
        "internal_failure",
        "cleanup_failed",
    }


def test_scrub_orchestration_error_rejects_non_reason_type():
    with pytest.raises(TypeError):
        ScrubOrchestrationError("not-a-reason")  # type: ignore[arg-type]


# ----------------------------------------------------------------------
# §42: no network — production module не импортирует сетевые клиенты
# ----------------------------------------------------------------------


def test_orchestrate_module_has_no_network_imports():
    import app.scrub.orchestrate as orchestrate_module

    source = Path(orchestrate_module.__file__).read_text(encoding="utf-8")
    for forbidden in ("import requests", "import httpx", "import urllib.request", "import socket"):
        assert forbidden not in source
