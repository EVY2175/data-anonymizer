"""
Тесты Stage 10C.1 — identifier expected-count evidence, Step A
(app.coverage.models.IdentifierColumnEvidence,
app.coverage.worksheet_policy.scan_process_worksheet/
prepare_workbook_for_anonymization).

Step B (сверка с фактическим output/store) НЕ реализуется здесь — это
Stage 10C.4.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import openpyxl
import pytest

from app.coverage.models import IdentifierColumnEvidence, WorksheetPolicy
from app.coverage.worksheet_policy import prepare_workbook_for_anonymization, scan_process_worksheet
from app.models.identifiers import IdentifierType
from app.models.rules import Action, FieldRule, FieldType

VALID_INN10 = "7707083893"
VALID_INN12 = "500100732259"
VALID_OGRN13 = "1047709000008"
VALID_OGRNIP15 = "304770900000126"
VALID_KPP = "770701001"


def _pseudonymize_rule(field_type: FieldType, column_name: str = "Col") -> FieldRule:
    return FieldRule(column_name=column_name, field_type=field_type, action=Action.PSEUDONYMIZE)


def _worksheet_with_column(header: str, values: list[object]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = header
    for idx, value in enumerate(values, start=2):
        ws.cell(row=idx, column=1, value=value)
    return ws


@pytest.fixture()
def tmp_xlsx_path(tmp_path: Path) -> Path:
    return tmp_path / "source.xlsx"


# ---------------------------------------------------------------------------
# A. IdentifierColumnEvidence — модель
# ---------------------------------------------------------------------------


def test_evidence_valid_construction() -> None:
    evidence = IdentifierColumnEvidence(
        worksheet_index=1, column=2, identifier_type=IdentifierType.INN, expected_count=3
    )
    assert evidence.worksheet_index == 1
    assert evidence.column == 2
    assert evidence.identifier_type is IdentifierType.INN
    assert evidence.expected_count == 3


@pytest.mark.parametrize("worksheet_index", [0, -1, True])
def test_evidence_rejects_invalid_worksheet_index(worksheet_index: object) -> None:
    with pytest.raises(ValueError):
        IdentifierColumnEvidence(
            worksheet_index=worksheet_index,  # type: ignore[arg-type]
            column=1,
            identifier_type=IdentifierType.INN,
            expected_count=0,
        )


@pytest.mark.parametrize("column", [0, -1, True])
def test_evidence_rejects_invalid_column(column: object) -> None:
    with pytest.raises(ValueError):
        IdentifierColumnEvidence(
            worksheet_index=1,
            column=column,  # type: ignore[arg-type]
            identifier_type=IdentifierType.INN,
            expected_count=0,
        )


def test_evidence_rejects_wrong_identifier_type() -> None:
    with pytest.raises(ValueError):
        IdentifierColumnEvidence(
            worksheet_index=1, column=1, identifier_type="inn", expected_count=0  # type: ignore[arg-type]
        )


@pytest.mark.parametrize("expected_count", [-1, True])
def test_evidence_rejects_invalid_expected_count(expected_count: object) -> None:
    with pytest.raises(ValueError):
        IdentifierColumnEvidence(
            worksheet_index=1,
            column=1,
            identifier_type=IdentifierType.INN,
            expected_count=expected_count,  # type: ignore[arg-type]
        )


def test_evidence_repr_contains_no_raw_values() -> None:
    evidence = IdentifierColumnEvidence(
        worksheet_index=1, column=1, identifier_type=IdentifierType.INN, expected_count=1
    )
    assert VALID_INN10 not in repr(evidence)


# ---------------------------------------------------------------------------
# B. scan_process_worksheet — Step A подсчёт
# ---------------------------------------------------------------------------


def test_zero_expected_for_all_blank_column() -> None:
    ws = _worksheet_with_column("INN", [None, "   ", ""])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(FieldType.INN, "INN")})
    assert result.identifier_evidence == ()


def test_one_expected() -> None:
    ws = _worksheet_with_column("INN", [VALID_INN10])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(FieldType.INN, "INN")})
    assert result.identifier_evidence == (
        IdentifierColumnEvidence(
            worksheet_index=1, column=1, identifier_type=IdentifierType.INN, expected_count=1
        ),
    )


def test_duplicates_counted_per_cell() -> None:
    ws = _worksheet_with_column("INN", [VALID_INN10, VALID_INN10, VALID_INN10])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(FieldType.INN, "INN")})
    assert result.identifier_evidence[0].expected_count == 3


def test_blank_cells_not_counted() -> None:
    ws = _worksheet_with_column("INN", [VALID_INN10, None, "   ", VALID_INN10])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(FieldType.INN, "INN")})
    assert result.identifier_evidence[0].expected_count == 2


def test_invalid_identifier_like_text_still_counted() -> None:
    # Непустое, но невалидное значение всё равно СЧИТАЕТСЯ: если бы Stage 8
    # реально обработал такой workbook, anonymizer поднял бы
    # InvalidIdentifierValueError раньше, чем Step B когда-либо был бы
    # достигнут — рассогласованного evidence возникнуть не может.
    ws = _worksheet_with_column("INN", ["not-an-identifier"])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(FieldType.INN, "INN")})
    assert result.identifier_evidence[0].expected_count == 1


@pytest.mark.parametrize(
    "field_type,value",
    [
        (FieldType.INN, VALID_INN10),
        (FieldType.INN, VALID_INN12),
        (FieldType.KPP, VALID_KPP),
        (FieldType.OGRN, VALID_OGRN13),
        (FieldType.OGRN, VALID_OGRNIP15),
    ],
)
def test_identifier_types_counted(field_type: FieldType, value: str) -> None:
    ws = _worksheet_with_column("Col", [value])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(field_type, "Col")})
    assert len(result.identifier_evidence) == 1
    assert result.identifier_evidence[0].expected_count == 1


def test_ogrn_and_ogrnip_share_identifier_type() -> None:
    ws = _worksheet_with_column("Col", [VALID_OGRN13, VALID_OGRNIP15])
    result = scan_process_worksheet(ws, 1, {1: _pseudonymize_rule(FieldType.OGRN, "Col")})
    assert len(result.identifier_evidence) == 1
    assert result.identifier_evidence[0].identifier_type is IdentifierType.OGRN
    assert result.identifier_evidence[0].expected_count == 2


def test_multiple_identifier_columns_on_one_sheet() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "INN"
    ws["B1"] = "KPP"
    ws["A2"] = VALID_INN10
    ws["B2"] = VALID_KPP
    result = scan_process_worksheet(
        ws,
        1,
        {
            1: _pseudonymize_rule(FieldType.INN, "INN"),
            2: _pseudonymize_rule(FieldType.KPP, "KPP"),
        },
    )
    assert len(result.identifier_evidence) == 2
    assert result.identifier_evidence[0].column == 1
    assert result.identifier_evidence[1].column == 2


def test_deterministic_ordering_by_worksheet_index_then_column() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "KPP"
    ws["B1"] = "INN"
    ws["A2"] = VALID_KPP
    ws["B2"] = VALID_INN10
    result = scan_process_worksheet(
        ws,
        1,
        {
            2: _pseudonymize_rule(FieldType.INN, "INN"),
            1: _pseudonymize_rule(FieldType.KPP, "KPP"),
        },
    )
    columns = [entry.column for entry in result.identifier_evidence]
    assert columns == sorted(columns)


# ---------------------------------------------------------------------------
# C. prepare_workbook_for_anonymization — агрегация по нескольким листам,
#    REMOVE_ALL исключён из evidence
# ---------------------------------------------------------------------------


def _save_and_close(wb: openpyxl.Workbook, path: Path) -> None:
    wb.save(path)
    wb.close()


def test_multiple_process_worksheets_aggregate_evidence(tmp_xlsx_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "INN"
    ws1["A2"] = VALID_INN10
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "KPP"
    ws2["A2"] = VALID_KPP
    _save_and_close(wb, tmp_xlsx_path)

    policies = {"First": WorksheetPolicy.PROCESS, "Second": WorksheetPolicy.PROCESS}
    rules = {
        "First": {1: _pseudonymize_rule(FieldType.INN, "INN")},
        "Second": {1: _pseudonymize_rule(FieldType.KPP, "KPP")},
    }
    with prepare_workbook_for_anonymization(tmp_xlsx_path, policies, rules) as result:
        assert len(result.identifier_evidence) == 2
        assert {e.worksheet_index for e in result.identifier_evidence} == {1, 2}


def test_remove_all_worksheet_excluded_from_evidence(tmp_xlsx_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data"
    ws1["A1"] = "INN"
    ws1["A2"] = VALID_INN10
    ws2 = wb.create_sheet("Junk")
    ws2["A1"] = "KPP"
    ws2["A2"] = VALID_KPP
    _save_and_close(wb, tmp_xlsx_path)

    policies = {"Data": WorksheetPolicy.PROCESS, "Junk": WorksheetPolicy.REMOVE_ALL}
    rules = {"Data": {1: _pseudonymize_rule(FieldType.INN, "INN")}}
    with prepare_workbook_for_anonymization(tmp_xlsx_path, policies, rules) as result:
        assert len(result.identifier_evidence) == 1
        assert result.identifier_evidence[0].identifier_type is IdentifierType.INN
