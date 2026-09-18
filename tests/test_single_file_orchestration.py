"""
Тесты Single-file anonymization orchestration (Stage 8.3) —
app.orchestration.single_file.anonymize_workbook.

Все workbook — synthetic (никаких реальных компаний/ФИО); ИНН — валидные
по чек-сумме, но синтетические (VALID_INN/VALID_INN_2, те же значения,
что уже используются в tests/test_identifier_anonymizer.py).

mapping_store/identifier_store переиспользуются между несколькими
вызовами anonymize_workbook в персистентных тестах — это то, что и
доказывает главный бизнес-инвариант Stage 8.3: pseudonym/token
stability зависит от persistent store, а не от job_id.
"""

from __future__ import annotations

import dataclasses
import os
from pathlib import Path

import openpyxl
import pytest

import app.orchestration.single_file as single_file_module
from app.anonymizer import InvalidPseudonymizationValueError
from app.excel.models import CellRecord, FlatTable
from app.excel.writer import WriterValidationError
from app.excel.writer import write_anonymized_workbook as real_write_anonymized_workbook
from app.mapping.identifier_base import IdentifierMappingConflictError
from app.mapping.identifier_memory import InMemoryIdentifierMappingStore
from app.mapping.memory import InMemoryMappingStore
from app.mapping.provenance_encrypted import EncryptedFileProvenanceStore
from app.models.identifiers import IdentifierMappingEntry, IdentifierType
from app.models.rules import Action, FieldRule, FieldType
from app.orchestration.single_file import (
    AnonymizationJobResult,
    OrchestrationValidationError,
    anonymize_workbook,
)

VALID_INN = "7707083893"
VALID_INN_2 = "7707083886"
PASSWORD = "test-provenance-password"


# ---------------------------------------------------------------------------
# Фикстуры/помощники
# ---------------------------------------------------------------------------


def _source_workbook(
    path: Path, sheet_name: str, headers: list[object], rows: list[list[object]]
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(path)


def _company_rule(column_name: str = "Company") -> FieldRule:
    return FieldRule(column_name=column_name, field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)


def _inn_rule(column_name: str = "INN") -> FieldRule:
    return FieldRule(column_name=column_name, field_type=FieldType.INN, action=Action.PSEUDONYMIZE)


@pytest.fixture()
def mapping_store() -> InMemoryMappingStore:
    return InMemoryMappingStore()


@pytest.fixture()
def identifier_store() -> InMemoryIdentifierMappingStore:
    return InMemoryIdentifierMappingStore()


@pytest.fixture()
def source_path(tmp_path: Path) -> Path:
    p = tmp_path / "source.xlsx"
    _source_workbook(p, "Data", ["Company"], [["ООО Ромашка"]])
    return p


@pytest.fixture()
def destination_path(tmp_path: Path) -> Path:
    return tmp_path / "output.xlsx"


@pytest.fixture()
def provenance_path(tmp_path: Path) -> Path:
    return tmp_path / "provenance.enc"


def _make_call_counters(monkeypatch: pytest.MonkeyPatch) -> dict[str, int]:
    """Оборачивает все leaf-вызовы orchestration-слоя счётчиками поверх реальных функций."""
    counts = {"read": 0, "token_hex": 0, "provenance_ctor": 0, "anonymize": 0, "writer": 0}

    real_read = single_file_module.read_flat_table

    def counting_read(*args: object, **kwargs: object) -> object:
        counts["read"] += 1
        return real_read(*args, **kwargs)

    monkeypatch.setattr(single_file_module, "read_flat_table", counting_read)

    real_token_hex = single_file_module.secrets.token_hex

    def counting_token_hex(*args: object, **kwargs: object) -> object:
        counts["token_hex"] += 1
        return real_token_hex(*args, **kwargs)

    monkeypatch.setattr(single_file_module.secrets, "token_hex", counting_token_hex)

    real_provenance_ctor = single_file_module.EncryptedFileProvenanceStore

    def counting_provenance_ctor(*args: object, **kwargs: object) -> object:
        counts["provenance_ctor"] += 1
        return real_provenance_ctor(*args, **kwargs)

    monkeypatch.setattr(single_file_module, "EncryptedFileProvenanceStore", counting_provenance_ctor)

    real_anonymize = single_file_module.anonymize_flat_table

    def counting_anonymize(*args: object, **kwargs: object) -> object:
        counts["anonymize"] += 1
        return real_anonymize(*args, **kwargs)

    monkeypatch.setattr(single_file_module, "anonymize_flat_table", counting_anonymize)

    real_writer = single_file_module.write_anonymized_workbook

    def counting_writer(*args: object, **kwargs: object) -> object:
        counts["writer"] += 1
        return real_writer(*args, **kwargs)

    monkeypatch.setattr(single_file_module, "write_anonymized_workbook", counting_writer)

    return counts


def _assert_zero_job_side_effects(counts: dict[str, int]) -> None:
    assert counts["token_hex"] == 0
    assert counts["provenance_ctor"] == 0
    assert counts["anonymize"] == 0
    assert counts["writer"] == 0


# ---------------------------------------------------------------------------
# Preflight — типы путей
# ---------------------------------------------------------------------------


def test_invalid_source_type_raises_type_error(
    destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(TypeError):
        anonymize_workbook(
            123, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_invalid_destination_type_raises_type_error(
    source_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(TypeError):
        anonymize_workbook(
            source_path, 123, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_invalid_provenance_path_type_raises_type_error(
    source_path: Path, destination_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(TypeError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=123, provenance_password=PASSWORD,
        )


# ---------------------------------------------------------------------------
# Preflight — расширения
# ---------------------------------------------------------------------------


def test_unsupported_source_extension_rejected(
    tmp_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    bad_source = tmp_path / "source.xlsm"
    bad_source.write_bytes(b"not a real workbook")
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            bad_source, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_unsupported_destination_extension_rejected(
    source_path: Path, tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, tmp_path / "output.xlsm", {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


# ---------------------------------------------------------------------------
# Preflight — source существование/тип файла
# ---------------------------------------------------------------------------


def test_missing_source_rejected(
    tmp_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            tmp_path / "does_not_exist.xlsx", destination_path, {"Data": {1: _company_rule()}},
            mapping_store, None, provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_source_directory_rejected(
    tmp_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    dir_as_source = tmp_path / "dir.xlsx"
    dir_as_source.mkdir()
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            dir_as_source, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


# ---------------------------------------------------------------------------
# Preflight — родительские каталоги
# ---------------------------------------------------------------------------


def test_missing_destination_parent_rejected(
    source_path: Path, tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    bad_dest = tmp_path / "no_such_dir" / "output.xlsx"
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, bad_dest, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert not (tmp_path / "no_such_dir").exists()


def test_missing_provenance_parent_rejected(
    source_path: Path, destination_path: Path, tmp_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    bad_provenance = tmp_path / "no_such_dir" / "provenance.enc"
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=bad_provenance, provenance_password=PASSWORD,
        )
    assert not (tmp_path / "no_such_dir").exists()


# ---------------------------------------------------------------------------
# Preflight — равенство путей
# ---------------------------------------------------------------------------


def test_source_equal_destination_rejected(
    source_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, source_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_provenance_equal_source_rejected(
    source_path: Path, destination_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=source_path, provenance_password=PASSWORD,
        )


def test_provenance_equal_destination_rejected(
    source_path: Path, destination_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=destination_path, provenance_password=PASSWORD,
        )


# ---------------------------------------------------------------------------
# Preflight — provenance collision
# ---------------------------------------------------------------------------


def test_provenance_already_exists_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    provenance_path.write_bytes(b"existing sidecar content, must not be touched")
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert provenance_path.read_bytes() == b"existing sidecar content, must not be touched"


# ---------------------------------------------------------------------------
# Preflight — структура rules_by_sheet
# ---------------------------------------------------------------------------


def test_non_mapping_rules_by_sheet_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, "not-a-mapping", mapping_store, None,  # type: ignore[arg-type]
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_empty_rules_by_sheet_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_non_str_sheet_name_key_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {123: {1: _company_rule()}}, mapping_store, None,  # type: ignore[dict-item]
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_whitespace_only_sheet_name_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"   ": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_non_mapping_column_rules_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": "not-a-mapping"}, mapping_store, None,  # type: ignore[dict-item]
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_bool_column_index_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    # bool — подкласс int; existing anonymizer._validate_rules отвергает
    # его тем же способом — preflight обязан обнаружить это раньше.
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {True: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_non_int_column_index_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {"one": _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


def test_non_field_rule_value_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: "not-a-field-rule"}}, mapping_store, None,  # type: ignore[dict-item]
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )


# ---------------------------------------------------------------------------
# Preflight — password
# ---------------------------------------------------------------------------


def test_non_str_password_raises_type_error(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(TypeError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=12345,  # type: ignore[arg-type]
        )


def test_type_error_message_does_not_include_password_value(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(TypeError) as exc_info:
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=999999999,  # type: ignore[arg-type]
        )
    assert "999999999" not in str(exc_info.value)


def test_empty_password_rejected(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password="",
        )


def test_whitespace_only_password_accepted(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    # crypto contract принимает whitespace-only пароль — preflight не
    # должен вводить дополнительное .strip()-правило.
    result = anonymize_workbook(
        source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
        provenance_path=provenance_path, provenance_password="   ",
    )
    assert isinstance(result, AnonymizationJobResult)


# ---------------------------------------------------------------------------
# Preflight side-effects — Reader/job_id/provenance/anonymizer/Writer НЕ вызваны
# ---------------------------------------------------------------------------


def test_preflight_source_equal_destination_has_zero_side_effects(
    source_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore, monkeypatch: pytest.MonkeyPatch
) -> None:
    counts = _make_call_counters(monkeypatch)
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, source_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert counts["read"] == 0
    _assert_zero_job_side_effects(counts)
    assert mapping_store.entries() == ()


def test_preflight_provenance_exists_has_zero_side_effects(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    provenance_path.write_bytes(b"pre-existing")
    counts = _make_call_counters(monkeypatch)
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert counts["read"] == 0
    _assert_zero_job_side_effects(counts)
    assert mapping_store.entries() == ()


def test_preflight_empty_rules_has_zero_side_effects(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    counts = _make_call_counters(monkeypatch)
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert counts["read"] == 0
    _assert_zero_job_side_effects(counts)


def test_preflight_invalid_password_has_zero_side_effects(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    counts = _make_call_counters(monkeypatch)
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password="",
        )
    assert counts["read"] == 0
    _assert_zero_job_side_effects(counts)


def test_preflight_missing_parent_has_zero_side_effects(
    source_path: Path, tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    counts = _make_call_counters(monkeypatch)
    bad_dest = tmp_path / "no_such_dir" / "output.xlsx"
    with pytest.raises(OrchestrationValidationError):
        anonymize_workbook(
            source_path, bad_dest, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert counts["read"] == 0
    _assert_zero_job_side_effects(counts)


# ---------------------------------------------------------------------------
# Read-all-before-job
# ---------------------------------------------------------------------------


def test_read_failure_on_second_sheet_has_zero_job_side_effects(
    tmp_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "multi_sheet_source.xlsx"
    _source_workbook(source, "SheetA", ["Company"], [["ООО Ромашка"]])

    counts = _make_call_counters(monkeypatch)
    rules = {
        "SheetA": {1: _company_rule()},
        "SheetB": {1: _company_rule()},  # SheetB отсутствует в книге
    }
    with pytest.raises(KeyError):
        anonymize_workbook(
            source, destination_path, rules, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    # SheetA прочитана успешно, попытка прочитать SheetB упала —
    # оба вызова были сделаны (2), но ни job_id, ни provenance, ни
    # anonymizer, ни Writer не были достигнуты.
    assert counts["read"] == 2
    _assert_zero_job_side_effects(counts)
    assert mapping_store.entries() == ()
    assert not provenance_path.exists()


# ---------------------------------------------------------------------------
# Success — single sheet / multi-sheet
# ---------------------------------------------------------------------------


def test_single_sheet_end_to_end_success(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    result = anonymize_workbook(
        source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
        provenance_path=provenance_path, provenance_password=PASSWORD,
    )

    assert isinstance(result, AnonymizationJobResult)
    assert result.output_path == destination_path
    assert result.provenance_path == provenance_path
    assert result.processed_sheet_names == ("Data",)

    assert destination_path.exists()
    assert provenance_path.exists()

    output = openpyxl.load_workbook(destination_path)
    assert output["Data"]["A2"].value.startswith("C_")
    assert output.custom_doc_props["DataAnonymizer.JobId"].value == result.job_id

    # source остаётся неизменным
    reopened_source = openpyxl.load_workbook(source_path)
    assert reopened_source["Data"]["A2"].value == "ООО Ромашка"


def test_multi_sheet_end_to_end_success(
    tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    source = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "Company"
    ws1["A2"] = "ООО Ромашка"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "Company"
    ws2["A2"] = "ООО Ромашка"  # та же компания на другом листе
    wb.save(source)

    destination = tmp_path / "output.xlsx"
    rules = {"First": {1: _company_rule()}, "Second": {1: _company_rule()}}

    result = anonymize_workbook(
        source, destination, rules, mapping_store, None,
        provenance_path=provenance_path, provenance_password=PASSWORD,
    )

    assert result.processed_sheet_names == ("First", "Second")

    output = openpyxl.load_workbook(destination)
    alias_first = output["First"]["A2"].value
    alias_second = output["Second"]["A2"].value
    assert alias_first == alias_second
    assert len(mapping_store.entries()) == 1


# ---------------------------------------------------------------------------
# Persistence — главный бизнес-инвариант Stage 8.3
# ---------------------------------------------------------------------------


def test_same_company_across_three_jobs_gets_same_alias(
    tmp_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    company = "ООО Ромашка"
    job_ids: list[str] = []
    provenance_paths: list[Path] = []
    aliases: list[object] = []

    for month in ("january", "february", "march"):
        source = tmp_path / f"{month}_source.xlsx"
        _source_workbook(source, "Data", ["Company"], [[company]])
        destination = tmp_path / f"{month}_output.xlsx"
        provenance = tmp_path / f"{month}_provenance.enc"

        result = anonymize_workbook(
            source, destination, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance, provenance_password=PASSWORD,
        )
        job_ids.append(result.job_id)
        provenance_paths.append(result.provenance_path)
        aliases.append(openpyxl.load_workbook(destination)["Data"]["A2"].value)

    assert aliases[0] == aliases[1] == aliases[2]
    assert len(set(job_ids)) == 3
    assert len(set(provenance_paths)) == 3
    assert len(mapping_store.entries()) == 1  # одна identity, не три


def test_same_identifier_across_two_jobs_gets_same_token(
    tmp_path: Path, mapping_store: InMemoryMappingStore, identifier_store: InMemoryIdentifierMappingStore
) -> None:
    job_ids: list[str] = []
    tokens: list[object] = []

    for suffix in ("a", "b"):
        source = tmp_path / f"source_{suffix}.xlsx"
        _source_workbook(source, "Data", ["INN"], [[VALID_INN]])
        destination = tmp_path / f"output_{suffix}.xlsx"
        provenance = tmp_path / f"provenance_{suffix}.enc"

        result = anonymize_workbook(
            source, destination, {"Data": {1: _inn_rule()}}, mapping_store, identifier_store,
            provenance_path=provenance, provenance_password=PASSWORD,
        )
        job_ids.append(result.job_id)
        tokens.append(openpyxl.load_workbook(destination)["Data"]["A2"].value)

    assert tokens[0] == tokens[1]
    assert job_ids[0] != job_ids[1]
    assert len(identifier_store.entries()) == 1


def test_output_job_id_matches_provenance_job_id(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    result = anonymize_workbook(
        source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
        provenance_path=provenance_path, provenance_password=PASSWORD,
    )

    output = openpyxl.load_workbook(destination_path)
    output_job_id = output.custom_doc_props["DataAnonymizer.JobId"].value

    reopened_provenance = EncryptedFileProvenanceStore(provenance_path, PASSWORD)

    assert output_job_id == result.job_id
    assert reopened_provenance.job_id == result.job_id


# ---------------------------------------------------------------------------
# Failure semantics
# ---------------------------------------------------------------------------


def test_anonymizer_failure_leaves_no_output_but_sidecar_exists(
    tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    source = tmp_path / "source.xlsx"
    _source_workbook(source, "Data", ["Company"], [[42]])  # невалидное значение под COMPANY
    destination = tmp_path / "output.xlsx"

    with pytest.raises(InvalidPseudonymizationValueError):
        anonymize_workbook(
            source, destination, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    assert not destination.exists()
    # sidecar уже создан к этому моменту (после job_id, до anonymize) —
    # допустимый независимый artifact, не удаляется.
    assert provenance_path.exists()
    reopened = EncryptedFileProvenanceStore(provenance_path, PASSWORD)
    assert reopened.entries() == ()


def test_identifier_add_many_failure_propagates(
    tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    identifier_store: InMemoryIdentifierMappingStore, monkeypatch: pytest.MonkeyPatch,
) -> None:
    import app.anonymizer as anonymizer_module

    identifier_store.add(
        IdentifierMappingEntry(
            token="INN_FIXEDTOKEN", identifier_value=VALID_INN_2, identifier_type=IdentifierType.INN
        )
    )

    def fake_generate_identifier_token(identifier_type: object, existing_tokens: object) -> str:
        return "INN_FIXEDTOKEN"

    monkeypatch.setattr(anonymizer_module, "generate_identifier_token", fake_generate_identifier_token)

    source = tmp_path / "source.xlsx"
    _source_workbook(source, "Data", ["INN"], [[VALID_INN]])  # новая identity, token уже занят
    destination = tmp_path / "output.xlsx"

    with pytest.raises(IdentifierMappingConflictError):
        anonymize_workbook(
            source, destination, {"Data": {1: _inn_rule()}}, mapping_store, identifier_store,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    assert not destination.exists()
    assert len(identifier_store.entries()) == 1  # исходная запись не пострадала
    assert provenance_path.exists()


def test_provenance_add_many_failure_after_identifier_success(
    tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    identifier_store: InMemoryIdentifierMappingStore, monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Точечный fault injection: первый os.replace (persist пустого
    # sidecar в конструкторе EncryptedFileProvenanceStore) проходит
    # успешно, второй (add_many после успешной анонимизации) падает.
    call_count = {"n": 0}
    real_replace = os.replace

    def selective_failing_replace(*args: object, **kwargs: object) -> object:
        call_count["n"] += 1
        if call_count["n"] == 1:
            return real_replace(*args, **kwargs)
        raise OSError("simulated provenance commit failure")

    monkeypatch.setattr(os, "replace", selective_failing_replace)

    # Прямое доказательство "Writer не вызван" (а не только косвенное
    # через отсутствие destination): реальная функция остаётся реальной
    # (вызов ей просто не должен быть достигнут), считается только факт
    # обращения к ней.
    writer_calls = {"n": 0}
    real_writer = single_file_module.write_anonymized_workbook

    def counting_writer(*args: object, **kwargs: object) -> object:
        writer_calls["n"] += 1
        return real_writer(*args, **kwargs)

    monkeypatch.setattr(single_file_module, "write_anonymized_workbook", counting_writer)

    source = tmp_path / "source.xlsx"
    _source_workbook(source, "Data", ["INN"], [[VALID_INN]])
    destination = tmp_path / "output.xlsx"

    with pytest.raises(OSError):
        anonymize_workbook(
            source, destination, {"Data": {1: _inn_rule()}}, mapping_store, identifier_store,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    monkeypatch.undo()

    assert writer_calls["n"] == 0
    assert not destination.exists()
    # Принятая асимметрия: identifier mapping уже durable.
    assert len(identifier_store.entries()) == 1
    assert provenance_path.exists()
    reopened = EncryptedFileProvenanceStore(provenance_path, PASSWORD)
    assert reopened.entries() == ()


def test_later_sheet_failure_leaves_earlier_sheet_committed(
    tmp_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    source = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "Company"
    ws1["A2"] = "ООО Ромашка"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "Company"
    ws2["A2"] = 42  # невалидное значение -> сбой на втором листе
    wb.save(source)

    destination = tmp_path / "output.xlsx"
    rules = {"First": {1: _company_rule()}, "Second": {1: _company_rule()}}

    with pytest.raises(InvalidPseudonymizationValueError):
        anonymize_workbook(
            source, destination, rules, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    assert not destination.exists()
    # Sheet First уже успел закоммититься (eager entity add) до сбоя на Second.
    assert len(mapping_store.entries()) == 1
    assert mapping_store.entries()[0].real_value == "ООО Ромашка"


def test_writer_validation_failure_end_to_end(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Реальный Writer, спровоцированный на реальный WriterValidationError:
    # к списку anonymized tables добавляется таблица для несуществующего
    # в source листа.
    def wrapped_writer(
        source_path_arg: object, tables: object, destination_path_arg: object, *, job_id: str
    ) -> None:
        bogus = FlatTable(
            sheet_name="GhostSheet",
            header_row=(CellRecord(row=1, column=1, value="X"),),
            rows=(),
        )
        return real_write_anonymized_workbook(
            source_path_arg, tuple(tables) + (bogus,), destination_path_arg, job_id=job_id
        )

    monkeypatch.setattr(single_file_module, "write_anonymized_workbook", wrapped_writer)

    with pytest.raises(WriterValidationError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    assert not destination_path.exists()
    assert provenance_path.exists()
    assert len(mapping_store.entries()) == 1  # Data-лист уже успешно анонимизирован


def test_writer_save_failure_end_to_end(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from openpyxl.workbook.workbook import Workbook

    def failing_save(self: Workbook, *args: object, **kwargs: object) -> None:
        raise OSError("simulated save failure")

    monkeypatch.setattr(Workbook, "save", failing_save)

    with pytest.raises(OSError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    monkeypatch.undo()
    assert not destination_path.exists()
    assert provenance_path.exists()
    assert len(mapping_store.entries()) == 1


def test_writer_fsync_failure_end_to_end(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # os.fsync глобален для процесса: первый вызов принадлежит созданию
    # sidecar (EncryptedFileProvenanceStore.__init__, до Writer) и должен
    # пройти успешно; падать должен только следующий вызов — Writer'а.
    call_count = {"n": 0}
    real_fsync = os.fsync

    def selective_failing_fsync(fd: int) -> None:
        call_count["n"] += 1
        if call_count["n"] == 1:
            return real_fsync(fd)
        raise OSError("simulated fsync failure")

    monkeypatch.setattr(os, "fsync", selective_failing_fsync)

    with pytest.raises(OSError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    monkeypatch.undo()
    assert not destination_path.exists()
    assert provenance_path.exists()


def test_writer_os_replace_failure_end_to_end(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Тот же приём: первый os.replace — sidecar creation, должен пройти;
    # второй — Writer'а, должен упасть.
    call_count = {"n": 0}
    real_replace = os.replace

    def selective_failing_replace(*args: object, **kwargs: object) -> object:
        call_count["n"] += 1
        if call_count["n"] == 1:
            return real_replace(*args, **kwargs)
        raise OSError("simulated os.replace failure")

    monkeypatch.setattr(os, "replace", selective_failing_replace)

    with pytest.raises(OSError):
        anonymize_workbook(
            source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )

    monkeypatch.undo()
    assert not destination_path.exists()
    assert provenance_path.exists()


# ---------------------------------------------------------------------------
# Existing destination
# ---------------------------------------------------------------------------


def test_existing_destination_is_atomically_replaced(
    source_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "PRE-EXISTING DESTINATION CONTENT"
    wb.save(destination_path)

    result = anonymize_workbook(
        source_path, destination_path, {"Data": {1: _company_rule()}}, mapping_store, None,
        provenance_path=provenance_path, provenance_password=PASSWORD,
    )

    output = openpyxl.load_workbook(destination_path)
    assert "Data" in output.sheetnames
    assert "Sheet" not in output.sheetnames
    assert isinstance(result, AnonymizationJobResult)


# ---------------------------------------------------------------------------
# Security
# ---------------------------------------------------------------------------


def test_result_has_only_expected_non_confidential_fields() -> None:
    field_names = {f.name for f in dataclasses.fields(AnonymizationJobResult)}
    assert field_names == {"job_id", "output_path", "provenance_path", "processed_sheet_names"}


def test_orchestration_module_never_calls_clear_on_stores() -> None:
    source_text = Path(single_file_module.__file__).read_text(encoding="utf-8")
    assert ".clear(" not in source_text


def test_missing_sheet_error_does_not_expose_company_value(
    tmp_path: Path, destination_path: Path, provenance_path: Path, mapping_store: InMemoryMappingStore
) -> None:
    secret_looking_value = "SuperSecretCompanyName999"
    source = tmp_path / "source.xlsx"
    _source_workbook(source, "Data", ["Company"], [[secret_looking_value]])

    with pytest.raises(KeyError) as exc_info:
        anonymize_workbook(
            source, destination_path,
            {"Data": {1: _company_rule()}, "NoSuchSheet": {1: _company_rule()}},
            mapping_store, None,
            provenance_path=provenance_path, provenance_password=PASSWORD,
        )
    assert secret_looking_value not in str(exc_info.value)
