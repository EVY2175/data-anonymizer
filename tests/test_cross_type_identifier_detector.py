"""
Тесты Stage 10C.1 — cross-type identifier detector и identifier KEEP
emptiness contract (app.coverage.worksheet_policy.scan_process_worksheet,
keep_violations).

Тестовые векторы ИНН/КПП/ОГРН/ОГРНИП с корректной контрольной суммой —
те же самые, уже проверенные векторы, что и tests/test_detectors.py (не
изобретаются заново).
"""

from __future__ import annotations

import datetime

import openpyxl
import pytest

from app.coverage.errors import ColumnSafetyReason
from app.coverage.worksheet_policy import scan_process_worksheet
from app.models.rules import Action, FieldRule, FieldType

VALID_INN10 = "7707083893"
VALID_OGRN13 = "1047709000008"
VALID_OGRNIP15 = "304770900000126"
VALID_KPP = "770701001"


def _worksheet_with_value(header: str, value: object):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = header
    ws["A2"] = value
    return ws


def _keep_rule(field_type: FieldType, column_name: str = "Column") -> FieldRule:
    return FieldRule(column_name=column_name, field_type=field_type, action=Action.KEEP)


# ---------------------------------------------------------------------------
# A. FINANCIAL/BRANCH/DEPARTMENT + KEEP + identifier-подобное значение
# ---------------------------------------------------------------------------


def test_financial_keep_with_inn_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_INN10)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_financial_keep_with_kpp_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_KPP)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_branch_keep_with_inn_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_INN10)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.BRANCH, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_department_keep_with_ogrn_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_OGRN13)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.DEPARTMENT, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_department_keep_with_ogrnip_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_OGRNIP15)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.DEPARTMENT, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


# ---------------------------------------------------------------------------
# B. INN/KPP/OGRN + KEEP (§11 identifier KEEP emptiness contract)
# ---------------------------------------------------------------------------


def test_inn_keep_with_valid_inn_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_INN10)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.INN, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_kpp_keep_with_valid_kpp_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_KPP)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.KPP, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_ogrn_keep_with_valid_ogrn_rejected() -> None:
    ws = _worksheet_with_value("Col", VALID_OGRN13)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.OGRN, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_inn_keep_with_blank_column_allowed() -> None:
    ws = _worksheet_with_value("Col", None)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.INN, "Col")})
    assert result.keep_violations == ()


def test_inn_keep_with_whitespace_only_allowed() -> None:
    ws = _worksheet_with_value("Col", "   ")
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.INN, "Col")})
    assert result.keep_violations == ()


def test_inn_keep_with_ordinary_text_allowed() -> None:
    ws = _worksheet_with_value("Col", "не идентификатор")
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.INN, "Col")})
    assert result.keep_violations == ()


# ---------------------------------------------------------------------------
# C. Типовая семантика детектора (int/str/blank/invalid)
# ---------------------------------------------------------------------------


def test_detector_accepts_int_representation() -> None:
    ws = _worksheet_with_value("Col", int(VALID_INN10))
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_detector_accepts_str_representation() -> None:
    ws = _worksheet_with_value("Col", VALID_INN10)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ((1, ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED),)


def test_detector_rejects_float_representation() -> None:
    # float ВСЕГДА отвергается детекторами Stage 5, даже если цифры
    # совпадают с валидным ИНН — та же type-policy, что и в
    # app.detectors (см. tests/test_detectors.py).
    ws = _worksheet_with_value("Col", float(VALID_INN10))
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ()


def test_detector_rejects_bool_representation() -> None:
    ws = _worksheet_with_value("Col", True)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ()


def test_detector_ignores_invalid_identifier_like_text() -> None:
    # Длина не совпадает ни с одним форматом (ИНН: 10/12, ОГРН: 13,
    # ОГРНИП: 15, КПП: 9) — ни один из четырёх валидаторов не сработает.
    ws = _worksheet_with_value("Col", "12345")
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ()


def test_detector_ignores_invalid_inn_checksum() -> None:
    # 10 ASCII-цифр (правильная длина ИНН), но некорректная контрольная
    # сумма — не идентификатор ни по одному из четырёх форматов.
    ws = _worksheet_with_value("Col", "1234567890")
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ()


@pytest.mark.parametrize(
    "value",
    [
        datetime.datetime(2024, 1, 1, 12, 0, 0),
        datetime.date(2024, 1, 1),
        datetime.time(12, 0, 0),
    ],
)
def test_detector_safely_ignores_real_excel_date_time_values(value: object) -> None:
    # datetime/date/time — реально допустимые в ячейке Excel/openpyxl типы
    # (в отличие от list/dict/object, которые openpyxl не примет как
    # cell.value вообще). Ни один из четырёх Stage5-валидаторов не должен
    # ни считать их identifier-подобными, ни поднять исключение.
    ws = _worksheet_with_value("Col", value)
    result = scan_process_worksheet(ws, 1, {1: _keep_rule(FieldType.FINANCIAL, "Col")})
    assert result.keep_violations == ()


@pytest.mark.parametrize("field_type", [FieldType.PHONE, FieldType.EMAIL, FieldType.ADDRESS])
def test_no_new_phone_email_address_detector_needed(field_type: FieldType) -> None:
    # PHONE/EMAIL/ADDRESS + KEEP запрещены на уровне validate_column_action
    # (KEEP_FORBIDDEN_FOR_TYPE) раньше, чем дело доходит до сканирования
    # содержимого — scan_process_worksheet для них никогда не вызывается
    # с KEEP-правилом в реальном workflow. Здесь фиксируется, что даже если
    # бы это произошло, cross-type detector для них не задействован
    # (_KEEP_REQUIRES_DETECTOR_SCAN не включает эти типы) — новый
    # phone/email/address-валидатор не создаётся.
    ws = _worksheet_with_value("Col", VALID_INN10)
    rule = FieldRule(column_name="Col", field_type=field_type, action=Action.KEEP)
    result = scan_process_worksheet(ws, 1, {1: rule})
    assert result.keep_violations == ()
