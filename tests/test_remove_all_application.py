"""
Тесты Stage 10C.1 — formula policy, REMOVE_ALL application и temp XLSX
lifecycle (app.coverage.worksheet_policy.prepare_workbook_for_anonymization).

Целевая среда проекта — Windows + PowerShell + Python 3.13, поэтому
отдельно проверяется, что source-файл можно немедленно переименовать/
удалить сразу после возврата из prepare_workbook_for_anonymization (что
на Windows невозможно при незакрытом файловом handle на source_path).
Это подтверждает, что ВЕСЬ pipeline в целом (openpyxl.load_workbook,
scan_process_worksheet, workbook.save во временный путь, workbook.close())
не удерживает файловый handle на исходном файле ни в одном из проверяемых
исходов — а не специфически действие workbook.close(): для режима загрузки
read_only=False оно документированно является no-op (openpyxl.Workbook.
close: "Only affects read-only and write-only modes"; hasattr(wb, "_archive")
== False для обычно загруженного workbook), и openpyxl.load_workbook сам
полностью читает файл и освобождает свой внутренний handle до возврата —
это подтверждено отдельной проверкой при независимом review.
"""

from __future__ import annotations

import hashlib
import os
import tempfile as _tempfile_module
from pathlib import Path

import openpyxl
import pytest

from app.coverage.errors import CoverageError, CoverageReason
from app.coverage.models import WorksheetPolicy
from app.coverage.worksheet_policy import prepare_workbook_for_anonymization
from app.models.rules import Action, FieldRule, FieldType

VALID_INN10 = "7707083893"
PASSWORD_SENTINEL = "SuperSecretPassword123!"


def _build_source(
    path: Path,
    *,
    formula_on: str | None = None,
    extra_sheet: bool = False,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Company"
    ws["A2"] = "original-company"
    if formula_on == "Data":
        ws["B1"] = "Formula"
        ws["B2"] = "=A2"
    if extra_sheet:
        junk = wb.create_sheet("Junk")
        junk["A1"] = "Whatever"
        junk["A2"] = "some-value"
        if formula_on == "Junk":
            junk["B1"] = "Formula"
            junk["B2"] = "=A2"
    wb.save(path)
    wb.close()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


_PROCESS_RULES = {"Data": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}


@pytest.fixture()
def source_path(tmp_path: Path) -> Path:
    return tmp_path / "source.xlsx"


# ---------------------------------------------------------------------------
# A. Formula policy
# ---------------------------------------------------------------------------


def test_formula_on_process_worksheet_rejected(source_path: Path) -> None:
    _build_source(source_path, formula_on="Data")
    rules = {
        "Data": {
            1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE),
            2: FieldRule(column_name="Formula", field_type=FieldType.FINANCIAL, action=Action.KEEP),
        }
    }
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is CoverageReason.FORMULA_REJECTED


def test_formula_on_remove_all_worksheet_allowed(source_path: Path) -> None:
    _build_source(source_path, formula_on="Junk", extra_sheet=True)
    policies = {"Data": WorksheetPolicy.PROCESS, "Junk": WorksheetPolicy.REMOVE_ALL}
    with prepare_workbook_for_anonymization(source_path, policies, _PROCESS_RULES) as result:
        wb = openpyxl.load_workbook(result.prepared_path)
        assert wb.sheetnames == ["Data"]
        wb.close()


def test_all_remove_all_rejected_before_mutation(source_path: Path) -> None:
    _build_source(source_path, extra_sheet=True)
    before = _sha256(source_path)
    policies = {"Data": WorksheetPolicy.REMOVE_ALL, "Junk": WorksheetPolicy.REMOVE_ALL}
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, policies, {}):
            pass
    assert exc_info.value.reason is CoverageReason.NO_PROCESS_WORKSHEET_REMAINS
    assert _sha256(source_path) == before


# ---------------------------------------------------------------------------
# B. REMOVE_ALL physical application
# ---------------------------------------------------------------------------


def test_remove_all_worksheet_physically_absent_from_prepared_workbook(source_path: Path) -> None:
    _build_source(source_path, extra_sheet=True)
    policies = {"Data": WorksheetPolicy.PROCESS, "Junk": WorksheetPolicy.REMOVE_ALL}
    with prepare_workbook_for_anonymization(source_path, policies, _PROCESS_RULES) as result:
        wb = openpyxl.load_workbook(result.prepared_path)
        assert "Junk" not in wb.sheetnames
        assert wb.sheetnames == ["Data"]
        wb.close()


# ---------------------------------------------------------------------------
# C. RAW неизменность (SHA до/после)
# ---------------------------------------------------------------------------


def test_raw_sha_unchanged_on_success(source_path: Path) -> None:
    _build_source(source_path)
    before = _sha256(source_path)
    with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
        pass
    assert _sha256(source_path) == before


def test_raw_sha_unchanged_on_validation_failure(source_path: Path) -> None:
    _build_source(source_path, extra_sheet=True)
    before = _sha256(source_path)
    with pytest.raises(CoverageError):
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass
    assert _sha256(source_path) == before


def test_raw_sha_unchanged_on_save_failure(source_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _build_source(source_path)
    before = _sha256(source_path)

    def _boom(self, filename):  # noqa: ANN001
        raise RuntimeError("simulated save failure")

    monkeypatch.setattr(openpyxl.Workbook, "save", _boom)

    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass
    assert exc_info.value.reason is CoverageReason.WORKBOOK_MUTATION_FAILED
    assert _sha256(source_path) == before


# ---------------------------------------------------------------------------
# D. Temp XLSX lifecycle
# ---------------------------------------------------------------------------


def test_prepared_path_has_xlsx_suffix(source_path: Path) -> None:
    _build_source(source_path)
    with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES) as result:
        assert result.prepared_path.suffix.lower() == ".xlsx"


def test_prepared_path_opens_via_openpyxl(source_path: Path) -> None:
    _build_source(source_path)
    with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES) as result:
        wb = openpyxl.load_workbook(result.prepared_path)
        assert wb.sheetnames == ["Data"]
        wb.close()


def test_prepared_path_deleted_after_context_exit(source_path: Path) -> None:
    _build_source(source_path)
    with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES) as result:
        prepared = result.prepared_path
        assert prepared.exists()
    assert not prepared.exists()


def test_prepared_path_deleted_after_exception_in_with_body(source_path: Path) -> None:
    _build_source(source_path)
    prepared_holder: list[Path] = []
    with pytest.raises(ValueError):
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES) as result:
            prepared_holder.append(result.prepared_path)
            raise ValueError("boom from with-body")
    assert not prepared_holder[0].exists()


def test_partial_temp_deleted_after_save_failure(source_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _build_source(source_path)
    created_paths: list[str] = []
    real_mkstemp = __import__("tempfile").mkstemp

    def _tracking_mkstemp(*args, **kwargs):  # noqa: ANN002, ANN003
        fd, name = real_mkstemp(*args, **kwargs)
        created_paths.append(name)
        return fd, name

    monkeypatch.setattr("app.coverage.worksheet_policy.tempfile.mkstemp", _tracking_mkstemp)

    def _boom(self, filename):  # noqa: ANN001
        raise RuntimeError("simulated save failure")

    monkeypatch.setattr(openpyxl.Workbook, "save", _boom)

    with pytest.raises(CoverageError):
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass

    assert len(created_paths) == 1
    assert not Path(created_paths[0]).exists()


def test_partial_temp_deleted_after_fsync_failure(source_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _build_source(source_path)
    created_paths: list[str] = []
    real_mkstemp = __import__("tempfile").mkstemp

    def _tracking_mkstemp(*args, **kwargs):  # noqa: ANN002, ANN003
        fd, name = real_mkstemp(*args, **kwargs)
        created_paths.append(name)
        return fd, name

    monkeypatch.setattr("app.coverage.worksheet_policy.tempfile.mkstemp", _tracking_mkstemp)

    def _boom_fsync(fd):  # noqa: ANN001
        raise OSError("simulated fsync failure")

    monkeypatch.setattr("app.coverage.worksheet_policy.os.fsync", _boom_fsync)

    with pytest.raises(CoverageError):
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass

    assert len(created_paths) == 1
    assert not Path(created_paths[0]).exists()


def test_source_can_be_renamed_immediately_after_success(source_path: Path, tmp_path: Path) -> None:
    # Windows-специфичная проверка: если бы pipeline в целом (не конкретно
    # workbook.close(), который для read_only=False — no-op) где-то держал
    # файловый handle на source_path открытым, os.replace ниже упал бы с
    # PermissionError/WinError 32.
    _build_source(source_path)
    with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
        pass
    renamed = tmp_path / "renamed-after-success.xlsx"
    os.replace(str(source_path), str(renamed))
    assert renamed.exists()


def test_source_can_be_renamed_immediately_after_validation_failure(source_path: Path, tmp_path: Path) -> None:
    _build_source(source_path, extra_sheet=True)
    with pytest.raises(CoverageError):
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass
    renamed = tmp_path / "renamed-after-validation-failure.xlsx"
    os.replace(str(source_path), str(renamed))
    assert renamed.exists()


def test_source_can_be_renamed_immediately_after_save_failure(
    source_path: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _build_source(source_path)

    def _boom(self, filename):  # noqa: ANN001
        raise RuntimeError("simulated save failure")

    monkeypatch.setattr(openpyxl.Workbook, "save", _boom)

    with pytest.raises(CoverageError):
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass

    renamed = tmp_path / "renamed-after-save-failure.xlsx"
    os.replace(str(source_path), str(renamed))
    assert renamed.exists()


def test_no_extra_files_left_behind_besides_source(source_path: Path, tmp_path: Path) -> None:
    _build_source(source_path)
    before = set(os.listdir(tmp_path))
    with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
        pass
    after = set(os.listdir(tmp_path))
    assert after == before


# ---------------------------------------------------------------------------
# E. NO_VISIBLE_PROCESS_WORKSHEET (закрытие MAJOR-1 независимого review):
# pre-mutation проверка на полном pipeline, а не только на save-failure.
# ---------------------------------------------------------------------------


def _build_visible_remove_all_hidden_process_source(path: Path, *, hidden_state: str) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Visible"
    ws1["A1"] = "Whatever"
    ws1["A2"] = "some-value"
    ws2 = wb.create_sheet("Hidden")
    ws2.sheet_state = hidden_state
    ws2["A1"] = "Company"
    ws2["A2"] = "original-company"
    wb.save(path)
    wb.close()


@pytest.mark.parametrize("hidden_state", ["hidden", "veryHidden"])
def test_no_visible_process_worksheet_raises_before_mutation(
    source_path: Path, hidden_state: str
) -> None:
    _build_visible_remove_all_hidden_process_source(source_path, hidden_state=hidden_state)
    before = _sha256(source_path)
    policies = {"Visible": WorksheetPolicy.REMOVE_ALL, "Hidden": WorksheetPolicy.PROCESS}
    rules = {"Hidden": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}

    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, policies, rules):
            pass

    assert exc_info.value.reason is CoverageReason.NO_VISIBLE_PROCESS_WORKSHEET
    # F: RAW SHA до/после одинаков.
    assert _sha256(source_path) == before


def test_no_visible_process_worksheet_does_not_call_apply_remove_all(
    source_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # G: доказательство PRE-MUTATION validation — _apply_remove_all (in-memory
    # мутация workbook.remove) не вызывается вообще, если ни один PROCESS-лист
    # не visible. Проверяется через spy на реальную функцию (не через mock,
    # который бы её подменил и тем самым исключил из проверки), поэтому её
    # реальное поведение (если бы она всё же была вызвана) осталось бы
    # действующим — тест доказывает именно порядок вызовов, а не подменяет
    # результат.
    _build_visible_remove_all_hidden_process_source(source_path, hidden_state="veryHidden")
    policies = {"Visible": WorksheetPolicy.REMOVE_ALL, "Hidden": WorksheetPolicy.PROCESS}
    rules = {"Hidden": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}

    import app.coverage.worksheet_policy as worksheet_policy_module

    real_apply_remove_all = worksheet_policy_module._apply_remove_all
    call_count = {"value": 0}

    def _spy_apply_remove_all(*args, **kwargs):  # noqa: ANN002, ANN003
        call_count["value"] += 1
        return real_apply_remove_all(*args, **kwargs)

    monkeypatch.setattr(worksheet_policy_module, "_apply_remove_all", _spy_apply_remove_all)

    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, policies, rules):
            pass

    assert exc_info.value.reason is CoverageReason.NO_VISIBLE_PROCESS_WORKSHEET
    assert call_count["value"] == 0


# ---------------------------------------------------------------------------
# F. os.close(mkstemp fd) failure (закрытие MINOR-1 независимого review)
# ---------------------------------------------------------------------------


def test_mkstemp_fd_close_failure_is_sanitized(
    source_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _build_source(source_path)

    captured_fd: dict[str, int] = {}
    real_mkstemp = _tempfile_module.mkstemp

    def _tracking_mkstemp(*args, **kwargs):  # noqa: ANN002, ANN003
        fd, name = real_mkstemp(*args, **kwargs)
        captured_fd["value"] = fd
        return fd, name

    monkeypatch.setattr("app.coverage.worksheet_policy.tempfile.mkstemp", _tracking_mkstemp)

    real_close = os.close

    def _boom_close_for_mkstemp_fd_only(fd: int) -> None:
        # Падает ТОЛЬКО на конкретном fd, выданном mkstemp внутри
        # тестируемого вызова — любой другой os.close (в т.ч. используемый
        # pytest/fixtures/самим тестом при cleanup) делегируется реальной
        # функции без изменений.
        if captured_fd.get("value") == fd:
            raise OSError(f"simulated close failure containing password {PASSWORD_SENTINEL}")
        return real_close(fd)

    monkeypatch.setattr("app.coverage.worksheet_policy.os.close", _boom_close_for_mkstemp_fd_only)

    save_called = {"value": False}

    def _tracking_save(self, filename):  # noqa: ANN001
        save_called["value"] = True

    monkeypatch.setattr(openpyxl.Workbook, "save", _tracking_save)

    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass

    exc = exc_info.value
    assert exc.reason is CoverageReason.WORKBOOK_MUTATION_FAILED

    # OD-7: sentinel из "утечного" OSError не должен просочиться никуда.
    import traceback

    text_blobs = [
        str(exc),
        repr(exc),
        repr(exc.args),
        "".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
    ]
    for blob in text_blobs:
        assert PASSWORD_SENTINEL not in blob
    assert exc.__context__ is None
    assert exc.__cause__ is None

    # workbook.save не должен был вызываться после сбоя закрытия mkstemp fd.
    assert save_called["value"] is False

    # temp path (пустой на момент сбоя) удалён best-effort.
    assert "value" in captured_fd  # mkstemp действительно был вызван


# ---------------------------------------------------------------------------
# G. Дополнительные regression-тесты для formula-детекции (MINOR-3
#    независимого review): формула в header row / hidden row / hidden column.
# ---------------------------------------------------------------------------


def test_formula_in_header_row_on_process_worksheet_rejected(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "=1+1"
    ws["A2"] = "value"
    wb.save(source_path)
    wb.close()

    rules = {"Data": {1: FieldRule(column_name="=1+1", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is CoverageReason.FORMULA_REJECTED


def test_formula_in_hidden_row_on_process_worksheet_rejected(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Company"
    ws["A2"] = "value"
    ws["A3"] = "=A2"
    ws.row_dimensions[3].hidden = True
    wb.save(source_path)
    wb.close()

    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass
    assert exc_info.value.reason is CoverageReason.FORMULA_REJECTED


def test_formula_in_hidden_column_without_rule_on_process_worksheet_rejected(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Company"
    ws["A2"] = "value"
    ws["B1"] = "Hidden"
    ws["B2"] = "=A2"
    ws.column_dimensions["B"].hidden = True
    wb.save(source_path)
    wb.close()

    # B не входит в rules_by_sheet вообще — формула всё равно обязана быть
    # обнаружена (formula-детекция безусловна, независимо от column rule).
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, _PROCESS_RULES):
            pass
    assert exc_info.value.reason is CoverageReason.FORMULA_REJECTED


# ---------------------------------------------------------------------------
# H. Merged non-anchor header (MINOR-3 независимого review)
# ---------------------------------------------------------------------------


def test_merged_non_anchor_header_column_rejected_fail_closed(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Merged"
    ws.merge_cells("A1:B1")
    ws["A2"] = "v1"
    ws["B2"] = "v2"
    wb.save(source_path)
    wb.close()

    # Колонка 2 — non-anchor часть merge (значение всегда None у openpyxl),
    # правило для неё объявляет тот же column_name, что и у якоря — это
    # НЕ должно "унаследовать" текст якоря: fail-closed, HEADER_MISMATCH.
    rules = {
        "Data": {
            1: FieldRule(column_name="Merged", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE),
            2: FieldRule(column_name="Merged", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE),
        }
    }
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is CoverageReason.HEADER_MISMATCH
