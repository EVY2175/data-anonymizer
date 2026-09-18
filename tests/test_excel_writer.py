"""
Тесты Excel Writer (Stage 8.2) — app.excel.writer.write_anonymized_workbook.

Все тестовые workbook содержат только синтетические данные (никаких
реальных ИНН/КПП/ОГРН/ФИО/названий компаний) — Writer работает с уже
готовыми FlatTable, а не с исходными идентификаторами, поэтому даже
"похожие на ИНН" значения здесь не нужны.

FlatTable-фикстуры строятся напрямую через CellRecord/FlatTable, без
app.anonymizer — Writer не знает и не должен знать про anonymizer/
mapping stores (см. test_no_store_dependency_in_writer_module).
"""

from __future__ import annotations

import ast
import os
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.packaging.custom import IntProperty, StringProperty
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from app.excel.models import CellRecord, FlatTable
from app.excel.writer import WriterValidationError, write_anonymized_workbook

JOB_ID = "job-stage-8-2-test-0011223344"
STALE_JOB_ID = "STALE_PREVIOUS_JOB_ID"


# ---------------------------------------------------------------------------
# Фикстуры/помощники
# ---------------------------------------------------------------------------


def _table(sheet_name: str, header: list[object], rows: list[list[object]]) -> FlatTable:
    header_row = tuple(CellRecord(row=1, column=i + 1, value=v) for i, v in enumerate(header))
    data_rows = tuple(
        tuple(CellRecord(row=r_idx + 2, column=c_idx + 1, value=v) for c_idx, v in enumerate(row))
        for r_idx, row in enumerate(rows)
    )
    return FlatTable(sheet_name=sheet_name, header_row=header_row, rows=data_rows)


def _simple_source(path: Path, sheet_name: str = "Data") -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["A1"] = "Name"
    ws["B1"] = "Code"
    ws["A2"] = "original-name-1"
    ws["B2"] = "code-1"
    ws["A3"] = "original-name-2"
    ws["B3"] = "code-2"
    wb.save(path)


@pytest.fixture()
def source_path(tmp_path: Path) -> Path:
    p = tmp_path / "source.xlsx"
    _simple_source(p)
    return p


@pytest.fixture()
def destination_path(tmp_path: Path) -> Path:
    return tmp_path / "output.xlsx"


# ---------------------------------------------------------------------------
# A. Запись изменённых значений / untouched cells
# ---------------------------------------------------------------------------


def test_writes_changed_values(source_path: Path, destination_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"], ["ANON-2", "code-2"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    wb = openpyxl.load_workbook(destination_path)
    ws = wb["Data"]
    assert ws["A2"].value == "ANON-1"
    assert ws["A3"].value == "ANON-2"


def test_untouched_cells_unchanged(source_path: Path, destination_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"], ["original-name-2", "code-2"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    wb = openpyxl.load_workbook(destination_path)
    ws = wb["Data"]
    assert ws["B2"].value == "code-1"
    assert ws["A3"].value == "original-name-2"
    assert ws["B3"].value == "code-2"


# ---------------------------------------------------------------------------
# B. Source unchanged (проверка на диске, не только in-memory)
# ---------------------------------------------------------------------------


def test_source_unchanged_on_disk(source_path: Path, destination_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"], ["ANON-2", "code-2"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    reopened_source = openpyxl.load_workbook(source_path)
    ws = reopened_source["Data"]
    assert ws["A2"].value == "original-name-1"
    assert ws["A3"].value == "original-name-2"
    assert reopened_source.custom_doc_props.names == []


# ---------------------------------------------------------------------------
# C. Formulas preserved
# ---------------------------------------------------------------------------


def test_untouched_formula_preserved(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "formula_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Len"
    ws["A2"] = "original-name-1"
    ws["B2"] = "=LEN(A2)"
    wb.save(source)

    table = _table("Data", ["Name", "Len"], [["ANON-1", "=LEN(A2)"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    wb2 = openpyxl.load_workbook(destination_path, data_only=False)
    ws2 = wb2["Data"]
    assert ws2["A2"].value == "ANON-1"
    assert ws2["B2"].value == "=LEN(A2)"
    assert ws2["B2"].data_type == "f"


# ---------------------------------------------------------------------------
# D. Formatting preserved on modified cell
# ---------------------------------------------------------------------------


def test_formatting_preserved_on_modified_cell(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "formatting_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    ws["A2"].font = Font(bold=True, color="FF0000")
    ws["A2"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws["A2"].border = Border(bottom=Side(style="thin"))
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].number_format = "@"
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    wb2 = openpyxl.load_workbook(destination_path)
    cell = wb2["Data"]["A2"]
    assert cell.value == "ANON-1"
    assert cell.font.bold is True
    assert cell.font.color.rgb == "00FF0000"
    assert cell.fill.fgColor.rgb == "00FFFF00"
    assert cell.border.bottom.style == "thin"
    assert cell.alignment.horizontal == "center"
    assert cell.number_format == "@"


# ---------------------------------------------------------------------------
# E-F. Multi-sheet / hidden sheet preservation
# ---------------------------------------------------------------------------


def test_multiple_sheets_preserved(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "multi_sheet_source.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data"
    ws1["A1"] = "Name"
    ws1["A2"] = "original-name-1"
    ws2 = wb.create_sheet("Other")
    ws2["A1"] = "untouched sheet content"
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    assert result.sheetnames == ["Data", "Other"]
    assert result["Other"]["A1"].value == "untouched sheet content"


def test_hidden_sheet_state_preserved(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "hidden_sheet_source.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data"
    ws1["A1"] = "Name"
    ws1["A2"] = "original-name-1"
    ws2 = wb.create_sheet("Hidden")
    ws2["A1"] = "hidden content"
    ws2.sheet_state = "hidden"
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    assert result["Hidden"].sheet_state == "hidden"
    assert result["Hidden"]["A1"].value == "hidden content"


# ---------------------------------------------------------------------------
# G/H1-H3. Merged cells
# ---------------------------------------------------------------------------


def _merged_source(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Header"
    ws["A2"] = "merged value"
    ws.merge_cells("A2:B2")
    wb.save(path)


def test_h1_non_top_left_merged_cell_with_none_is_safe_noop(
    tmp_path: Path, destination_path: Path
) -> None:
    source = tmp_path / "merged_h1.xlsx"
    _merged_source(source)

    # B2 — non-top-left часть merge A2:B2; None соответствует тому, что
    # реально читает reader для этой позиции (см. app.excel.reader).
    header = (CellRecord(row=1, column=1, value="Header"),)
    rows = ((CellRecord(row=2, column=1, value="merged value"), CellRecord(row=2, column=2, value=None)),)
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    ws = result["Data"]
    assert ws["A2"].value == "merged value"
    assert ("A2:B2") in {str(r) for r in ws.merged_cells.ranges}


def test_h2_non_top_left_merged_cell_with_value_hard_fails(
    tmp_path: Path, destination_path: Path
) -> None:
    source = tmp_path / "merged_h2.xlsx"
    _merged_source(source)

    header = (CellRecord(row=1, column=1, value="Header"),)
    rows = (
        (
            CellRecord(row=2, column=1, value="merged value"),
            CellRecord(row=2, column=2, value="unexpected-non-none-value"),
        ),
    )
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    assert not destination_path.exists()


def test_h2_merged_cell_failure_does_not_apply_earlier_valid_record(
    tmp_path: Path, destination_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """
    Усиленная версия H2: FlatTable содержит СНАЧАЛА обычную валидную
    запись (row=2), которую Writer мог бы изменить, и ТОЛЬКО ПОТОМ
    non-top-left merged-cell конфликт (row=3, column=2). Доказывает, что
    ошибка обнаруживается в рамках единой validation-фазы ДО начала
    mutation phase — т.е. валидная запись, обработанная validation-циклом
    раньше конфликтной, всё равно не применяется ни к in-memory workbook,
    ни (тем более) к какому-либо сохранённому файлу.
    """
    source = tmp_path / "merged_h2_transactional.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Header"
    ws["A2"] = "original-plain-value"
    ws["A3"] = "merged value"
    ws.merge_cells("A3:B3")
    wb.save(source)

    from openpyxl.workbook.workbook import Workbook

    save_called = {"n": 0}
    real_save = Workbook.save

    def counting_save(self: Workbook, *args: object, **kwargs: object) -> None:
        save_called["n"] += 1
        return real_save(self, *args, **kwargs)

    monkeypatch.setattr(Workbook, "save", counting_save)

    header = (CellRecord(row=1, column=1, value="Header"),)
    rows = (
        (CellRecord(row=2, column=1, value="ANON-should-not-be-applied"),),
        (
            CellRecord(row=3, column=1, value="merged value"),
            CellRecord(row=3, column=2, value="unexpected-non-none-value"),
        ),
    )
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    assert save_called["n"] == 0
    assert not destination_path.exists()

    reopened_source = openpyxl.load_workbook(source)
    ws2 = reopened_source["Data"]
    assert ws2["A2"].value == "original-plain-value"
    assert ws2["A3"].value == "merged value"
    assert "A3:B3" in {str(r) for r in ws2.merged_cells.ranges}


def test_h3_top_left_merged_cell_write_succeeds(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "merged_h3.xlsx"
    _merged_source(source)

    header = (CellRecord(row=1, column=1, value="Header"),)
    rows = ((CellRecord(row=2, column=1, value="ANON-merged"), CellRecord(row=2, column=2, value=None)),)
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    ws = result["Data"]
    assert ws["A2"].value == "ANON-merged"
    assert ("A2:B2") in {str(r) for r in ws.merged_cells.ranges}


# ---------------------------------------------------------------------------
# I1. Bounds test
# ---------------------------------------------------------------------------


def test_i1_coordinate_outside_source_bounds_hard_fails(
    source_path: Path, destination_path: Path
) -> None:
    # source_path имеет max_row=3, max_column=2 (см. _simple_source).
    header = (CellRecord(row=1, column=1, value="Name"), CellRecord(row=1, column=2, value="Code"))
    rows = (
        (CellRecord(row=2, column=1, value="ANON-1"), CellRecord(row=2, column=2, value="code-1")),
        (CellRecord(row=50, column=1, value="out-of-bounds"), CellRecord(row=50, column=2, value="x")),
    )
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    assert not destination_path.exists()

    # source не должен был структурно расшириться.
    reopened = openpyxl.load_workbook(source_path)
    assert reopened["Data"].max_row == 3
    assert reopened["Data"].max_column == 2


# ---------------------------------------------------------------------------
# J/45. Транзакционная предвалидация нескольких FlatTable
# ---------------------------------------------------------------------------


def test_transactional_prevalidation_no_partial_output(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "two_sheets_source.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "Name"
    ws1["A2"] = "original-first"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "Name"
    ws2["A2"] = "original-second"
    wb.save(source)

    valid_table = _table("First", ["Name"], [["ANON-first"]])
    # Второй table ссылается на несуществующий лист -> ошибка валидации
    # должна быть обнаружена ДО мутации первого (валидного) table.
    invalid_table = _table("NoSuchSheet", ["Name"], [["ANON-second"]])

    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(
            source, [valid_table, invalid_table], destination_path, job_id=JOB_ID
        )

    assert not destination_path.exists()

    reopened_source = openpyxl.load_workbook(source)
    assert reopened_source["First"]["A2"].value == "original-first"
    assert reopened_source["Second"]["A2"].value == "original-second"


# ---------------------------------------------------------------------------
# K1-K2. Duplicate coordinates / duplicate sheet_name
# ---------------------------------------------------------------------------


def test_k1_duplicate_coordinates_within_one_table_hard_fails(
    source_path: Path, destination_path: Path
) -> None:
    header = (CellRecord(row=1, column=1, value="Name"), CellRecord(row=1, column=2, value="Code"))
    rows = (
        (CellRecord(row=2, column=1, value="ANON-1"), CellRecord(row=2, column=2, value="code-1")),
        (CellRecord(row=2, column=1, value="DUPLICATE"), CellRecord(row=2, column=2, value="code-x")),
    )
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    assert not destination_path.exists()


def test_k2_duplicate_flat_table_sheet_name_hard_fails(
    source_path: Path, destination_path: Path
) -> None:
    table_a = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    table_b = _table("Data", ["Name", "Code"], [["ANON-2", "code-2"]])

    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table_a, table_b], destination_path, job_id=JOB_ID)

    assert not destination_path.exists()


# ---------------------------------------------------------------------------
# L/47. Missing sheet
# ---------------------------------------------------------------------------


def test_missing_sheet_hard_fails(source_path: Path, destination_path: Path) -> None:
    table = _table("DoesNotExist", ["Name"], [["ANON-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)
    assert not destination_path.exists()


# ---------------------------------------------------------------------------
# M/49. Extension tests
# ---------------------------------------------------------------------------


def test_xlsx_source_and_destination_accepted(source_path: Path, destination_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)
    assert destination_path.exists()


def test_xlsm_source_rejected(tmp_path: Path, destination_path: Path) -> None:
    fake_xlsm = tmp_path / "source.xlsm"
    fake_xlsm.write_bytes(b"not a real xlsm, extension check happens first")
    table = _table("Data", ["Name"], [["ANON-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(fake_xlsm, [table], destination_path, job_id=JOB_ID)


def test_xls_source_rejected(tmp_path: Path, destination_path: Path) -> None:
    fake_xls = tmp_path / "source.xls"
    fake_xls.write_bytes(b"not a real xls, extension check happens first")
    table = _table("Data", ["Name"], [["ANON-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(fake_xls, [table], destination_path, job_id=JOB_ID)


def test_xlsm_destination_rejected(source_path: Path, tmp_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], tmp_path / "output.xlsm", job_id=JOB_ID)


# ---------------------------------------------------------------------------
# N/50. Destination parent missing
# ---------------------------------------------------------------------------


def test_missing_destination_parent_raises_and_does_not_create_directory(
    source_path: Path, tmp_path: Path
) -> None:
    missing_parent_dest = tmp_path / "no_such_dir" / "output.xlsx"
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(FileNotFoundError):
        write_anonymized_workbook(source_path, [table], missing_parent_dest, job_id=JOB_ID)
    assert not (tmp_path / "no_such_dir").exists()


# ---------------------------------------------------------------------------
# O/51. Existing destination
# ---------------------------------------------------------------------------


def test_existing_destination_is_atomically_replaced(
    source_path: Path, destination_path: Path
) -> None:
    # Старый destination содержит лист "Sheet" (имя активного листа по
    # умолчанию у openpyxl.Workbook()) — source_path же содержит только
    # лист "Data" (см. _simple_source). Полная атомарная замена
    # destination содержимым, построенным из source, должна означать,
    # что после Writer лист "Sheet" отсутствует вовсе, а "Data" содержит
    # анонимизированное значение — прямая, однозначная проверка вместо
    # избыточной/вводящей в заблуждение проверки чужого диапазона.
    wb = openpyxl.Workbook()
    wb.active["A1"] = "PRE-EXISTING DESTINATION CONTENT"
    wb.save(destination_path)

    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    assert "Data" in result.sheetnames
    assert "Sheet" not in result.sheetnames
    assert result["Data"]["A2"].value == "ANON-1"


def test_failure_before_replace_leaves_existing_destination_unchanged(
    source_path: Path, destination_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "PRE-EXISTING DESTINATION CONTENT"
    wb.save(destination_path)

    def failing_replace(*args: object, **kwargs: object) -> None:
        raise OSError("simulated os.replace failure")

    monkeypatch.setattr(os, "replace", failing_replace)

    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(OSError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    monkeypatch.undo()
    result = openpyxl.load_workbook(destination_path)
    assert result.active["A1"].value == "PRE-EXISTING DESTINATION CONTENT"


# ---------------------------------------------------------------------------
# P/48. source == destination
# ---------------------------------------------------------------------------


def test_source_equal_destination_rejected(source_path: Path) -> None:
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(
            source_path, [_table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])], source_path, job_id=JOB_ID
        )

    reopened = openpyxl.load_workbook(source_path)
    assert reopened["Data"]["A2"].value == "original-name-1"


def test_source_equal_destination_via_different_relative_spelling_rejected(
    source_path: Path,
) -> None:
    equivalent_path = source_path.parent / "." / source_path.name
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(
            source_path,
            [_table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])],
            equivalent_path,
            job_id=JOB_ID,
        )


# ---------------------------------------------------------------------------
# job_id validation
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("bad_job_id", [None, 123, True, 3.14, [], {}])
def test_non_str_job_id_rejected(source_path: Path, destination_path: Path, bad_job_id: object) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=bad_job_id)  # type: ignore[arg-type]


@pytest.mark.parametrize("bad_job_id", ["", "   "])
def test_empty_or_whitespace_job_id_rejected(
    source_path: Path, destination_path: Path, bad_job_id: str
) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=bad_job_id)


def test_job_id_not_normalized_before_storage(source_path: Path, destination_path: Path) -> None:
    padded_job_id = "  job-with-outer-whitespace-preserved-exactly  "
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=padded_job_id)

    result = openpyxl.load_workbook(destination_path)
    assert result.custom_doc_props["DataAnonymizer.JobId"].value == padded_job_id


# ---------------------------------------------------------------------------
# job_id round-trip (metadata)
# ---------------------------------------------------------------------------


def test_job_id_round_trip_single_cycle(source_path: Path, destination_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    assert result.custom_doc_props.names.count("DataAnonymizer.JobId") == 1
    prop = result.custom_doc_props["DataAnonymizer.JobId"]
    assert isinstance(prop, StringProperty)
    assert prop.value == JOB_ID


def test_job_id_survives_second_save_reopen_cycle(
    source_path: Path, destination_path: Path, tmp_path: Path
) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    reopened = openpyxl.load_workbook(destination_path)
    second_cycle_path = tmp_path / "second_cycle.xlsx"
    reopened.save(second_cycle_path)

    final = openpyxl.load_workbook(second_cycle_path)
    assert final.custom_doc_props.names.count("DataAnonymizer.JobId") == 1
    assert final.custom_doc_props["DataAnonymizer.JobId"].value == JOB_ID


# ---------------------------------------------------------------------------
# Stale job_id replacement + unrelated metadata preservation
# ---------------------------------------------------------------------------


def test_stale_job_id_replaced_and_unrelated_property_preserved(
    tmp_path: Path, destination_path: Path
) -> None:
    source = tmp_path / "stale_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    wb.custom_doc_props.append(StringProperty(name="DataAnonymizer.JobId", value=STALE_JOB_ID))
    wb.custom_doc_props.append(StringProperty(name="UnrelatedProp", value="keep-me"))
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    names = result.custom_doc_props.names
    assert names.count("DataAnonymizer.JobId") == 1
    assert result.custom_doc_props["DataAnonymizer.JobId"].value == JOB_ID
    assert result.custom_doc_props["DataAnonymizer.JobId"].value != STALE_JOB_ID
    assert result.custom_doc_props["UnrelatedProp"].value == "keep-me"


def test_malformed_type_prior_job_id_property_replaced(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "malformed_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    wb.custom_doc_props.append(IntProperty(name="DataAnonymizer.JobId", value=999))
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    prop = result.custom_doc_props["DataAnonymizer.JobId"]
    assert isinstance(prop, StringProperty)
    assert prop.value == JOB_ID


def _write_source_with_duplicate_job_id_properties(
    path: Path, stale_values: list[str], unrelated_value: str = "keep-me"
) -> None:
    """
    Формирует synthetic .xlsx, docProps/custom.xml которой физически
    содержит НЕСКОЛЬКО custom properties с ОДНИМ И ТЕМ ЖЕ именем
    DataAnonymizer.JobId — сценарий, недостижимый через
    CustomPropertyList.append() (который сам запрещает дублирование
    имён), но возможный для книги, отредактированной не через openpyxl.
    Только для тестовой фикстуры — production-код такой манипуляции не
    содержит и не должен содержать.
    """
    plain = path.parent / f".{path.stem}.plain.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    wb.save(plain)

    properties_xml = "".join(
        f'<property name="DataAnonymizer.JobId" '
        f'fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i + 2}">'
        f"<vt:lpwstr>{value}</vt:lpwstr></property>"
        for i, value in enumerate(stale_values)
    ) + (
        f'<property name="UnrelatedProp" '
        f'fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{len(stale_values) + 2}">'
        f"<vt:lpwstr>{unrelated_value}</vt:lpwstr></property>"
    )
    custom_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes" '
        'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties">'
        f"{properties_xml}</Properties>"
    )

    with zipfile.ZipFile(plain, "r") as zin:
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
        rels = zin.read("_rels/.rels").decode("utf-8")

    content_types = content_types.replace(
        "</Types>",
        '<Override PartName="/docProps/custom.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
    )
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdCustomTest" '
        "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties\" "
        'Target="docProps/custom.xml"/></Relationships>',
    )

    with zipfile.ZipFile(plain, "r") as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            if item.filename in ("[Content_Types].xml", "_rels/.rels"):
                continue
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("[Content_Types].xml", content_types)
        zout.writestr("_rels/.rels", rels)
        zout.writestr("docProps/custom.xml", custom_xml)

    plain.unlink()


def test_duplicate_stale_job_id_properties_normalized_to_exactly_one(
    tmp_path: Path, destination_path: Path
) -> None:
    source = tmp_path / "duplicate_job_id_source.xlsx"
    _write_source_with_duplicate_job_id_properties(source, ["stale-job-1", "stale-job-2"])

    # PRECONDITION: source действительно содержит >= 2 property с exact
    # именем DataAnonymizer.JobId — без этого тест ничего бы не доказывал.
    preloaded = openpyxl.load_workbook(source)
    assert preloaded.custom_doc_props.names.count("DataAnonymizer.JobId") == 2

    fresh_job_id = "fresh-job-normalized"
    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=fresh_job_id)

    result = openpyxl.load_workbook(destination_path)
    matches = [name for name in result.custom_doc_props.names if name == "DataAnonymizer.JobId"]
    assert len(matches) == 1

    prop = result.custom_doc_props["DataAnonymizer.JobId"]
    assert isinstance(prop, StringProperty)
    assert prop.value == fresh_job_id

    assert result.custom_doc_props["UnrelatedProp"].value == "keep-me"

    # Writer нормализует только OUTPUT — исходные stale duplicates на
    # диске должны остаться нетронутыми.
    reopened_source = openpyxl.load_workbook(source)
    assert reopened_source.custom_doc_props.names.count("DataAnonymizer.JobId") == 2
    assert {
        reopened_source.custom_doc_props.names[i]
        for i in range(len(reopened_source.custom_doc_props.names))
    } == {"DataAnonymizer.JobId", "UnrelatedProp"}


def test_triple_duplicate_stale_job_id_properties_normalized_to_exactly_one(
    tmp_path: Path, destination_path: Path
) -> None:
    source = tmp_path / "triple_duplicate_job_id_source.xlsx"
    _write_source_with_duplicate_job_id_properties(source, ["stale-1", "stale-2", "stale-3"])

    preloaded = openpyxl.load_workbook(source)
    assert preloaded.custom_doc_props.names.count("DataAnonymizer.JobId") == 3

    fresh_job_id = "fresh-job-triple"
    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=fresh_job_id)

    result = openpyxl.load_workbook(destination_path)
    assert result.custom_doc_props.names.count("DataAnonymizer.JobId") == 1
    assert result.custom_doc_props["DataAnonymizer.JobId"].value == fresh_job_id
    assert result.custom_doc_props["UnrelatedProp"].value == "keep-me"


# ---------------------------------------------------------------------------
# no-store-dependency / no-confidential-metadata
# ---------------------------------------------------------------------------


def test_no_store_dependency_in_writer_module() -> None:
    """
    Проверяет реальные import-операторы модуля (через ast), а не просто
    текстовое вхождение — docstring модуля намеренно ОБЪЯСНЯЕТ, что
    Writer не зависит от ProvenanceStore/IdentifierMappingStore/
    MappingStore, поэтому наивный текстовый grep дал бы ложное
    срабатывание на само это объяснение.
    """
    import app.excel.writer as writer_module

    tree = ast.parse(Path(writer_module.__file__).read_text(encoding="utf-8"))
    imported_modules: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported_modules.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module is not None:
            imported_modules.add(node.module)

    assert not any(name.startswith("app.mapping") for name in imported_modules)
    assert not any(name.startswith("app.security") for name in imported_modules)
    assert "app.anonymizer" not in imported_modules


def test_only_job_id_present_in_output_metadata(source_path: Path, destination_path: Path) -> None:
    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    # Единственная custom property в output — DataAnonymizer.JobId, т.к.
    # source не содержал никаких других custom properties.
    assert result.custom_doc_props.names == ["DataAnonymizer.JobId"]


def test_error_messages_do_not_expose_cell_value(destination_path: Path) -> None:
    secret_looking_value = "SuperSecretRawIdentifierValue999"
    source = destination_path.parent / "merged_secret_source.xlsx"
    _merged_source(source)

    header = (CellRecord(row=1, column=1, value="Header"),)
    rows = (
        (
            CellRecord(row=2, column=1, value="merged value"),
            CellRecord(row=2, column=2, value=secret_looking_value),
        ),
    )
    table = FlatTable(sheet_name="Data", header_row=header, rows=rows)

    with pytest.raises(WriterValidationError) as exc_info:
        write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    assert secret_looking_value not in str(exc_info.value)


# ---------------------------------------------------------------------------
# Empty tables / invalid table type
# ---------------------------------------------------------------------------


def test_empty_tables_rejected(source_path: Path, destination_path: Path) -> None:
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [], destination_path, job_id=JOB_ID)
    assert not destination_path.exists()


def test_non_flat_table_element_rejected(source_path: Path, destination_path: Path) -> None:
    with pytest.raises(TypeError):
        write_anonymized_workbook(source_path, ["not-a-flat-table"], destination_path, job_id=JOB_ID)  # type: ignore[list-item]


# ---------------------------------------------------------------------------
# Validation happens strictly before mutation / save / metadata / replace
# ---------------------------------------------------------------------------


def test_validation_failure_happens_before_any_save_call(
    source_path: Path, destination_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl.workbook.workbook import Workbook

    save_called = {"n": 0}
    real_save = Workbook.save

    def counting_save(self: Workbook, *args: object, **kwargs: object) -> None:
        save_called["n"] += 1
        return real_save(self, *args, **kwargs)

    monkeypatch.setattr(Workbook, "save", counting_save)

    table = _table("DoesNotExist", ["Name"], [["ANON-1"]])
    with pytest.raises(WriterValidationError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    assert save_called["n"] == 0


# ---------------------------------------------------------------------------
# Save / os.replace / fsync failure injection
# ---------------------------------------------------------------------------


def test_save_failure_propagates_and_leaves_no_destination(
    source_path: Path, destination_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl.workbook.workbook import Workbook

    def failing_save(self: Workbook, *args: object, **kwargs: object) -> None:
        raise OSError("simulated save failure")

    monkeypatch.setattr(Workbook, "save", failing_save)

    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(OSError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    monkeypatch.undo()
    assert not destination_path.exists()

    leftover_temp_files = [
        p for p in destination_path.parent.iterdir() if p.name.startswith(f".{destination_path.name}.")
    ]
    assert leftover_temp_files == []

    reopened_source = openpyxl.load_workbook(source_path)
    assert reopened_source["Data"]["A2"].value == "original-name-1"


def test_os_replace_failure_propagates_and_cleans_up_temp(
    source_path: Path, destination_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def failing_replace(*args: object, **kwargs: object) -> None:
        raise OSError("simulated os.replace failure")

    monkeypatch.setattr(os, "replace", failing_replace)

    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(OSError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    monkeypatch.undo()
    assert not destination_path.exists()

    leftover_temp_files = [
        p for p in destination_path.parent.iterdir() if p.name.startswith(f".{destination_path.name}.")
    ]
    assert leftover_temp_files == []


def test_fsync_failure_propagates_before_replace(
    source_path: Path, destination_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    replace_called = {"n": 0}
    real_replace = os.replace

    def counting_replace(*args: object, **kwargs: object) -> None:
        replace_called["n"] += 1
        return real_replace(*args, **kwargs)

    def failing_fsync(fd: int) -> None:
        raise OSError("simulated fsync failure")

    monkeypatch.setattr(os, "fsync", failing_fsync)
    monkeypatch.setattr(os, "replace", counting_replace)

    table = _table("Data", ["Name", "Code"], [["ANON-1", "code-1"]])
    with pytest.raises(OSError):
        write_anonymized_workbook(source_path, [table], destination_path, job_id=JOB_ID)

    monkeypatch.undo()
    assert replace_called["n"] == 0
    assert not destination_path.exists()

    leftover_temp_files = [
        p for p in destination_path.parent.iterdir() if p.name.startswith(f".{destination_path.name}.")
    ]
    assert leftover_temp_files == []


# ---------------------------------------------------------------------------
# Другие preservation-проверки (row height / column width / freeze panes /
# data validation)
# ---------------------------------------------------------------------------


def test_row_height_column_width_freeze_panes_preserved(
    tmp_path: Path, destination_path: Path
) -> None:
    source = tmp_path / "dimensions_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    ws.row_dimensions[1].height = 30
    ws.column_dimensions["A"].width = 25
    ws.freeze_panes = "A2"
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    ws2 = result["Data"]
    assert ws2.row_dimensions[1].height == 30
    assert ws2.column_dimensions["A"].width == 25
    assert ws2.freeze_panes == "A2"


def test_data_validation_preserved(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "data_validation_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    dv = DataValidation(type="list", formula1='"A,B,C"')
    ws.add_data_validation(dv)
    dv.add("A2")
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    assert len(result["Data"].data_validations.dataValidation) == 1


def test_core_property_preserved(tmp_path: Path, destination_path: Path) -> None:
    source = tmp_path / "core_property_source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["A2"] = "original-name-1"
    wb.properties.title = "Synthetic Test Workbook"
    wb.save(source)

    table = _table("Data", ["Name"], [["ANON-1"]])
    write_anonymized_workbook(source, [table], destination_path, job_id=JOB_ID)

    result = openpyxl.load_workbook(destination_path)
    assert result.properties.title == "Synthetic Test Workbook"
