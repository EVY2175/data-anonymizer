"""
Тесты Stage 10C.1 — WorksheetPolicy, validate_worksheet_coverage,
validate_process_rules (app.coverage.worksheet_policy).

Все workbook-фикстуры строятся напрямую через openpyxl.Workbook() в
памяти (без файлового I/O) — validate_worksheet_coverage принимает уже
открытый workbook, файловый lifecycle тестируется отдельно в
test_remove_all_application.py.
"""

from __future__ import annotations

import openpyxl
import pytest

from app.coverage.errors import CoverageError, CoverageReason
from app.coverage.models import WorksheetPolicy
from app.coverage.worksheet_policy import validate_process_rules, validate_worksheet_coverage
from app.models.rules import Action, FieldRule, FieldType


def _workbook(sheet_names: list[str], hidden: dict[str, str] | None = None) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        wb.create_sheet(name)
    if hidden:
        for name, state in hidden.items():
            wb[name].sheet_state = state
    return wb


def _rule(column_name: str = "Company", field_type: FieldType = FieldType.COMPANY) -> FieldRule:
    return FieldRule(column_name=column_name, field_type=field_type, action=Action.PSEUDONYMIZE)


# ---------------------------------------------------------------------------
# A. WorksheetPolicy — enum
# ---------------------------------------------------------------------------


def test_worksheet_policy_has_exactly_two_members() -> None:
    assert {member.value for member in WorksheetPolicy} == {"process", "remove_all"}
    assert len(WorksheetPolicy) == 2


# ---------------------------------------------------------------------------
# B. validate_worksheet_coverage — успешные сценарии
# ---------------------------------------------------------------------------


def test_all_process_covered_passes() -> None:
    wb = _workbook(["Data"])
    validate_worksheet_coverage(wb, {"Data": WorksheetPolicy.PROCESS})


def test_process_and_remove_all_mixed_passes() -> None:
    wb = _workbook(["Data", "Junk"])
    validate_worksheet_coverage(
        wb, {"Data": WorksheetPolicy.PROCESS, "Junk": WorksheetPolicy.REMOVE_ALL}
    )


def test_hidden_worksheet_with_policy_passes() -> None:
    wb = _workbook(["Data", "Hidden"], hidden={"Hidden": "hidden"})
    validate_worksheet_coverage(
        wb, {"Data": WorksheetPolicy.PROCESS, "Hidden": WorksheetPolicy.REMOVE_ALL}
    )


def test_very_hidden_worksheet_with_policy_passes() -> None:
    wb = _workbook(["Data", "Secret"], hidden={"Secret": "veryHidden"})
    validate_worksheet_coverage(
        wb, {"Data": WorksheetPolicy.PROCESS, "Secret": WorksheetPolicy.PROCESS}
    )


# ---------------------------------------------------------------------------
# C. validate_worksheet_coverage — HARD FAIL сценарии
# ---------------------------------------------------------------------------


def test_missing_worksheet_policy_raises() -> None:
    wb = _workbook(["Data", "Extra"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {"Data": WorksheetPolicy.PROCESS})
    assert exc_info.value.reason is CoverageReason.MISSING_WORKSHEET_POLICY


def test_hidden_worksheet_missing_policy_raises() -> None:
    wb = _workbook(["Data", "Hidden"], hidden={"Hidden": "hidden"})
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {"Data": WorksheetPolicy.PROCESS})
    assert exc_info.value.reason is CoverageReason.MISSING_WORKSHEET_POLICY


def test_very_hidden_worksheet_missing_policy_raises() -> None:
    wb = _workbook(["Data", "Secret"], hidden={"Secret": "veryHidden"})
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {"Data": WorksheetPolicy.PROCESS})
    assert exc_info.value.reason is CoverageReason.MISSING_WORKSHEET_POLICY


def test_policy_for_unknown_worksheet_raises() -> None:
    wb = _workbook(["Data"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(
            wb, {"Data": WorksheetPolicy.PROCESS, "Ghost": WorksheetPolicy.REMOVE_ALL}
        )
    assert exc_info.value.reason is CoverageReason.POLICY_FOR_UNKNOWN_WORKSHEET


def test_all_remove_all_raises_no_process_remains() -> None:
    wb = _workbook(["Data", "Junk"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(
            wb, {"Data": WorksheetPolicy.REMOVE_ALL, "Junk": WorksheetPolicy.REMOVE_ALL}
        )
    assert exc_info.value.reason is CoverageReason.NO_PROCESS_WORKSHEET_REMAINS


def test_single_remove_all_worksheet_raises_no_process_remains() -> None:
    wb = _workbook(["Only"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {"Only": WorksheetPolicy.REMOVE_ALL})
    assert exc_info.value.reason is CoverageReason.NO_PROCESS_WORKSHEET_REMAINS


# ---------------------------------------------------------------------------
# C2. NO_VISIBLE_PROCESS_WORKSHEET (закрытие MAJOR-1 независимого review):
# PROCESS-листы существуют, но ни один из них не visible.
# ---------------------------------------------------------------------------


def test_remaining_process_hidden_raises_no_visible_process_worksheet() -> None:
    # A: visible sheet = REMOVE_ALL, remaining PROCESS = hidden.
    wb = _workbook(["Visible", "Hidden"], hidden={"Hidden": "hidden"})
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(
            wb, {"Visible": WorksheetPolicy.REMOVE_ALL, "Hidden": WorksheetPolicy.PROCESS}
        )
    assert exc_info.value.reason is CoverageReason.NO_VISIBLE_PROCESS_WORKSHEET


def test_remaining_process_very_hidden_raises_no_visible_process_worksheet() -> None:
    # B: visible sheet = REMOVE_ALL, remaining PROCESS = veryHidden.
    wb = _workbook(["Visible", "Secret"], hidden={"Secret": "veryHidden"})
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(
            wb, {"Visible": WorksheetPolicy.REMOVE_ALL, "Secret": WorksheetPolicy.PROCESS}
        )
    assert exc_info.value.reason is CoverageReason.NO_VISIBLE_PROCESS_WORKSHEET


def test_remaining_process_hidden_and_very_hidden_raises_no_visible_process_worksheet() -> None:
    # C: visible sheet = REMOVE_ALL, remaining PROCESS-листы = hidden + veryHidden
    # (несколько PROCESS-листов, ни один не visible).
    wb = _workbook(
        ["Visible", "Hidden", "Secret"],
        hidden={"Hidden": "hidden", "Secret": "veryHidden"},
    )
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(
            wb,
            {
                "Visible": WorksheetPolicy.REMOVE_ALL,
                "Hidden": WorksheetPolicy.PROCESS,
                "Secret": WorksheetPolicy.PROCESS,
            },
        )
    assert exc_info.value.reason is CoverageReason.NO_VISIBLE_PROCESS_WORKSHEET


def test_remove_all_hidden_remaining_process_visible_allowed() -> None:
    # D: REMOVE_ALL hidden/veryHidden, remaining PROCESS — visible (обычный
    # случай: скрытые служебные листы удаляются, видимый лист обрабатывается).
    wb = _workbook(["Visible", "Hidden"], hidden={"Hidden": "hidden"})
    validate_worksheet_coverage(
        wb, {"Visible": WorksheetPolicy.PROCESS, "Hidden": WorksheetPolicy.REMOVE_ALL}
    )


def test_process_visible_and_process_hidden_allowed() -> None:
    # E: PROCESS visible + PROCESS hidden — допустимо, т.к. среди PROCESS
    # есть хотя бы один visible.
    wb = _workbook(["Visible", "Hidden"], hidden={"Hidden": "hidden"})
    validate_worksheet_coverage(
        wb, {"Visible": WorksheetPolicy.PROCESS, "Hidden": WorksheetPolicy.PROCESS}
    )


# ---------------------------------------------------------------------------
# D. validate_worksheet_coverage — некорректный вход
# ---------------------------------------------------------------------------


def test_worksheet_policies_not_a_mapping_raises() -> None:
    wb = _workbook(["Data"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, ["Data"])  # type: ignore[arg-type]
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


def test_worksheet_policies_wrong_value_type_raises() -> None:
    wb = _workbook(["Data"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {"Data": "process"})  # type: ignore[dict-item]
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


def test_worksheet_policies_empty_key_raises() -> None:
    wb = _workbook(["Data"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {"": WorksheetPolicy.PROCESS})
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


def test_worksheet_policies_non_string_key_raises() -> None:
    wb = _workbook(["Data"])
    with pytest.raises(CoverageError) as exc_info:
        validate_worksheet_coverage(wb, {1: WorksheetPolicy.PROCESS})  # type: ignore[dict-item]
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


# ---------------------------------------------------------------------------
# E. validate_process_rules — успешные сценарии
# ---------------------------------------------------------------------------


def test_process_with_rules_and_remove_all_without_passes() -> None:
    validate_process_rules(
        {"Data": WorksheetPolicy.PROCESS, "Junk": WorksheetPolicy.REMOVE_ALL},
        {"Data": {1: _rule()}},
    )


def test_process_with_empty_ruleset_passes_here() -> None:
    # Пустой ruleset синтаксически допустим на уровне 10C.1 — его
    # достаточность (соответствие реальным колонкам) остаётся
    # ответственностью Stage8 anonymizer._validate_rules.
    validate_process_rules({"Data": WorksheetPolicy.PROCESS}, {"Data": {}})


# ---------------------------------------------------------------------------
# F. validate_process_rules — HARD FAIL сценарии
# ---------------------------------------------------------------------------


def test_process_without_rules_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules({"Data": WorksheetPolicy.PROCESS}, {})
    assert exc_info.value.reason is CoverageReason.MISSING_PROCESS_RULES


def test_rules_for_nonexistent_worksheet_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules(
            {"Data": WorksheetPolicy.PROCESS},
            {"Data": {1: _rule()}, "Ghost": {1: _rule()}},
        )
    assert exc_info.value.reason is CoverageReason.RULES_FOR_UNKNOWN_WORKSHEET


def test_rules_for_remove_all_worksheet_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules(
            {"Data": WorksheetPolicy.PROCESS, "Junk": WorksheetPolicy.REMOVE_ALL},
            {"Data": {1: _rule()}, "Junk": {1: _rule()}},
        )
    assert exc_info.value.reason is CoverageReason.RULES_FOR_REMOVE_ALL_WORKSHEET


def test_rules_by_sheet_not_a_mapping_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules({"Data": WorksheetPolicy.PROCESS}, [1, 2, 3])  # type: ignore[arg-type]
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


def test_rules_by_sheet_column_key_not_positive_int_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules({"Data": WorksheetPolicy.PROCESS}, {"Data": {0: _rule()}})
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


def test_rules_by_sheet_column_key_bool_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules({"Data": WorksheetPolicy.PROCESS}, {"Data": {True: _rule()}})
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE


def test_rules_by_sheet_value_not_field_rule_raises() -> None:
    with pytest.raises(CoverageError) as exc_info:
        validate_process_rules({"Data": WorksheetPolicy.PROCESS}, {"Data": {1: "not-a-rule"}})
    assert exc_info.value.reason is CoverageReason.INVALID_POLICY_VALUE
