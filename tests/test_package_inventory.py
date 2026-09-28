"""
Тесты Stage 10C.2.2 — Package Policy / Inventory
(app.scrub.inventory.inspect_package_policy).

Все фикстуры — synthetic (openpyxl-генерация + прямая ZIP/XML-модификация
через zipfile/строковые замены), без единого реального confidential
значения. Стиль/помощники — по образцу tests/test_package_preflight.py.

Ряд REJECT/INVALID-сценариев (calcChain, sharedStrings, customXml,
threaded comments, произвольные relationship/Content-Type мутации)
openpyxl не умеет генерировать напрямую — они строятся ручной модификацией
реального openpyxl-вывода (`[Content_Types].xml`/`.rels`/новые части),
что является production-path тестом (через `inspect_package_policy`),
а не тестом внутренних helper'ов.
"""

from __future__ import annotations

import io
import re
import traceback
import zipfile
from pathlib import Path
from unittest import mock

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table

from app.scrub.errors import PackageScrubError, PackageScrubReason
from app.scrub.inventory import inspect_package_policy
from app.scrub.models import PackageInventory

PASSWORD_SENTINEL = "SENTINEL_SECRET_PATH_COMPONENT"

_CT_TABLE = "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"
_CT_CALC_CHAIN = "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"
_CT_SHARED_STRINGS = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
_REL_CALC_CHAIN = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain"


# ---------------------------------------------------------------------------
# Фикстуры/помощники
# ---------------------------------------------------------------------------


def _minimal_workbook_parts() -> dict[str, bytes]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "v"
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    with zipfile.ZipFile(buf) as z:
        return {name: z.read(name) for name in z.namelist()}


def _write_xlsx(path: Path, parts: dict[str, bytes], extra: list[tuple[str, bytes]] | None = None) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
        for name, data in extra or []:
            z.writestr(name, data)


@pytest.fixture()
def base_parts() -> dict[str, bytes]:
    return _minimal_workbook_parts()


def _assert_no_sentinel_leak(exc: BaseException, sentinel: str = PASSWORD_SENTINEL) -> None:
    blobs = [
        str(exc),
        repr(exc),
        repr(exc.args),
        "".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
    ]
    for blob in blobs:
        assert sentinel not in blob
    assert exc.__context__ is None
    assert exc.__cause__ is None


# ---------------------------------------------------------------------------
# A. PASS — обычные аналитические workbook (реальные openpyxl-фичи)
# ---------------------------------------------------------------------------


def test_minimal_workbook_pass(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "minimal.xlsx"
    _write_xlsx(p, base_parts)
    result = inspect_package_policy(p)
    assert isinstance(result, PackageInventory)
    assert result == PackageInventory(
        has_shared_strings=False,
        has_calc_chain=False,
        comment_part_count=0,
        table_part_count=0,
        hyperlink_relationship_count=0,
        has_core_properties=True,
        has_app_properties=True,
        has_custom_properties=False,
    )


def test_styled_workbook_with_common_features_pass(tmp_path: Path) -> None:
    # formulas, merged cells, freeze panes, autofilter, conditional
    # formatting, data validation, defined names, hidden worksheet -- ни
    # одна из этих возможностей не создаёт отдельный OPC part, поэтому
    # inventory не должен на них ложно реагировать.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 1
    ws["A2"] = 2
    ws["B1"] = "=A1+A2"
    ws.merge_cells("C1:D2")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:B2"
    ws.conditional_formatting.add("A1:A2", CellIsRule(operator="greaterThan", formula=["1"]))
    dv = DataValidation(type="list", formula1='"X,Y"')
    ws.add_data_validation(dv)
    dv.add("E1")
    wb.defined_names["MyRange"] = openpyxl.workbook.defined_name.DefinedName(
        "MyRange", attr_text="Data!$A$1:$A$2"
    )
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "x"

    p = tmp_path / "styled.xlsx"
    wb.save(p)
    wb.close()

    result = inspect_package_policy(p)
    assert result.table_part_count == 0
    assert result.comment_part_count == 0
    assert result.hyperlink_relationship_count == 0


def test_synthetic_shared_strings_pass(tmp_path: Path, base_parts) -> None:
    shared_strings_xml = (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'count="0" uniqueCount="0"/>'
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        f'<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT_SHARED_STRINGS}"/></Types>',
    )
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    wb_rels2 = wb_rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdSharedStrings" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" '
        'Target="sharedStrings.xml"/></Relationships>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "shared_strings.xlsx"
    _write_xlsx(p, parts, [("xl/sharedStrings.xml", shared_strings_xml)])

    result = inspect_package_policy(p)
    assert result.has_shared_strings is True


def test_synthetic_calc_chain_pass(tmp_path: Path, base_parts) -> None:
    calc_chain_xml = (
        '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<c r="A1" i="1"/></calcChain>'
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>", f'<Override PartName="/xl/calcChain.xml" ContentType="{_CT_CALC_CHAIN}"/></Types>'
    )
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    wb_rels2 = wb_rels.replace(
        "</Relationships>",
        f'<Relationship Id="rIdCalcChain" Type="{_REL_CALC_CHAIN}" Target="calcChain.xml"/></Relationships>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "calc_chain.xlsx"
    _write_xlsx(p, parts, [("xl/calcChain.xml", calc_chain_xml)])

    result = inspect_package_policy(p)
    assert result.has_calc_chain is True


def test_comments_and_vml_pass(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "v"
    ws["A1"].comment = Comment("hi", "author")
    p = tmp_path / "comment_vml.xlsx"
    wb.save(p)
    wb.close()

    result = inspect_package_policy(p)
    assert result.comment_part_count == 1


def test_comment_without_vml_pass(tmp_path: Path, base_parts) -> None:
    # Comments part без сопутствующей VML-части -- допустимо (VML нужен
    # только вместе с comments, обратное не требуется).
    comments_xml = (
        '<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<authors><author>a</author></authors>"
        '<commentList><comment ref="A1" authorId="0"><text><t>hi</t></text></comment></commentList>'
        "</comments>"
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        '<Override PartName="/xl/comments/comment1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"/></Types>',
    )
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdComments" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" '
        'Target="/xl/comments/comment1.xml"/></Relationships>'
    ).encode("utf-8")
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "comment_no_vml.xlsx"
    _write_xlsx(
        p,
        parts,
        [
            ("xl/comments/comment1.xml", comments_xml),
            ("xl/worksheets/_rels/sheet1.xml.rels", ws_rels),
        ],
    )

    result = inspect_package_policy(p)
    assert result.comment_part_count == 1


def test_multiple_tables_on_one_worksheet_pass(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 6):
        ws.cell(row=r, column=1, value=f"v{r}")
        ws.cell(row=r, column=2, value=f"w{r}")
    ws.add_table(Table(displayName="Table1", ref="A1:A5"))
    ws.add_table(Table(displayName="Table2", ref="B1:B5"))
    p = tmp_path / "two_tables.xlsx"
    wb.save(p)
    wb.close()

    result = inspect_package_policy(p)
    # Ключевая корректировка контракта: table_part_count > 1 на ОДНОМ
    # worksheet -- НЕ равно worksheet_count.
    assert result.table_part_count == 2


def test_external_and_internal_hyperlinks_pass(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "x"
    ws["D5"] = "target"
    ws["C1"].value = "link1"
    ws["C1"].hyperlink = "https://example.com/"
    ws["C2"].value = "link2"
    ws["C2"].hyperlink = "https://example.org/"
    # Настоящий внутренний переход -- НЕ создаёт relationship вовсе.
    ws["C3"].value = "internal"
    ws["C3"].hyperlink = Hyperlink(ref="C3", location="Data!D5", display="internal")
    p = tmp_path / "hyperlinks.xlsx"
    wb.save(p)
    wb.close()

    result = inspect_package_policy(p)
    assert result.hyperlink_relationship_count == 2


def test_hyperlink_relationship_count_greater_than_one_per_sheet(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for i, col in enumerate(["C1", "C2", "C3", "C4"]):
        ws[col].value = f"link{i}"
        ws[col].hyperlink = f"https://example.com/{i}"
    p = tmp_path / "many_hyperlinks.xlsx"
    wb.save(p)
    wb.close()

    result = inspect_package_policy(p)
    assert result.hyperlink_relationship_count == 4


def test_core_app_custom_properties_pass(tmp_path: Path, base_parts) -> None:
    root_rels = base_parts["_rels/.rels"].decode("utf-8")
    root_rels2 = root_rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdCustomProps" '
        "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties\" "
        'Target="docProps/custom.xml"/></Relationships>',
    )
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        '<Override PartName="/docProps/custom.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
    )
    custom_xml = (
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
    ).encode("utf-8")
    parts = dict(base_parts)
    parts["_rels/.rels"] = root_rels2.encode("utf-8")
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "custom_props.xlsx"
    _write_xlsx(p, parts, [("docProps/custom.xml", custom_xml)])

    result = inspect_package_policy(p)
    assert result.has_core_properties is True
    assert result.has_app_properties is True
    assert result.has_custom_properties is True


# ---------------------------------------------------------------------------
# B. REJECT — hard-reject features
# ---------------------------------------------------------------------------


def test_chart_rejected(tmp_path: Path) -> None:
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 6):
        ws.cell(row=r, column=1, value=r)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=5))
    ws.add_chart(chart, "H1")
    p = tmp_path / "chart.xlsx"
    wb.save(p)
    wb.close()

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


@pytest.mark.parametrize(
    ("rel_type", "target", "extra"),
    [
        (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing",
            "/xl/drawings/drawing1.xml",
            [("xl/drawings/drawing1.xml", b"<x/>")],
        ),
        (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
            "/xl/embeddings/oleObject1.bin",
            [("xl/embeddings/oleObject1.bin", b"\x00")],
        ),
        (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable",
            "/xl/pivotTables/pivotTable1.xml",
            [("xl/pivotTables/pivotTable1.xml", b"<x/>")],
        ),
    ],
)
def test_worksheet_hard_reject_relationship_types(
    tmp_path: Path, base_parts, rel_type: str, target: str, extra
) -> None:
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rIdX" Type="{rel_type}" Target="{target}"/></Relationships>'
    ).encode("utf-8")
    p = tmp_path / "worksheet_hard_reject.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels)] + extra)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


@pytest.mark.parametrize(
    ("rel_type", "target"),
    [
        ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartsheet", "/xl/chartsheets/sheet2.xml"),
        ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink", "/xl/externalLinks/externalLink1.xml"),
        ("http://schemas.openxmlformats.org/officeDocument/2006/relationships/vbaProject", "vbaProject.bin"),
    ],
)
def test_workbook_hard_reject_relationship_types(tmp_path: Path, base_parts, rel_type: str, target: str) -> None:
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    wb_rels2 = wb_rels.replace(
        "</Relationships>", f'<Relationship Id="rIdX" Type="{rel_type}" Target="{target}"/></Relationships>'
    )
    parts = dict(base_parts)
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "workbook_hard_reject.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_threaded_comments_microsoft_namespace_rejected(tmp_path: Path, base_parts) -> None:
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdTC" '
        'Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
        'Target="/xl/threadedComments/threadedComment1.xml"/></Relationships>'
    ).encode("utf-8")
    p = tmp_path / "threaded.xlsx"
    _write_xlsx(
        p,
        base_parts,
        [
            ("xl/worksheets/_rels/sheet1.xml.rels", ws_rels),
            ("xl/threadedComments/threadedComment1.xml", b"<x/>"),
        ],
    )
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_customxml_root_relationship_rejected(tmp_path: Path, base_parts) -> None:
    root_rels = base_parts["_rels/.rels"].decode("utf-8")
    root_rels2 = root_rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdCustomXml" '
        "Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml\" "
        'Target="customXml/item1.xml"/></Relationships>',
    )
    parts = dict(base_parts)
    parts["_rels/.rels"] = root_rels2.encode("utf-8")
    p = tmp_path / "customxml.xlsx"
    _write_xlsx(p, parts, [("customXml/item1.xml", b"<root/>")])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_unknown_relationship_type_rejected(tmp_path: Path, base_parts) -> None:
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdU" Type="http://example.com/totally-unknown" Target="foo.xml"/>'
        "</Relationships>"
    ).encode("utf-8")
    p = tmp_path / "unknown_rel_type.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_unknown_content_type_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        '<Override PartName="/xl/mystery.xml" ContentType="application/x-totally-unknown"/></Types>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "unknown_ct.xlsx"
    _write_xlsx(p, parts, [("xl/mystery.xml", b"<x/>")])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_macro_enabled_workbook_content_type_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    )
    assert ct2 != ct
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "macro_enabled.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_unknown_default_extension_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>'
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "unknown_default.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_standalone_vml_without_comments_rejected(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "v"
    ws["A1"].comment = Comment("hi", "author")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    with zipfile.ZipFile(buf) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    ws_rels = parts["xl/worksheets/_rels/sheet1.xml.rels"].decode("utf-8")
    ws_rels2 = re.sub(r'<Relationship [^>]*Type="[^"]*relationships/comments"[^>]*/>', "", ws_rels)
    assert ws_rels2 != ws_rels
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = ws_rels2.encode("utf-8")
    p = tmp_path / "standalone_vml.xlsx"
    _write_xlsx(p, parts)

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_generic_application_xml_orphan_rejected(tmp_path: Path, base_parts) -> None:
    # Default xml->application/xml -- допустимая декларация сама по себе,
    # но НЕ разрешает произвольный XML part без структурной роли.
    p = tmp_path / "generic_xml_orphan.xlsx"
    _write_xlsx(p, base_parts, [("xl/evil.xml", b"<x/>")])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_unaccounted_physical_member_rejected(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "unaccounted.xlsx"
    _write_xlsx(p, base_parts, [("xl/mysterious/stray.bin", b"\x00\x01")])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_unexpected_rels_location_rejected(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "unexpected_rels.xlsx"
    _write_xlsx(
        p,
        base_parts,
        [
            (
                "xl/tables/_rels/table1.xml.rels",
                b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
            )
        ],
    )
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_non_hyperlink_external_relationship_rejected(tmp_path: Path, base_parts) -> None:
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdFakeExt" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/table" '
        'Target="http://evil.example/x" TargetMode="External"/></Relationships>'
    ).encode("utf-8")
    p = tmp_path / "non_hyperlink_external.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


# ---------------------------------------------------------------------------
# C. REJECT — orphan approved-role parts
# ---------------------------------------------------------------------------


def test_orphan_table_rejected(tmp_path: Path, base_parts) -> None:
    table_xml = (
        '<table id="1" name="Table1" displayName="Table1" ref="A1:A5" headerRowCount="1" '
        'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<autoFilter ref="A1:A5"/><tableColumns count="1"><tableColumn id="1" name="v1"/></tableColumns>'
        "</table>"
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace("</Types>", f'<Override PartName="/xl/tables/table1.xml" ContentType="{_CT_TABLE}"/></Types>')
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "orphan_table.xlsx"
    _write_xlsx(p, parts, [("xl/tables/table1.xml", table_xml)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_orphan_comments_rejected(tmp_path: Path, base_parts) -> None:
    comments_xml = (
        '<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<authors><author>a</author></authors><commentList/></comments>"
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        '<Override PartName="/xl/comments/comment1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"/></Types>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "orphan_comments.xlsx"
    _write_xlsx(p, parts, [("xl/comments/comment1.xml", comments_xml)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_orphan_vml_rejected(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "v"
    ws["A1"].comment = Comment("hi", "author")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    with zipfile.ZipFile(buf) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    # Убираем ТОЛЬКО vmlDrawing-relationship, оставляя comments -- VML-часть
    # остаётся физически, но становится orphan (ни одна relationship на
    # неё не ссылается).
    ws_rels = parts["xl/worksheets/_rels/sheet1.xml.rels"].decode("utf-8")
    ws_rels2 = re.sub(r'<Relationship [^>]*Type="[^"]*relationships/vmlDrawing"[^>]*/>', "", ws_rels)
    assert ws_rels2 != ws_rels
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = ws_rels2.encode("utf-8")
    p = tmp_path / "orphan_vml.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_orphan_shared_strings_rejected(tmp_path: Path, base_parts) -> None:
    shared_strings_xml = (
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="0" uniqueCount="0"/>'
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        f'<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT_SHARED_STRINGS}"/></Types>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "orphan_shared_strings.xlsx"
    _write_xlsx(p, parts, [("xl/sharedStrings.xml", shared_strings_xml)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


def test_orphan_calc_chain_rejected(tmp_path: Path, base_parts) -> None:
    calc_chain_xml = (
        '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<c r="A1" i="1"/></calcChain>'
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>", f'<Override PartName="/xl/calcChain.xml" ContentType="{_CT_CALC_CHAIN}"/></Types>'
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "orphan_calc_chain.xlsx"
    _write_xlsx(p, parts, [("xl/calcChain.xml", calc_chain_xml)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.UNSUPPORTED_PACKAGE_CONTENT


# ---------------------------------------------------------------------------
# D. INVALID — malformed/structural inconsistency
# ---------------------------------------------------------------------------


def test_duplicate_relationship_id_rejected(tmp_path: Path, base_parts) -> None:
    root_rels = base_parts["_rels/.rels"].decode("utf-8")
    m = re.search(r'<Relationship [^>]*Id="rId1"[^>]*/>', root_rels)
    assert m is not None
    root_rels2 = root_rels.replace("</Relationships>", m.group(0) + "</Relationships>")
    parts = dict(base_parts)
    parts["_rels/.rels"] = root_rels2.encode("utf-8")
    p = tmp_path / "dup_rel_id.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


@pytest.mark.parametrize(
    "mutate",
    [
        lambda xml: re.sub(r' Id="rId1"', "", xml, count=1),
        lambda xml: re.sub(r' Type="[^"]*officeDocument"', "", xml, count=1),
        lambda xml: re.sub(r' Target="xl/workbook\.xml"', "", xml, count=1),
    ],
    ids=["missing_id", "missing_type", "missing_target"],
)
def test_relationship_missing_required_field_rejected(tmp_path: Path, base_parts, mutate) -> None:
    root_rels = base_parts["_rels/.rels"].decode("utf-8")
    mutated = mutate(root_rels)
    assert mutated != root_rels
    parts = dict(base_parts)
    parts["_rels/.rels"] = mutated.encode("utf-8")
    p = tmp_path / "missing_field.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_invalid_target_mode_value_rejected(tmp_path: Path, base_parts) -> None:
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdBadMode" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="x" TargetMode="Weird"/></Relationships>'
    ).encode("utf-8")
    p = tmp_path / "invalid_targetmode.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_override_partname_identical_value_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    m = re.search(r'<Override PartName="/xl/styles\.xml"[^>]*/>', ct)
    assert m is not None
    ct2 = ct.replace("</Types>", m.group(0) + "</Types>")
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "dup_override_same.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_override_partname_conflicting_value_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.theme+xml"/></Types>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "dup_override_conflict.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_default_extension_identical_value_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>", '<Default Extension="xml" ContentType="application/xml"/></Types>'
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "dup_default_same.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_default_extension_conflicting_value_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace("</Types>", '<Default Extension="xml" ContentType="application/x-other"/></Types>')
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "dup_default_conflict.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_role_content_type_mismatch_rejected(tmp_path: Path, base_parts) -> None:
    # table relationship -> целится в workbook.xml (валидная часть с ДРУГИМ
    # content-type) -- role cross-check mismatch.
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdBadTable" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/table" '
        'Target="/xl/workbook.xml"/></Relationships>'
    ).encode("utf-8")
    p = tmp_path / "role_mismatch.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_dangling_internal_target_rejected(tmp_path: Path, base_parts) -> None:
    ws_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdDangling" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/table" '
        'Target="/xl/tables/does_not_exist.xml"/></Relationships>'
    ).encode("utf-8")
    p = tmp_path / "dangling_target.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_ambiguous_physical_target_rejected(tmp_path: Path) -> None:
    # Один и тот же physical target table part, достигнутый ДВУМЯ разными
    # relationship-записями (в т.ч. с разных worksheet) -- структурная
    # неоднозначность.
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "S0"
    for r in range(1, 6):
        ws1.cell(row=r, column=1, value=f"v{r}")
    ws1.add_table(Table(displayName="Table1", ref="A1:A5"))
    ws2 = wb.create_sheet("S1")
    for r in range(1, 6):
        ws2.cell(row=r, column=1, value=f"v{r}")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    with zipfile.ZipFile(buf) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    table_target = "/xl/tables/table1.xml"
    ws2_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdSharedTable" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/table" '
        f'Target="{table_target}"/></Relationships>'
    ).encode("utf-8")
    parts["xl/worksheets/_rels/sheet2.xml.rels"] = ws2_rels
    p = tmp_path / "ambiguous_target.xlsx"
    _write_xlsx(p, parts)

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_calc_chain_relationship_rejected(tmp_path: Path, base_parts) -> None:
    calc_chain_xml = (
        '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<c r="A1" i="1"/></calcChain>'
    ).encode("utf-8")
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>", f'<Override PartName="/xl/calcChain.xml" ContentType="{_CT_CALC_CHAIN}"/></Types>'
    )
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    wb_rels2 = wb_rels.replace(
        "</Relationships>",
        f'<Relationship Id="rIdCalcChain1" Type="{_REL_CALC_CHAIN}" Target="calcChain.xml"/>'
        f'<Relationship Id="rIdCalcChain2" Type="{_REL_CALC_CHAIN}" Target="calcChain.xml"/></Relationships>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "dup_calc_chain_rel.xlsx"
    _write_xlsx(p, parts, [("xl/calcChain.xml", calc_chain_xml)])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_missing_office_document_relationship_rejected(tmp_path: Path, base_parts) -> None:
    root_rels = base_parts["_rels/.rels"].decode("utf-8")
    root_rels2 = re.sub(r'<Relationship [^>]*Type="[^"]*officeDocument"[^>]*/>', "", root_rels)
    assert root_rels2 != root_rels
    parts = dict(base_parts)
    parts["_rels/.rels"] = root_rels2.encode("utf-8")
    p = tmp_path / "missing_office_document.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_office_document_relationship_rejected(tmp_path: Path, base_parts) -> None:
    root_rels = base_parts["_rels/.rels"].decode("utf-8")
    m = re.search(r'<Relationship [^>]*Type="[^"]*officeDocument"[^>]*/>', root_rels)
    assert m is not None
    root_rels2 = root_rels.replace(
        "</Relationships>", m.group(0).replace('Id="rId1"', 'Id="rIdDupOfficeDoc"') + "</Relationships>"
    )
    parts = dict(base_parts)
    parts["_rels/.rels"] = root_rels2.encode("utf-8")
    p = tmp_path / "dup_office_document.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_missing_styles_relationship_rejected(tmp_path: Path, base_parts) -> None:
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    wb_rels2 = re.sub(r'<Relationship [^>]*Type="[^"]*relationships/styles"[^>]*/>', "", wb_rels)
    assert wb_rels2 != wb_rels
    parts = dict(base_parts)
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "missing_styles.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_styles_relationship_rejected(tmp_path: Path, base_parts) -> None:
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    m = re.search(r'<Relationship [^>]*Type="[^"]*relationships/styles"[^>]*/>', wb_rels)
    assert m is not None
    wb_rels2 = wb_rels.replace(
        "</Relationships>", m.group(0).replace('Id="rId3"', 'Id="rIdDupStyles"') + "</Relationships>"
    )
    parts = dict(base_parts)
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "dup_styles.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_missing_theme_relationship_rejected(tmp_path: Path, base_parts) -> None:
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    wb_rels2 = re.sub(r'<Relationship [^>]*Type="[^"]*relationships/theme"[^>]*/>', "", wb_rels)
    assert wb_rels2 != wb_rels
    parts = dict(base_parts)
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "missing_theme.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_duplicate_theme_relationship_rejected(tmp_path: Path, base_parts) -> None:
    wb_rels = base_parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
    m = re.search(r'<Relationship [^>]*Type="[^"]*relationships/theme"[^>]*/>', wb_rels)
    assert m is not None
    wb_rels2 = wb_rels.replace(
        "</Relationships>", m.group(0).replace('Id="rId4"', 'Id="rIdDupTheme"') + "</Relationships>"
    )
    parts = dict(base_parts)
    parts["xl/_rels/workbook.xml.rels"] = wb_rels2.encode("utf-8")
    p = tmp_path / "dup_theme.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_missing_workbook_rels_rejected(tmp_path: Path, base_parts) -> None:
    parts = {k: v for k, v in base_parts.items() if k != "xl/_rels/workbook.xml.rels"}
    p = tmp_path / "missing_workbook_rels.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_missing_root_rels_rejected(tmp_path: Path, base_parts) -> None:
    parts = {k: v for k, v in base_parts.items() if k != "_rels/.rels"}
    p = tmp_path / "missing_root_rels.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_missing_content_types_rejected(tmp_path: Path, base_parts) -> None:
    parts = {k: v for k, v in base_parts.items() if k != "[Content_Types].xml"}
    p = tmp_path / "missing_content_types.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_malformed_content_types_xml_rejected(tmp_path: Path, base_parts) -> None:
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = b"<not-valid-xml"
    p = tmp_path / "malformed_content_types.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_override_nonexistent_physical_part_rejected(tmp_path: Path, base_parts) -> None:
    ct = base_parts["[Content_Types].xml"].decode("utf-8")
    ct2 = ct.replace(
        "</Types>",
        f'<Override PartName="/xl/tables/table1.xml" ContentType="{_CT_TABLE}"/></Types>',
    )
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = ct2.encode("utf-8")
    p = tmp_path / "override_nonexistent.xlsx"
    _write_xlsx(p, parts)  # table1.xml НЕ добавлен физически
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


# ---------------------------------------------------------------------------
# E. Test quality / structural proofs / OD-7
# ---------------------------------------------------------------------------


def test_no_full_dom_parse_used(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "valid.xlsx"
    _write_xlsx(p, base_parts)

    def _boom(*args, **kwargs):  # noqa: ANN002, ANN003
        raise AssertionError("inventory не должен использовать ET.parse/ET.fromstring (full-DOM)")

    with mock.patch("app.scrub.inventory.ET.parse", _boom), mock.patch(
        "app.scrub.inventory.ET.fromstring", _boom
    ):
        result = inspect_package_policy(p)
    assert result.has_core_properties is True


def test_openpyxl_never_called(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "valid.xlsx"
    _write_xlsx(p, base_parts)

    def _boom(*args, **kwargs):  # noqa: ANN002, ANN003
        raise AssertionError("inspect_package_policy не должен вызывать openpyxl.load_workbook")

    with mock.patch("openpyxl.load_workbook", _boom):
        result = inspect_package_policy(p)
    assert result.has_core_properties is True


def test_no_extraction_no_temp_files(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "valid.xlsx"
    _write_xlsx(p, base_parts)
    before = set(tmp_path.iterdir())
    inspect_package_policy(p)
    after = set(tmp_path.iterdir())
    assert before == after  # никаких новых файлов/директорий не появилось


def test_sentinel_does_not_leak_on_malformed_content_types(tmp_path: Path, base_parts) -> None:
    sensitive_dir = tmp_path / PASSWORD_SENTINEL
    sensitive_dir.mkdir()
    parts = dict(base_parts)
    parts["[Content_Types].xml"] = b"<not-valid-xml"
    p = sensitive_dir / "malformed.xlsx"
    _write_xlsx(p, parts)
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    _assert_no_sentinel_leak(exc_info.value)


def test_sentinel_does_not_leak_on_unaccounted_member(tmp_path: Path, base_parts) -> None:
    sensitive_dir = tmp_path / PASSWORD_SENTINEL
    sensitive_dir.mkdir()
    p = sensitive_dir / "unaccounted.xlsx"
    _write_xlsx(p, base_parts, [(f"xl/{PASSWORD_SENTINEL}.xml", b"<x/>")])
    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    _assert_no_sentinel_leak(exc_info.value)


@pytest.mark.parametrize("reason", list(PackageScrubReason))
def test_package_scrub_error_messages_never_contain_sentinel(reason: PackageScrubReason) -> None:
    exc = PackageScrubError(reason)
    _assert_no_sentinel_leak(exc)


# ---------------------------------------------------------------------------
# F. PackageInventory model validation
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "field_name", ["comment_part_count", "table_part_count", "hyperlink_relationship_count"]
)
@pytest.mark.parametrize("bad_value", [-1, True])
def test_package_inventory_rejects_invalid_int_fields(field_name: str, bad_value: object) -> None:
    kwargs = dict(
        has_shared_strings=False,
        has_calc_chain=False,
        comment_part_count=0,
        table_part_count=0,
        hyperlink_relationship_count=0,
        has_core_properties=False,
        has_app_properties=False,
        has_custom_properties=False,
    )
    kwargs[field_name] = bad_value
    with pytest.raises(ValueError):
        PackageInventory(**kwargs)


@pytest.mark.parametrize(
    "field_name",
    ["has_shared_strings", "has_calc_chain", "has_core_properties", "has_app_properties", "has_custom_properties"],
)
def test_package_inventory_rejects_non_bool_flag_fields(field_name: str) -> None:
    kwargs = dict(
        has_shared_strings=False,
        has_calc_chain=False,
        comment_part_count=0,
        table_part_count=0,
        hyperlink_relationship_count=0,
        has_core_properties=False,
        has_app_properties=False,
        has_custom_properties=False,
    )
    kwargs[field_name] = 1  # int, не bool
    with pytest.raises(ValueError):
        PackageInventory(**kwargs)


def test_package_inventory_accepts_all_defaults() -> None:
    result = PackageInventory(
        has_shared_strings=False,
        has_calc_chain=False,
        comment_part_count=0,
        table_part_count=0,
        hyperlink_relationship_count=0,
        has_core_properties=False,
        has_app_properties=False,
        has_custom_properties=False,
    )
    assert result.comment_part_count == 0


# ---------------------------------------------------------------------------
# G. Correction Pass after Independent Security Review — MAJOR-1
# ---------------------------------------------------------------------------


def test_hyperlink_processing_never_calls_resolve_and_account(tmp_path: Path, base_parts) -> None:
    # Структурное доказательство (MAJOR-1 fix): hyperlink-relationships
    # только считаются, никогда не резолвятся/не сохраняются как
    # _RelationshipRecord. _resolve_and_account обязан вызываться лишь
    # ограниченное число раз (root officeDocument + workbook styles/theme)
    # НЕЗАВИСИМО от количества hyperlink-записей на листе.
    n = 10_000
    rel_tpl = (
        '<Relationship Id="rId{i}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="https://example.com/{i}"/>'
    )
    body = "".join(rel_tpl.format(i=i) for i in range(n))
    ws_rels_xml = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + body
        + "</Relationships>"
    ).encode("utf-8")
    p = tmp_path / "many_hyperlinks_structural.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels_xml)])

    from app.scrub import inventory as inv_module

    original = inv_module._resolve_and_account
    call_count = {"n": 0}

    def counting_wrapper(*args, **kwargs):
        call_count["n"] += 1
        return original(*args, **kwargs)

    with mock.patch("app.scrub.inventory._resolve_and_account", counting_wrapper):
        result = inspect_package_policy(p)

    assert result.hyperlink_relationship_count == n
    # root(officeDocument+core+app)=3 + workbook(worksheet+styles+theme)=3
    # -- ровно 6 для этой фикстуры, ни одного вызова на сами
    # hyperlink-записи, сколько бы их ни было (граница с запасом,
    # чтобы не быть хрупкой к мелким вариациям фикстуры).
    assert call_count["n"] <= 10, (
        f"_resolve_and_account вызван {call_count['n']} раз при {n} hyperlinks -- "
        "hyperlink-обработка не должна вызывать резолюцию/хранение записи"
    )


def test_worksheet_hyperlinks_large_count_still_correct(tmp_path: Path, base_parts) -> None:
    # Функциональная (не только memory-структурная) регрессия: большое
    # число hyperlinks по-прежнему корректно считается после MAJOR-1 fix.
    n = 5_000
    rel_tpl = (
        '<Relationship Id="rId{i}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="https://example.com/{i}"/>'
    )
    body = "".join(rel_tpl.format(i=i) for i in range(n))
    ws_rels_xml = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + body
        + "</Relationships>"
    ).encode("utf-8")
    p = tmp_path / "many_hyperlinks_functional.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels_xml)])

    result = inspect_package_policy(p)
    assert result.hyperlink_relationship_count == n


def test_duplicate_id_still_detected_among_many_hyperlinks(tmp_path: Path, base_parts) -> None:
    # Duplicate-Id detection обязана сохраняться после потоковой
    # переработки, даже когда дубликат встречается среди большого числа
    # записей.
    n = 5_000
    rel_tpl = (
        '<Relationship Id="rId{i}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="https://example.com/{i}"/>'
    )
    body = "".join(rel_tpl.format(i=i) for i in range(n))
    # Дубликат последнего Id.
    body += rel_tpl.format(i=n - 1)
    ws_rels_xml = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + body
        + "</Relationships>"
    ).encode("utf-8")
    p = tmp_path / "many_hyperlinks_duplicate.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels_xml)])

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


def test_malformed_tail_after_many_hyperlinks_still_detected(tmp_path: Path, base_parts) -> None:
    # Preserve XML tail validation (Correction Pass §11): потоковая
    # переработка не должна повторить историческую ошибку Stage10C.2.1 --
    # malformed-хвост ПОСЛЕ большого количества валидных hyperlink-записей
    # обязан обнаруживаться (документ дочитывается до EOF на успешном
    # пути, ранний break допустим только при уже доказанной ошибке).
    n = 1_000
    rel_tpl = (
        '<Relationship Id="rId{i}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="https://example.com/{i}"/>'
    )
    body = "".join(rel_tpl.format(i=i) for i in range(n))
    ws_rels_xml = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + body
        + "<<<not xml at all>>>&garbage;"
    ).encode("utf-8")
    p = tmp_path / "many_hyperlinks_malformed_tail.xlsx"
    _write_xlsx(p, base_parts, [("xl/worksheets/_rels/sheet1.xml.rels", ws_rels_xml)])

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE


# ---------------------------------------------------------------------------
# H. Correction Pass after Independent Security Review — MINOR-1
# ---------------------------------------------------------------------------


def test_inspect_package_policy_calls_full_preflight_first(tmp_path: Path, base_parts) -> None:
    p = tmp_path / "valid.xlsx"
    _write_xlsx(p, base_parts)

    from app.scrub import inventory as inv_module

    call_order: list[str] = []
    original_preflight = inv_module.preflight_xlsx_package
    original_after = inv_module._inspect_package_policy_after_preflight

    def spy_preflight(path):
        call_order.append("preflight")
        return original_preflight(path)

    def spy_after(path):
        call_order.append("after_preflight")
        return original_after(path)

    with mock.patch("app.scrub.inventory.preflight_xlsx_package", spy_preflight), mock.patch(
        "app.scrub.inventory._inspect_package_policy_after_preflight", spy_after
    ):
        inspect_package_policy(p)

    assert call_order == ["preflight", "after_preflight"]


def test_inspect_package_policy_never_reaches_inventory_logic_if_preflight_fails(
    tmp_path: Path,
) -> None:
    # Программная гарантия (не только docstring-конвенция): если
    # preflight_xlsx_package падает, package-policy логика вообще не
    # должна начинать выполняться.
    p = tmp_path / "not_a_zip.xlsx"
    p.write_bytes(b"not a zip file at all")

    from app.scrub import inventory as inv_module

    called = {"after_preflight": False}
    original_after = inv_module._inspect_package_policy_after_preflight

    def spy_after(path):
        called["after_preflight"] = True
        return original_after(path)

    with mock.patch("app.scrub.inventory._inspect_package_policy_after_preflight", spy_after):
        with pytest.raises(PackageScrubError) as exc_info:
            inspect_package_policy(p)

    assert exc_info.value.reason is PackageScrubReason.INVALID_INPUT_PACKAGE
    assert called["after_preflight"] is False


def test_inspect_package_policy_enforces_max_nonempty_cells(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # MINOR-1: прямой вызов inspect_package_policy(path) обязан
    # самостоятельно (программно) обеспечивать проверку MAX_NONEMPTY_CELLS
    # (Stage10C.2.1 Section B) -- небольшая representative-фикстура с
    # заниженным лимитом, без гигантского реального файла.
    monkeypatch.setattr("app.scrub.preflight.MAX_NONEMPTY_CELLS", 5)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 11):  # 10 непустых ячеек > лимита 5
        ws.cell(row=r, column=1, value=f"v{r}")
    p = tmp_path / "too_many_cells.xlsx"
    wb.save(p)
    wb.close()

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.RESOURCE_LIMIT_EXCEEDED


def test_inspect_package_policy_max_worksheets_still_enforced(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Сохранена и MAX_WORKSHEETS-проверка (унаследована независимо от
    # MINOR-1 fix через переиспользуемый _read_workbook_sheet_rids).
    monkeypatch.setattr("app.scrub.preflight.MAX_WORKSHEETS", 2)

    wb = openpyxl.Workbook()
    wb.active.title = "S0"
    wb.create_sheet("S1")
    wb.create_sheet("S2")
    p = tmp_path / "too_many_sheets.xlsx"
    wb.save(p)
    wb.close()

    with pytest.raises(PackageScrubError) as exc_info:
        inspect_package_policy(p)
    assert exc_info.value.reason is PackageScrubReason.RESOURCE_LIMIT_EXCEEDED
