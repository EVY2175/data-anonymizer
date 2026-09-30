"""
Тесты Stage 10C.2.4 — Strict Post-Save Package Validation.

Методология: положительные тесты получают РЕАЛЬНЫЙ scrubbed-output через
app.scrub.mutate.scrub_workbook_object_model(source) и проверяют, что
validate_scrubbed_workbook_package(result) успешно проходит и возвращает
корректный PackageValidationResult. Отрицательные тесты получают
РЕАЛЬНЫЙ scrubbed-output и хирургически мутируют РОВНО ОДНО свойство
пакета (один байт-паттерн в одной XML-части или структуру ZIP), затем
проверяют, что validate_scrubbed_workbook_package поднимает
PackageScrubError с ожидаемым reason. Искусственные ZIP "с нуля"
используются только там, где получить нужный сценарий через реальный
scrub принципиально невозможно (整 malformed XML, отсутствующая часть).
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import NamedStyle
from openpyxl.styles.colors import Color
from openpyxl.styles.fills import GradientFill, PatternFill, Stop
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.scrub.errors import PackageScrubError, PackageScrubReason
from app.scrub.mutate import scrub_workbook_object_model
import app.scrub.postvalidate as postvalidate
from app.scrub.postvalidate import PackageValidationResult, validate_scrubbed_workbook_package

SECRET = "SECRET_CONFIDENTIAL_PACKAGE_VALUE"


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------


def _save_and_scrub(tmp_path: Path, wb, name: str = "src.xlsx") -> Path:
    src = tmp_path / name
    wb.save(src)
    wb.close()
    return scrub_workbook_object_model(src)


def _read_part(path: Path, part_name: str) -> bytes:
    with zipfile.ZipFile(path) as z:
        return z.read(part_name)


def _mutate_zip(src: Path, dst: Path, edits: dict, add: dict | None = None) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in edits:
                data = edits[item.filename](data)
            zout.writestr(item, data)
        if add:
            for name, data in add.items():
                zout.writestr(name, data)
    return dst


def _mutate_one_part(tmp_path: Path, scrubbed: Path, part_name: str, transform, out_name: str) -> Path:
    dst = tmp_path / out_name
    return _mutate_zip(scrubbed, dst, {part_name: transform})


def _assert_rejected(path: Path, expected_reason: PackageScrubReason) -> None:
    with pytest.raises(PackageScrubError) as excinfo:
        validate_scrubbed_workbook_package(path)
    assert excinfo.value.reason == expected_reason


# ----------------------------------------------------------------------
# Positive: базовая форма / result-поля / SHA неизменность источника
# ----------------------------------------------------------------------


def test_minimal_single_sheet_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert isinstance(r, PackageValidationResult)
    assert r.worksheet_count == 1
    assert r.zip_entry_count >= 9
    assert r.package_size_bytes > 0
    assert r.formula_cell_count == 0


def test_multi_sheet_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "One"
    wb.active["A1"] = 1
    wb.create_sheet("Two")["A1"] = 2
    wb.create_sheet("Three")["A1"] = 3
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.worksheet_count == 3


def test_hidden_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = 2
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.worksheet_count == 2


def test_merged_cells(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.merge_cells("A1:B2")
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_dimension_variants(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["C5"] = 1
    ws["A1"] = 2
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_freeze_panes(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.freeze_panes = "B2"
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_page_margins(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.page_margins.left = 0.3
    ws.page_margins.top = 1.5
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_print_options(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.print_options.horizontalCentered = True
    ws.print_options.gridLines = True
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_page_setup_without_rid(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)
    raw = _read_part(result, "xl/worksheets/sheet1.xml").decode("utf-8")
    assert ' id="' not in raw.replace('r:id', '')
    assert "r:id" not in raw


def test_plain_inline_strings(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello world"
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_whitespace_inline_string_xml_space(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "  leading and trailing  "
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)
    raw = _read_part(result, "xl/worksheets/sheet1.xml").decode("utf-8")
    assert 'xml:space="preserve"' in raw


def test_rich_text(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = CellRichText([TextBlock(InlineFont(b=True, color="FF0000"), "bold red")])
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_rich_text_whitespace(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = CellRichText([TextBlock(InlineFont(b=True), "  padded  ")])
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)
    raw = _read_part(result, "xl/worksheets/sheet1.xml").decode("utf-8")
    assert 'xml:space="preserve"' in raw


def test_rich_text_allowed_inline_font_fields(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    font = InlineFont(
        rFont="Calibri Light",
        b=True,
        i=True,
        u="double",
        strike=True,
        sz=14,
        color="00FF00FF",
        vertAlign="superscript",
        family=2,
        charset=1,
        scheme="minor",
    )
    ws["A1"] = CellRichText([TextBlock(font, "styled run")])
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_normal_formula(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = 2
    ws["C1"] = "=A1+B1"
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.formula_cell_count == 1


def test_array_formula_structure(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = ArrayFormula("A1:A1", "=SUM(B1:B2)")
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.formula_cell_count == 1


def test_datatable_formula_structure(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B3"] = DataTableFormula(ref="B3:B3", r1="C1")
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.formula_cell_count == 1


def test_trusted_builtin_style(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ns = NamedStyle(name="Currency")
    ns.builtinId = 4
    wb.add_named_style(ns)
    ws["A1"] = 1
    ws["A1"].style = "Currency"
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_dastyle_custom_style(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ns = NamedStyle(name=SECRET)
    wb.add_named_style(ns)
    ws["A1"] = 1
    ws["A1"].style = SECRET
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)
    raw = _read_part(result, "xl/styles.xml").decode("utf-8")
    assert SECRET not in raw
    assert "DAStyle" in raw


def test_allowed_custom_number_format(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 5
    ws["A1"].number_format = "dd.mm.yyyy hh:mm"
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_exactly_max_worksheets(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i in range(200):
        wb.create_sheet(f"S{i}")["A1"] = i
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.worksheet_count == 200


def test_bounded_nonempty_cell_limit_boundary(tmp_path, monkeypatch):
    monkeypatch.setattr(postvalidate, "MAX_NONEMPTY_CELLS", 5)
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(5):
        ws.cell(row=i + 1, column=1, value=i)
    result = _save_and_scrub(tmp_path, wb, "ok.xlsx")
    r = validate_scrubbed_workbook_package(result)
    assert r is not None

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    for i in range(6):
        ws2.cell(row=i + 1, column=1, value=i)
    result2 = _save_and_scrub(tmp_path, wb2, "over.xlsx")
    with pytest.raises(PackageScrubError) as excinfo:
        validate_scrubbed_workbook_package(result2)
    assert excinfo.value.reason == PackageScrubReason.RESOURCE_LIMIT_EXCEEDED


def test_result_fields_correct(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["B1"] = "=A1+1"
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    with zipfile.ZipFile(result) as z:
        assert r.zip_entry_count == len(z.infolist())
    assert r.package_size_bytes == result.stat().st_size
    assert r.worksheet_count == 1
    assert r.formula_cell_count == 1


def test_source_sha256_unchanged_by_validation(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    result = _save_and_scrub(tmp_path, wb)
    before = hashlib.sha256(result.read_bytes()).hexdigest()
    validate_scrubbed_workbook_package(result)
    after = hashlib.sha256(result.read_bytes()).hexdigest()
    assert before == after


def test_type_error_on_invalid_path_type():
    with pytest.raises(TypeError):
        validate_scrubbed_workbook_package(12345)  # type: ignore[arg-type]


def test_wrong_extension_rejected(tmp_path):
    fake = tmp_path / "not_xlsx.txt"
    fake.write_bytes(b"hello")
    with pytest.raises(PackageScrubError) as excinfo:
        validate_scrubbed_workbook_package(fake)
    assert excinfo.value.reason == PackageScrubReason.INVALID_INPUT_PACKAGE


# ----------------------------------------------------------------------
# Negative: ZIP / package structure
# ----------------------------------------------------------------------


@pytest.fixture
def scrubbed(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["B1"] = 5
    ws["C1"] = "=A1&B1"
    ns = NamedStyle(name="Currency")
    ns.builtinId = 4
    wb.add_named_style(ns)
    ws["A1"].style = "Currency"
    ws["D1"] = CellRichText([TextBlock(InlineFont(b=True), "rich")])
    return _save_and_scrub(tmp_path, wb)


def test_extra_unexpected_part_rejected(tmp_path, scrubbed):
    out = _mutate_zip(scrubbed, tmp_path / "extra.xlsx", {}, add={"xl/extra_secret.xml": b"<root/>"})
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_orphan_worksheet_part_rejected(tmp_path, scrubbed):
    """
    Физически присутствующая worksheet-подобная часть с корректным
    Content-Type Override, но БЕЗ relationship-ссылки из
    xl/_rels/workbook.xml.rels — triple cross-check (физическая часть ↔
    content-type ↔ relationship) обязан отклонить orphan, а не молча
    проигнорировать лишний файл.
    """
    with zipfile.ZipFile(scrubbed) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    orphan_path = "xl/worksheets/sheet99.xml"
    data[orphan_path] = data["xl/worksheets/sheet1.xml"]
    ct = data["[Content_Types].xml"].decode("utf-8")
    ct = ct.replace(
        "</Types>",
        f'<Override PartName="/{orphan_path}" '
        f'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
    )
    data["[Content_Types].xml"] = ct.encode("utf-8")
    dst = tmp_path / "orphan.xlsx"
    with zipfile.ZipFile(dst, "w") as zout:
        for name, content in data.items():
            zout.writestr(name, content)
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_directory_marker_rejected(tmp_path, scrubbed):
    dst = tmp_path / "dirmarker.xlsx"
    with zipfile.ZipFile(scrubbed) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/worksheets/", b"")
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_missing_content_types_rejected(tmp_path, scrubbed):
    dst = tmp_path / "missing_ct.xlsx"
    with zipfile.ZipFile(scrubbed) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            if item.filename == "[Content_Types].xml":
                continue
            zout.writestr(item, zin.read(item.filename))
    with pytest.raises(PackageScrubError):
        validate_scrubbed_workbook_package(dst)


def test_malformed_worksheet_xml_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path, scrubbed, "xl/worksheets/sheet1.xml", lambda d: d[:-5], "malformed.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: [Content_Types].xml
# ----------------------------------------------------------------------


def test_unknown_default_extension_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "[Content_Types].xml",
        lambda d: d.replace(
            b"</Types>",
            b'<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/></Types>',
        ),
        "bad_default.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_forbidden_override_content_type_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "[Content_Types].xml",
        lambda d: d.replace(
            b"</Types>",
            b'<Override PartName="/xl/sharedStrings.xml" '
            b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        ),
        "bad_override.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_duplicate_override_rejected(tmp_path, scrubbed):
    def dup(data: bytes) -> bytes:
        marker = b'<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml" />'
        return data.replace(marker, marker + marker, 1)

    out = _mutate_one_part(tmp_path, scrubbed, "[Content_Types].xml", dup, "dup_override.xlsx")
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: relationships
# ----------------------------------------------------------------------


def test_root_rels_external_target_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "_rels/.rels",
        lambda d: d.replace(b"Target=\"docProps/core.xml\"", b'Target="http://evil.example/x" TargetMode="External"'),
        "root_external.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_root_rels_duplicate_office_document_rejected(tmp_path, scrubbed):
    def dup(data: bytes) -> bytes:
        marker = (
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            b'Target="xl/workbook.xml" Id="rId1" />'
        )
        return data.replace(marker, marker.replace(b'Id="rId1"', b'Id="rIdX"') + marker, 1)

    out = _mutate_one_part(tmp_path, scrubbed, "_rels/.rels", dup, "root_dup.xlsx")
    _assert_rejected(out, PackageScrubError and PackageScrubReason.POST_VALIDATION_FAILED)


def test_workbook_rels_sharedstrings_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/_rels/workbook.xml.rels",
        lambda d: d.replace(
            b"</Relationships>",
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" '
            b'Target="sharedStrings.xml" Id="rIdShared" /></Relationships>',
        ),
        "shared_strings_rel.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_workbook_rels_missing_theme_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/_rels/workbook.xml.rels",
        lambda d: d.replace(
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" '
            b'Target="theme/theme1.xml" Id="rId3" />',
            b"",
        ),
        "missing_theme_rel.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: workbook.xml
# ----------------------------------------------------------------------


def test_defined_names_with_content_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/workbook.xml",
        lambda d: d.replace(
            b"<definedNames />",
            b'<definedNames><definedName name="Secret">Sheet1!$A$1</definedName></definedNames>',
        ),
        "defined_names.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_external_references_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/workbook.xml",
        lambda d: d.replace(b"<calcPr", b"<externalReferences/><calcPr"),
        "external_refs.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_duplicate_sheet_name_case_insensitive_rejected(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.create_sheet("Extra")["A1"] = 1
    src = tmp_path / "dup_src.xlsx"
    wb.save(src)
    wb.close()
    result = scrub_workbook_object_model(src)
    out = _mutate_one_part(
        tmp_path,
        result,
        "xl/workbook.xml",
        lambda d: d.replace(b'name="Extra"', b'name="DATA"'),
        "dup_case.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_workbook_pr_codename_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/workbook.xml",
        lambda d: d.replace(b"<workbookPr />", f'<workbookPr codeName="{SECRET}" />'.encode()),
        "workbookpr_codename.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: worksheet.xml
# ----------------------------------------------------------------------


def test_sheetpr_codename_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"<sheetPr>", f'<sheetPr codeName="{SECRET}">'.encode())
        if b"<sheetPr>" in d
        else d.replace(b"<sheetPr", f'<sheetPr codeName="{SECRET}"'.encode(), 1),
        "sheetpr_codename.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_pagesetup_rid_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"</worksheet>", b'<pageSetup r:id="rIdSecret"/></worksheet>')
        if b"<pageSetup" not in d
        else d.replace(b"<pageSetup", b'<pageSetup r:id="rIdSecret"', 1),
        "pagesetup_rid.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_extlst_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"</worksheet>", b"<extLst/></worksheet>"),
        "extlst.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_alternate_content_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"</worksheet>", b'<mc:AlternateContent xmlns:mc="x"/></worksheet>'),
        "mc_alt.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_unknown_worksheet_attribute_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"<dimension ", b'<dimension secret="1" ', 1),
        "unknown_attr.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: cells
# ----------------------------------------------------------------------


def test_shared_string_type_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b't="inlineStr"', b't="s"', 1),
        "shared_string.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_shared_formula_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"<f>", b'<f t="shared" si="0">', 1),
        "shared_formula.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_cell_style_index_out_of_bounds_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b'<c r="B1"', b'<c r="B1" s="999"', 1),
        "cell_s_oob.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_v_element_with_attribute_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"<v>5</v>", b'<v secret="1">5</v>'),
        "v_attr.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: rich text
# ----------------------------------------------------------------------


def test_rich_text_unknown_rfont_rejected(tmp_path, scrubbed):
    raw = _read_part(scrubbed, "xl/worksheets/sheet1.xml")
    assert b"<rPr>" in raw
    mutated = raw.replace(b"<rPr>", b'<rPr><rFont val="EvilFont"/>', 1)
    dst = tmp_path / "bad_rfont.xlsx"
    _mutate_zip(scrubbed, dst, {"xl/worksheets/sheet1.xml": lambda d: mutated})
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_rich_text_unknown_child_rejected(tmp_path, scrubbed):
    raw = _read_part(scrubbed, "xl/worksheets/sheet1.xml")
    mutated = raw.replace(b"<rPr>", b'<rPr><phoneticPr/>', 1)
    dst = tmp_path / "phonetic.xlsx"
    _mutate_zip(scrubbed, dst, {"xl/worksheets/sheet1.xml": lambda d: mutated})
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: styles.xml
# ----------------------------------------------------------------------


def test_unknown_font_name_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<name val="Calibri" />', b'<name val="EvilFont" />', 1),
        "bad_font.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_unknown_custom_numfmt_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(
            b'<numFmts count="0" />',
            f'<numFmts count="1"><numFmt numFmtId="164" formatCode="{SECRET}"/></numFmts>'.encode(),
        ),
        "bad_numfmt.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_spoofed_builtin_name_rejected(tmp_path, scrubbed):
    raw = _read_part(scrubbed, "xl/styles.xml")
    assert b'name="Currency" xfId="1" builtinId="4" hidden="0"' in raw
    mutated = raw.replace(
        b'name="Currency" xfId="1" builtinId="4" hidden="0"',
        f'name="{SECRET}" xfId="1" builtinId="4" hidden="0"'.encode(),
    )
    dst = tmp_path / "spoofed_name.xlsx"
    _mutate_zip(scrubbed, dst, {"xl/styles.xml": lambda d: mutated})
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_dastyle_name_bad_shape_rejected(tmp_path):
    import re as _re

    wb = openpyxl.Workbook()
    ws = wb.active
    ns = NamedStyle(name=SECRET)
    wb.add_named_style(ns)
    ws["A1"] = 1
    ws["A1"].style = SECRET
    result = _save_and_scrub(tmp_path, wb)

    raw = _read_part(result, "xl/styles.xml")
    m = _re.search(rb'cellStyle name="(DAStyle\d{4})"', raw)
    assert m is not None
    dstyle_name = m.group(1)
    mutated = raw.replace(dstyle_name, b"DAStyleXXXX")
    dst = tmp_path / "bad_dastyle.xlsx"
    _mutate_zip(result, dst, {"xl/styles.xml": lambda d: mutated})
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_custom_table_style_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'defaultTableStyle="TableStyleMedium9"', b'defaultTableStyle="Custom1"'),
        "custom_table_style.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_font_id_out_of_bounds_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'fontId="0" fillId="0" borderId="0" pivotButton', b'fontId="999" fillId="0" borderId="0" pivotButton', 1),
        "font_oob.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Negative: package-fixed (theme/app hash, core.xml)
# ----------------------------------------------------------------------


def test_theme_hash_mismatch_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path, scrubbed, "xl/theme/theme1.xml", lambda d: d + b"<!-- tamper -->", "theme_tamper.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_app_xml_hash_mismatch_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path, scrubbed, "docProps/app.xml", lambda d: d.replace(b"3.1", b"9.9"), "app_tamper.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_core_xml_bad_created_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "docProps/core.xml",
        lambda d: d.replace(b"2000-01-01T00:00:00Z", b"2020-05-05T00:00:00Z", 1),
        "core_bad_created.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_core_xml_custom_element_rejected(tmp_path, scrubbed):
    out = _mutate_one_part(
        tmp_path,
        scrubbed,
        "docProps/core.xml",
        lambda d: d.replace(
            b"</cp:coreProperties>",
            f'<dc:creator xmlns:dc="http://purl.org/dc/elements/1.1/">{SECRET}</dc:creator></cp:coreProperties>'.encode(),
        ),
        "core_custom.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_custom_xml_part_rejected(tmp_path, scrubbed):
    dst = tmp_path / "custom_xml.xlsx"
    _mutate_zip(
        scrubbed,
        dst,
        {},
        add={"customXml/item1.xml": f"<root>{SECRET}</root>".encode()},
    )
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_shared_strings_part_rejected(tmp_path, scrubbed):
    dst = tmp_path / "shared_strings_part.xlsx"
    _mutate_zip(
        scrubbed,
        dst,
        {},
        add={"xl/sharedStrings.xml": b'<sst xmlns="x" count="0" uniqueCount="0"/>'},
    )
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_comments_part_rejected(tmp_path, scrubbed):
    dst = tmp_path / "comments_part.xlsx"
    _mutate_zip(
        scrubbed,
        dst,
        {},
        add={"xl/comments1.xml": b'<comments xmlns="x"/>'},
    )
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


def test_drawing_part_rejected(tmp_path, scrubbed):
    dst = tmp_path / "drawing_part.xlsx"
    _mutate_zip(
        scrubbed,
        dst,
        {},
        add={"xl/drawings/drawing1.xml": b'<xdr:wsDr xmlns:xdr="x"/>'},
    )
    _assert_rejected(dst, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# Error-privacy
# ----------------------------------------------------------------------


def test_error_message_never_leaks_secret_from_cell_value(tmp_path):
    # Значения ячеек — зона ответственности Stage8 anonymization (уже
    # выполненного ДО scrub), а не Stage10C.2.3/.2.4: scrub_workbook_object_model
    # умышленно не трогает cell values, поэтому SECRET здесь ожидаемо
    # переживает scrub как обычное (уже якобы анонимизированное) значение —
    # тест проверяет ТОЛЬКО, что сообщение об ошибке post-validation (для
    # НЕСВЯЗАННОГО grammar-нарушения в отдельной части пакета) не включает
    # это значение, а не то, что оно стирается откуда-либо ещё.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = SECRET
    result = _save_and_scrub(tmp_path, wb)

    dst = tmp_path / "with_secret_in_extra_part.xlsx"
    _mutate_zip(result, dst, {}, add={"xl/extra_secret.xml": SECRET.encode()})
    try:
        validate_scrubbed_workbook_package(dst)
        assert False, "expected PackageScrubError"
    except PackageScrubError as exc:
        assert SECRET not in str(exc)
        assert SECRET not in repr(exc)
        assert exc.__context__ is None
        assert exc.__cause__ is None


def test_error_message_never_leaks_secret_from_font_name(tmp_path, scrubbed):
    dst = tmp_path / "font_secret.xlsx"
    _mutate_zip(
        scrubbed,
        dst,
        {"xl/styles.xml": lambda d: d.replace(b'<name val="Calibri" />', f'<name val="{SECRET}" />'.encode(), 1)},
    )
    try:
        validate_scrubbed_workbook_package(dst)
        assert False, "expected PackageScrubError"
    except PackageScrubError as exc:
        assert SECRET not in str(exc)
        assert SECRET not in repr(exc)


# ----------------------------------------------------------------------
# PackageValidationResult contract
# ----------------------------------------------------------------------


def test_result_rejects_bool_for_int_fields():
    with pytest.raises(ValueError):
        PackageValidationResult(
            zip_entry_count=True,  # type: ignore[arg-type]
            package_size_bytes=100,
            worksheet_count=1,
            formula_cell_count=0,
        )


def test_result_rejects_zip_entry_count_below_minimum():
    with pytest.raises(ValueError):
        PackageValidationResult(
            zip_entry_count=8, package_size_bytes=100, worksheet_count=1, formula_cell_count=0
        )


def test_result_rejects_zero_package_size():
    with pytest.raises(ValueError):
        PackageValidationResult(
            zip_entry_count=9, package_size_bytes=0, worksheet_count=1, formula_cell_count=0
        )


def test_result_rejects_zero_worksheet_count():
    with pytest.raises(ValueError):
        PackageValidationResult(
            zip_entry_count=9, package_size_bytes=100, worksheet_count=0, formula_cell_count=0
        )


def test_result_rejects_negative_formula_count():
    with pytest.raises(ValueError):
        PackageValidationResult(
            zip_entry_count=9, package_size_bytes=100, worksheet_count=1, formula_cell_count=-1
        )


# ----------------------------------------------------------------------
# BLOCKER-1 Correction Pass — regression tests
# ----------------------------------------------------------------------
#
# Фикстура richer_scrubbed ниже — более насыщенный, чем `scrubbed`, реальный
# scrub-output (formula/rich-text/named style/несколько ячеек), используемый
# как база для text/tail/comment/PI/mixed-content инъекций во всех 7 типах
# XML-частей пакета.


@pytest.fixture
def richer_scrubbed(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["B1"] = 5
    ws["C1"] = "=A1&B1"
    ns = NamedStyle(name="Currency")
    ns.builtinId = 4
    wb.add_named_style(ns)
    ws["A1"].style = "Currency"
    ws["D1"] = CellRichText([TextBlock(InlineFont(b=True), "rich")])
    return _save_and_scrub(tmp_path, wb, "richer_src.xlsx")


_ALL_PACKAGE_PARTS = (
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
    "xl/worksheets/sheet1.xml",
    "xl/styles.xml",
    "docProps/core.xml",
)


def _assert_rejected_with_privacy(path: Path, expected_reason: PackageScrubReason) -> None:
    with pytest.raises(PackageScrubError) as excinfo:
        validate_scrubbed_workbook_package(path)
    assert excinfo.value.reason == expected_reason
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None


# ---- §11: comments/PI matrix (7 частей × 2 инъекции = 14 тестов) ----


@pytest.mark.parametrize("part_name", _ALL_PACKAGE_PARTS)
def test_xml_comment_before_root_rejected(tmp_path, richer_scrubbed, part_name):
    injection = f"<!--{SECRET}-->".encode("utf-8")
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, part_name, lambda d, inj=injection: inj + d, f"comment_{part_name.replace('/', '_')}.xlsx"
    )
    _assert_rejected_with_privacy(out, PackageScrubReason.POST_VALIDATION_FAILED)


@pytest.mark.parametrize("part_name", _ALL_PACKAGE_PARTS)
def test_xml_pi_before_root_rejected(tmp_path, richer_scrubbed, part_name):
    injection = f"<?evil {SECRET}?>".encode("utf-8")
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, part_name, lambda d, inj=injection: inj + d, f"pi_{part_name.replace('/', '_')}.xlsx"
    )
    _assert_rejected_with_privacy(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §12: non-whitespace elem.text matrix ----


def test_text_content_types_root_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "[Content_Types].xml", lambda d: d.replace(b"<Default", secret + b"<Default", 1), "t1.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_workbook_root_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/workbook.xml", lambda d: d.replace(b"<workbookPr", secret + b"<workbookPr", 1), "t2.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_worksheet_root_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/worksheets/sheet1.xml", lambda d: d.replace(b"<sheetPr", secret + b"<sheetPr", 1), "t3.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_sheetdata_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/worksheets/sheet1.xml", lambda d: d.replace(b"<sheetData>", b"<sheetData>" + secret, 1), "t4.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_dimension_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b'<dimension ref="A1:D1" />', b'<dimension ref="A1:D1">' + secret + b"</dimension>"),
        "t5.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_row_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/worksheets/sheet1.xml", lambda d: d.replace(b'<row r="1">', b'<row r="1">' + secret, 1), "t6.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_styles_root_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b"<numFmts", secret + b"<numFmts", 1), "t7.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_fonts_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b"<fonts", secret + b"<fonts", 1), "t8.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_relationships_root_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "_rels/.rels", lambda d: d.replace(b"<Relationship", secret + b"<Relationship", 1), "t9.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_text_coreproperties_root_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "docProps/core.xml", lambda d: d.replace(b"<dcterms:created", secret + b"<dcterms:created", 1), "t10.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §12: non-whitespace elem.tail matrix ----


def test_tail_dimension_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b'<dimension ref="A1:D1" />', b'<dimension ref="A1:D1" />' + secret),
        "tail1.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_tail_row_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/worksheets/sheet1.xml", lambda d: d.replace(b"</row>", b"</row>" + secret, 1), "tail2.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_tail_relationship_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "_rels/.rels",
        lambda d: d.replace(b"</Relationships>", secret + b"</Relationships>"),
        "tail3.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_tail_font_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b"</fonts>", b"</fonts>" + secret, 1), "tail4.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_tail_created_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "docProps/core.xml",
        lambda d: d.replace(b"</dcterms:created>", b"</dcterms:created>" + secret, 1),
        "tail5.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §13: positive whitespace (pretty-printed XML must still PASS) ----


def _pretty_print_whitespace(data: bytes) -> bytes:
    """Вставляет форматирующие пробелы/переносы строк МЕЖДУ тегами (никогда
    внутрь текстовых узлов <v>/<f>/<t>/dcterms:*) — заменяет "><" на ">\n  <".
    """
    return data.replace(b"><", b">\n  <")


@pytest.mark.parametrize(
    "part_name",
    (
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
        "xl/styles.xml",
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/core.xml",
    ),
)
def test_pretty_printed_whitespace_passes(tmp_path, richer_scrubbed, part_name):
    out = _mutate_one_part(tmp_path, richer_scrubbed, part_name, _pretty_print_whitespace, f"pretty_{part_name.replace('/', '_')}.xlsx")
    validate_scrubbed_workbook_package(out)


# ---- §15: mixed content (non-whitespace text просачивается МИМО дочернего
# структурного элемента, а не внутрь текстового) ----


def test_mixed_content_row_before_cell_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b'<row r="1"><c r="A1"', b'<row r="1">' + secret + b'<c r="A1"', 1),
        "mixed1.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_mixed_content_sheetdata_before_row_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b'<sheetData><row', b"<sheetData>" + secret + b"<row", 1),
        "mixed2.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_mixed_content_font_before_name_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b"<font><name", b"<font>" + secret + b"<name", 1),
        "mixed3.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_mixed_content_worksheet_before_sheetpr_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/worksheets/sheet1.xml",
        lambda d: d.replace(b"<sheetPr>", secret + b"<sheetPr>", 1),
        "mixed4.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_mixed_content_workbook_before_workbookpr_rejected(tmp_path, richer_scrubbed):
    secret = SECRET.encode()
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/workbook.xml",
        lambda d: d.replace(b"<workbookPr", secret + b"<workbookPr", 1),
        "mixed5.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §18: MINOR-1 count consistency ----


def test_fonts_count_correct_passes(tmp_path, richer_scrubbed):
    validate_scrubbed_workbook_package(richer_scrubbed)


def test_fonts_count_minus_one_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b'<fonts count="2">', b'<fonts count="1">', 1), "cnt1.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fonts_count_plus_one_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b'<fonts count="2">', b'<fonts count="3">', 1), "cnt2.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fonts_count_negative_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b'<fonts count="2">', b'<fonts count="-1">', 1), "cnt3.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fonts_count_non_integer_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b'<fonts count="2">', b'<fonts count="abc">', 1), "cnt4.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fills_count_mismatch_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b'<fills count="2">', b'<fills count="1">', 1), "cnt5.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_cellxfs_count_mismatch_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/styles.xml", lambda d: d.replace(b'<cellXfs count="2">', b'<cellXfs count="1">', 1), "cnt6.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_cellstyles_count_mismatch_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<cellStyles count="2">', b'<cellStyles count="1">', 1),
        "cnt7.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_numfmts_count_mismatch_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(
            b'<numFmts count="0" />',
            b'<numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy-mm-dd"/></numFmts>'.replace(b'count="1"', b'count="2"'),
        ),
        "cnt8.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_dxfs_declared_nonzero_count_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b"</cellStyles>", b'</cellStyles><dxfs count="1"/>', 1),
        "cnt9.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §20: MINOR-2 color semantics ----


def test_color_rgb_only_passes(tmp_path):
    wb = openpyxl.Workbook()
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color

    wb.active["A1"] = 1
    wb.active["A1"].font = Font(color=Color(rgb="FF112233"))
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_color_indexed_only_passes(tmp_path):
    wb = openpyxl.Workbook()
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color

    wb.active["A1"] = 1
    wb.active["A1"].font = Font(color=Color(indexed=5))
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_color_theme_only_passes(tmp_path):
    wb = openpyxl.Workbook()
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color

    wb.active["A1"] = 1
    wb.active["A1"].font = Font(color=Color(theme=3))
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_color_auto_only_passes(tmp_path):
    wb = openpyxl.Workbook()
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color

    wb.active["A1"] = 1
    wb.active["A1"].font = Font(color=Color(auto=True))
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_color_theme_with_tint_passes(tmp_path):
    wb = openpyxl.Workbook()
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color

    wb.active["A1"] = 1
    wb.active["A1"].font = Font(color=Color(theme=3, tint=0.5))
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_color_rgb_plus_theme_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color theme="1" rgb="FF00FF00" />', 1),
        "color1.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_color_rgb_plus_indexed_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color rgb="FF00FF00" indexed="3" />', 1),
        "color2.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_color_theme_plus_indexed_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color theme="1" indexed="3" />', 1),
        "color3.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_color_type_attribute_present_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color theme="1" type="theme" />', 1),
        "color4.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_color_type_only_without_mode_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color type="theme" />', 1),
        "color5.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_color_zero_modes_only_tint_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color tint="0.5" />', 1),
        "color6.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_color_auto_plus_rgb_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b'<color theme="1" />', b'<color auto="1" rgb="FF00FF00" />', 1),
        "color7.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §22: MINOR-3 sheet title structural grammar ----


def test_sheet_title_one_char_passes(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.active["A1"] = 1
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_sheet_title_31_chars_passes(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "A" * 31
    wb.active["A1"] = 1
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_sheet_title_unicode_passes(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Отчёт Q3"
    wb.active["A1"] = 1
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_sheet_title_with_spaces_passes(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "My Sheet Name"
    wb.active["A1"] = 1
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


def test_sheet_title_case_distinct_names_pass(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = 1
    wb.create_sheet("Report")["A1"] = 2
    result = _save_and_scrub(tmp_path, wb)
    r = validate_scrubbed_workbook_package(result)
    assert r.worksheet_count == 2


def test_sheet_title_more_than_31_chars_survives_pipeline_and_passes(tmp_path):
    """
    MINOR-3 Correction Pass: эмпирически подтверждено (см. production
    докстринг _sheet_name_shape), что openpyxl==3.1.5 лишь предупреждает
    (UserWarning), но НЕ отклоняет имя листа длиннее 31 символа, и такое
    имя РЕАЛЬНО переживает полный scrub_workbook_object_model pipeline —
    поэтому 2.4 сознательно НЕ вводит верхнюю границу длины.
    """
    import warnings

    wb = openpyxl.Workbook()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb.active.title = "A" * 40
    wb.active["A1"] = 1
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


@pytest.mark.parametrize("forbidden_char", list("[]:*?/\\"))
def test_sheet_title_forbidden_char_rejected(tmp_path, richer_scrubbed, forbidden_char):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/workbook.xml",
        lambda d, ch=forbidden_char: d.replace(b'name="Sheet"', f'name="She{ch}et"'.encode("utf-8"), 1),
        f"badname_{ord(forbidden_char)}.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_sheet_title_control_char_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/workbook.xml",
        lambda d: d.replace(b'name="Sheet"', "name=\"She{}et\"".format(chr(0x7F)).encode("utf-8"), 1),
        "badname_ctrl.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_sheet_title_empty_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path, richer_scrubbed, "xl/workbook.xml", lambda d: d.replace(b'name="Sheet"', b'name=""', 1), "badname_empty.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ----------------------------------------------------------------------
# GradientFill Grammar v1 (OD-10C2.4-R1) — Implementation Pass
# ----------------------------------------------------------------------


def _build_gradient_fill_workbook(tmp_path: Path, fill, name: str = "gf_src.xlsx") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = fill
    return _save_and_scrub(tmp_path, wb, name)


@pytest.fixture
def gradient_scrubbed(tmp_path) -> Path:
    gf = GradientFill(type="linear", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    return _build_gradient_fill_workbook(tmp_path, gf)


# ---- §15: positive tests ----


def test_gradient_simple_linear_two_rgb_stops(tmp_path):
    gf = GradientFill(type="linear", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_linear_with_degree(tmp_path):
    gf = GradientFill(type="linear", degree=45, stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_degree_scientific_notation_small(tmp_path):
    gf = GradientFill(type="linear", degree=1e-07, stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    raw = _read_part(result, "xl/styles.xml").decode()
    assert 'degree="1e-07"' in raw
    validate_scrubbed_workbook_package(result)


def test_gradient_degree_scientific_notation_large(tmp_path):
    gf = GradientFill(type="linear", degree=1e20, stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    raw = _read_part(result, "xl/styles.xml").decode()
    assert 'degree="1e+20"' in raw
    validate_scrubbed_workbook_package(result)


def test_gradient_negative_degree(tmp_path):
    gf = GradientFill(type="linear", degree=-45, stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_path_type(tmp_path):
    gf = GradientFill(type="path", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_path_with_edges(tmp_path):
    gf = GradientFill(
        type="path", left=0.1, right=0.2, top=0.3, bottom=0.4, stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00"))
    )
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_degree_plus_edges_simultaneously(tmp_path):
    """
    Frozen contract §6: type-dependent restrictions НЕ вводятся —
    реальный openpyxl output может содержать degree одновременно с
    left/right/top/bottom независимо от type, и это легитимно.
    """
    gf = GradientFill(
        type="linear", degree=45, left=0.1, right=0.9, stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00"))
    )
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_boundary_position_zero(tmp_path):
    gf = GradientFill(type="linear", stop=(Stop(Color(rgb="FF0000FF"), 0.0), Stop(Color(rgb="FF00FF00"), 0.0000001)))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    raw = _read_part(result, "xl/styles.xml").decode()
    assert 'position="0"' in raw
    validate_scrubbed_workbook_package(result)


def test_gradient_boundary_position_one(tmp_path):
    gf = GradientFill(type="linear", stop=(Stop(Color(rgb="FF0000FF"), 0.9999999), Stop(Color(rgb="FF00FF00"), 1.0)))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    raw = _read_part(result, "xl/styles.xml").decode()
    assert 'position="1"' in raw
    validate_scrubbed_workbook_package(result)


def test_gradient_theme_color(tmp_path):
    gf = GradientFill(type="linear", stop=(Stop(Color(theme=3), 0.0), Stop(Color(theme=4), 1.0)))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_indexed_color(tmp_path):
    gf = GradientFill(type="linear", stop=(Stop(Color(indexed=5), 0.0), Stop(Color(indexed=7), 1.0)))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_auto_color(tmp_path):
    gf = GradientFill(type="linear", stop=(Stop(Color(auto=True), 0.0), Stop(Color(rgb="FF00FF00"), 1.0)))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_color_with_tint(tmp_path):
    gf = GradientFill(type="linear", stop=(Stop(Color(theme=3, tint=0.5), 0.0), Stop(Color(rgb="FF00FF00"), 1.0)))
    result = _build_gradient_fill_workbook(tmp_path, gf)
    validate_scrubbed_workbook_package(result)


def test_gradient_zero_stops(tmp_path):
    gf = GradientFill(type="linear", stop=())
    result = _build_gradient_fill_workbook(tmp_path, gf)
    raw = _read_part(result, "xl/styles.xml").decode()
    assert "<gradientFill" in raw and "<stop" not in raw
    validate_scrubbed_workbook_package(result)


def test_gradient_multiple_stops(tmp_path):
    stops = [Stop(Color(rgb="FF000000"), i / 9) for i in range(10)]
    gf = GradientFill(type="linear", stop=stops)
    result = _build_gradient_fill_workbook(tmp_path, gf)
    r = validate_scrubbed_workbook_package(result)
    assert r is not None


def test_gradient_fillid_reference_on_cell(tmp_path):
    """
    MINOR-B Final Pre-Commit Pass: заменяет предыдущую misleading-версию
    (assertion на worksheet_count не имела отношения к fillId). Доказывает
    ПРЯМО: (1) genuine GradientFill — физически конкретный <fill> в
    <fills> с известным индексом; (2) <cellXfs> реально содержит
    fillId, ссылающийся ИМЕННО на этот индекс (не произвольный/дефолтный);
    (3) validator PASS именно с этой ссылкой.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = GradientFill(type="linear", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _save_and_scrub(tmp_path, wb)

    raw = _read_part(result, "xl/styles.xml").decode()
    fills_block = re.search(r"<fills count=\"(\d+)\">(.*?)</fills>", raw).group(0)
    fill_entries = re.findall(r"<fill>(.*?)</fill>", fills_block)
    gradient_indices = [i for i, entry in enumerate(fill_entries) if entry.startswith("<gradientFill")]
    assert len(gradient_indices) == 1, "ожидается ровно один <gradientFill> среди <fill>"
    gradient_fill_index = gradient_indices[0]

    cell_xfs_block = re.search(r"<cellXfs.*?</cellXfs>", raw).group(0)
    referencing_xfs = re.findall(r'fillId="(\d+)"', cell_xfs_block)
    assert str(gradient_fill_index) in referencing_xfs, (
        "ни один <xf> в <cellXfs> не ссылается на реальный индекс gradientFill-а"
    )

    r = validate_scrubbed_workbook_package(result)
    assert r.worksheet_count == 1


def test_gradient_fillid_upper_bound_rejected(tmp_path):
    """
    Тот же genuine GradientFill output, что и в предыдущем тесте —
    fillId, установленный РОВНО в actual_fill_count (за пределами
    валидного диапазона), должен быть отклонён.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = GradientFill(type="linear", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _save_and_scrub(tmp_path, wb)

    raw = _read_part(result, "xl/styles.xml").decode()
    actual_fill_count = int(re.search(r'<fills count="(\d+)">', raw).group(1))
    m_grad_xf = re.search(r'(<xf [^>]*fillId=")(\d+)("[^>]*xfId="0" />)', raw)
    assert m_grad_xf is not None

    def bump_to_upper_bound(d: bytes) -> bytes:
        old = m_grad_xf.group(0).encode()
        new = (m_grad_xf.group(1) + str(actual_fill_count) + m_grad_xf.group(3)).encode()
        return d.replace(old, new, 1)

    out = _mutate_one_part(tmp_path, result, "xl/styles.xml", bump_to_upper_bound, "gf_fillid_upper.xlsx")
    with zipfile.ZipFile(out) as z:
        mutated_raw = z.read("xl/styles.xml").decode()
    assert f'fillId="{actual_fill_count}"' in mutated_raw
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_fillid_negative_rejected(tmp_path):
    """Тот же genuine GradientFill output — fillId=-1 должен быть отклонён (MINOR-A)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = GradientFill(type="linear", stop=(Color(rgb="FF0000FF"), Color(rgb="FF00FF00")))
    result = _save_and_scrub(tmp_path, wb)

    raw = _read_part(result, "xl/styles.xml").decode()
    m_grad_xf = re.search(r'(<xf [^>]*fillId=")(\d+)("[^>]*xfId="0" />)', raw)
    assert m_grad_xf is not None

    def set_negative(d: bytes) -> bytes:
        old = m_grad_xf.group(0).encode()
        new = (m_grad_xf.group(1) + "-1" + m_grad_xf.group(3)).encode()
        return d.replace(old, new, 1)

    out = _mutate_one_part(tmp_path, result, "xl/styles.xml", set_negative, "gf_fillid_negative.xlsx")
    with zipfile.ZipFile(out) as z:
        mutated_raw = z.read("xl/styles.xml").decode()
    assert 'fillId="-1"' in mutated_raw
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- MINOR-A Final Pre-Commit Pass: negative style-index IDs ----


@pytest.mark.parametrize("attr", ["fontId", "fillId", "borderId", "numFmtId"])
def test_negative_style_index_rejected(tmp_path, richer_scrubbed, attr):
    """
    fontId/fillId/borderId/numFmtId в <xf> внутри <cellXfs> — _int_shape
    допускает знак "-", поэтому только upper-bound (">= count") НЕ
    отклоняет отрицательные значения (эмпирически подтверждено Focused
    Independent Security Re-Review). Мутирует РЕАЛЬНЫЙ <xf> внутри
    <cellXfs> (не <cellStyleXfs>, где bounds-check исторически не
    применяется вообще — отдельный, задокументированный, не относящийся
    к этому фиксу gap).
    """
    raw = _read_part(richer_scrubbed, "xl/styles.xml").decode()
    cell_xfs_block = re.search(r"<cellXfs.*?</cellXfs>", raw).group(0)
    assert f'{attr}="0"' in cell_xfs_block, f"фикстура должна содержать {attr}=\"0\" внутри <cellXfs>"

    def make_negative(d: bytes, a=attr) -> bytes:
        pattern = re.compile((a + r'="0"([^>]*xfId=)').encode())
        return pattern.sub((a + '="-1"\\1').encode(), d, count=1)

    out = _mutate_one_part(tmp_path, richer_scrubbed, "xl/styles.xml", make_negative, f"neg_style_{attr}.xlsx")
    with zipfile.ZipFile(out) as z:
        mutated_raw = z.read("xl/styles.xml").decode()
    assert f'{attr}="-1"' in mutated_raw, "мутация не применилась — тест был бы ложноположительным"
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_positive_style_index_zero_still_passes(tmp_path, richer_scrubbed):
    """
    Boundary-регрессия: fontId/fillId/borderId="0" (уже присутствует в
    реальном выводе фикстуры) обязан по-прежнему проходить после
    MINOR-A — фикс добавляет только "< 0", не трогая "== 0".
    """
    validate_scrubbed_workbook_package(richer_scrubbed)


def test_patternfill_regression_passes(tmp_path):
    """PatternFill grammar НЕ должна быть ослаблена добавлением gradientFill."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A1"].fill = PatternFill(patternType="solid", fgColor=Color(rgb="FF00FF00"))
    result = _save_and_scrub(tmp_path, wb)
    validate_scrubbed_workbook_package(result)


# ---- §16: negative tests (real scrub output + surgical mutation) ----


def _mutate_gradient(tmp_path: Path, gradient_scrubbed: Path, transform, out_name: str) -> Path:
    return _mutate_one_part(tmp_path, gradient_scrubbed, "xl/styles.xml", transform, out_name)


def test_gradient_unknown_attr_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<gradientFill type="linear">', b'<gradientFill type="linear" evil="1">'), "n1.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_unknown_child_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<gradientFill type="linear">', b'<gradientFill type="linear"><evil/>'), "n2.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_missing_type_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(tmp_path, gradient_scrubbed, lambda d: d.replace(b' type="linear"', b""), "n3.xlsx")
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_type_radial_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(tmp_path, gradient_scrubbed, lambda d: d.replace(b'type="linear"', b'type="radial"'), "n4.xlsx")
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


@pytest.mark.parametrize(
    "bad_degree",
    ["", "nan", "NaN", "inf", "-inf", "Infinity", "1e7", "+1", ".5", "1.", "1_000", " 1", "1 "],
)
def test_gradient_degree_invalid_lexical_form_rejected(tmp_path, gradient_scrubbed, bad_degree):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d, v=bad_degree: d.replace(b'<gradientFill type="linear">', f'<gradientFill type="linear" degree="{v}">'.encode()),
        f"n5_{abs(hash(bad_degree))}.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


@pytest.mark.parametrize("edge_attr", ["left", "right", "top", "bottom"])
def test_gradient_edge_invalid_lexical_form_rejected(tmp_path, gradient_scrubbed, edge_attr):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d, a=edge_attr: d.replace(
            b'<gradientFill type="linear">', f'<gradientFill type="linear" {a}="1_000">'.encode()
        ),
        f"n6_{edge_attr}.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_missing_position_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', b"<stop>", 1), "n7.xlsx")
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_position_negative_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', b'<stop position="-1">', 1), "n8.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_position_above_one_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', b'<stop position="2">', 1), "n9.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_position_nan_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', b'<stop position="nan">', 1), "n10.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_position_scientific_out_of_range_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', b'<stop position="1e7">', 1), "n11.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_missing_color_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b'<stop position="0"><color rgb="FF0000FF" /></stop>', b'<stop position="0"></stop>'),
        "n12.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_two_colors_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(
            b'<stop position="0"><color rgb="FF0000FF" /></stop>',
            b'<stop position="0"><color rgb="FF0000FF" /><color rgb="FF00FF00" /></stop>',
        ),
        "n13.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_unknown_stop_attr_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', b'<stop position="0" evil="1">', 1), "n14.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_unknown_stop_child_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(
            b'<stop position="0"><color rgb="FF0000FF" /></stop>',
            b'<stop position="0"><color rgb="FF0000FF" /><evil/></stop>',
        ),
        "n15.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_invalid_color_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<color rgb="FF0000FF" />', b'<color rgb="ZZZZZZ" />', 1), "n16.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_two_color_modes_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<color rgb="FF0000FF" />', b'<color rgb="FF0000FF" theme="1" />', 1), "n17.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_color_type_attr_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<color rgb="FF0000FF" />', b'<color rgb="FF0000FF" type="rgb" />', 1), "n18.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fill_pattern_plus_gradient_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b"<fill><gradientFill", b"<fill><patternFill/><gradientFill"), "n19.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fill_two_gradientfill_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b"</gradientFill></fill>", b'</gradientFill><gradientFill type="linear"/></fill>'),
        "n20.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fill_two_patternfill_rejected(tmp_path, richer_scrubbed):
    out = _mutate_one_part(
        tmp_path,
        richer_scrubbed,
        "xl/styles.xml",
        lambda d: d.replace(b"<fill><patternFill /></fill>", b"<fill><patternFill /><patternFill /></fill>", 1),
        "n21.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_fill_zero_children_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(
            b'<fill><gradientFill type="linear"><stop position="0"><color rgb="FF0000FF" /></stop>'
            b'<stop position="1"><color rgb="FF00FF00" /></stop></gradientFill></fill>',
            b"<fill></fill>",
        ),
        "n22.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_non_whitespace_text_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b'<gradientFill type="linear">', f'<gradientFill type="linear">{SECRET}'.encode()),
        "n23.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_tail_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b"</gradientFill>", f"</gradientFill>{SECRET}".encode()), "n24.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_text_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path, gradient_scrubbed, lambda d: d.replace(b'<stop position="0">', f'<stop position="0">{SECRET}'.encode()), "n25.xlsx"
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_stop_tail_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(tmp_path, gradient_scrubbed, lambda d: d.replace(b"</stop>", f"</stop>{SECRET}".encode(), 1), "n26.xlsx")
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_color_text_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b'<color rgb="FF0000FF" />', f'<color rgb="FF0000FF">{SECRET}</color>'.encode(), 1),
        "n27.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_comment_inside_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b'<gradientFill type="linear">', f'<gradientFill type="linear"><!--{SECRET}-->'.encode()),
        "n28.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


def test_gradient_pi_inside_rejected(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b'<gradientFill type="linear">', f'<gradientFill type="linear"><?evil {SECRET}?>'.encode()),
        "n29.xlsx",
    )
    _assert_rejected(out, PackageScrubReason.POST_VALIDATION_FAILED)


# ---- §18: error privacy for representative GradientFill failures ----


def test_gradient_error_privacy(tmp_path, gradient_scrubbed):
    out = _mutate_gradient(
        tmp_path,
        gradient_scrubbed,
        lambda d: d.replace(b'<gradientFill type="linear">', f'<gradientFill type="linear" evil="{SECRET}">'.encode()),
        "privacy1.xlsx",
    )
    with pytest.raises(PackageScrubError) as excinfo:
        validate_scrubbed_workbook_package(out)
    assert excinfo.value.reason == PackageScrubReason.POST_VALIDATION_FAILED
    assert SECRET not in str(excinfo.value)
    assert SECRET not in repr(excinfo.value)
    assert excinfo.value.__context__ is None
    assert excinfo.value.__cause__ is None
