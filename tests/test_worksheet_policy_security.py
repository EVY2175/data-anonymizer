"""
Тесты Stage 10C.1 — OD-7: ни одно исключение/результат/evidence не должно
раскрывать пароль, raw company/person/identifier значение, имя листа или
column_name/заголовок.

Проверяются: str(exc), repr(exc), exc.args, полный traceback-текст,
exc.__context__/exc.__cause__ (в частности — что санитизированные
CoverageError(WORKBOOK_MUTATION_FAILED) НЕ несут исходное, потенциально
"утечное" исключение нижнего слоя как __context__), а также repr()
результатов/evidence моделей.
"""

from __future__ import annotations

import traceback
from pathlib import Path

import openpyxl
import pytest

from app.coverage.errors import ColumnSafetyError, ColumnSafetyReason, CoverageError, CoverageReason
from app.coverage.models import WorksheetPolicy
from app.coverage.worksheet_policy import prepare_workbook_for_anonymization, scan_process_worksheet
from app.models.rules import Action, FieldRule, FieldType

PASSWORD_SENTINEL = "SuperSecretPassword123!"
COMPANY_SENTINEL = "ООО СекретнаяКомпания"
PERSON_SENTINEL = "Иванов Секретович"
RAW_INN_SENTINEL = "7707083893"  # валидный ИНН (checksum), тот же вектор что и в tests/test_detectors.py
WORKSHEET_NAME_SENTINEL = "СекретныйЛист"
HEADER_SENTINEL = "СекретныйЗаголовок"

_ALL_SENTINELS = [
    PASSWORD_SENTINEL,
    COMPANY_SENTINEL,
    PERSON_SENTINEL,
    RAW_INN_SENTINEL,
    WORKSHEET_NAME_SENTINEL,
    HEADER_SENTINEL,
]


def _assert_no_sentinel_leak(exc: BaseException, sentinels: list[str] = _ALL_SENTINELS) -> None:
    text_blobs = [
        str(exc),
        repr(exc),
        repr(exc.args),
        "".join(traceback.format_exception(type(exc), exc, exc.__traceback__)),
    ]
    if exc.__context__ is not None:
        text_blobs.append(str(exc.__context__))
        text_blobs.append(repr(exc.__context__))
    if exc.__cause__ is not None:
        text_blobs.append(str(exc.__cause__))
        text_blobs.append(repr(exc.__cause__))

    for sentinel in sentinels:
        for blob in text_blobs:
            assert sentinel not in blob, f"sentinel {sentinel!r} leaked into {blob!r}"


@pytest.fixture()
def source_path(tmp_path: Path) -> Path:
    return tmp_path / "source.xlsx"


# ---------------------------------------------------------------------------
# A. CoverageError — MISSING_WORKSHEET_POLICY (имя листа = sentinel)
# ---------------------------------------------------------------------------


def test_missing_worksheet_policy_does_not_leak_worksheet_name(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.create_sheet(WORKSHEET_NAME_SENTINEL)
    wb.save(source_path)
    wb.close()

    rules = {"Data": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is CoverageReason.MISSING_WORKSHEET_POLICY
    _assert_no_sentinel_leak(exc_info.value)


# ---------------------------------------------------------------------------
# B. CoverageError — HEADER_MISMATCH (заголовок = sentinel)
# ---------------------------------------------------------------------------


def test_header_mismatch_does_not_leak_header_text(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = HEADER_SENTINEL
    ws["A2"] = "value"
    wb.save(source_path)
    wb.close()

    rules = {"Data": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is CoverageReason.HEADER_MISMATCH
    _assert_no_sentinel_leak(exc_info.value)


# ---------------------------------------------------------------------------
# C. ColumnSafetyError — KEEP_FORBIDDEN_FOR_TYPE (company/person sentinel)
# ---------------------------------------------------------------------------


def test_keep_forbidden_does_not_leak_company_value(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Company"
    ws["A2"] = COMPANY_SENTINEL
    wb.save(source_path)
    wb.close()

    rules = {"Data": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.KEEP)}}
    with pytest.raises(ColumnSafetyError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is ColumnSafetyReason.KEEP_FORBIDDEN_FOR_TYPE
    _assert_no_sentinel_leak(exc_info.value)


def test_keep_forbidden_person_does_not_leak_person_value() -> None:
    with pytest.raises(ColumnSafetyError) as exc_info:
        from app.coverage.worksheet_policy import validate_column_action

        validate_column_action(
            FieldRule(column_name=PERSON_SENTINEL, field_type=FieldType.PERSON, action=Action.KEEP)
        )
    _assert_no_sentinel_leak(exc_info.value)


# ---------------------------------------------------------------------------
# D. ColumnSafetyError — CROSS_TYPE_IDENTIFIER_DETECTED (raw INN sentinel)
# ---------------------------------------------------------------------------


def test_cross_type_detected_does_not_leak_raw_identifier(source_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Financial"
    ws["A2"] = RAW_INN_SENTINEL
    wb.save(source_path)
    wb.close()

    rules = {"Data": {1: FieldRule(column_name="Financial", field_type=FieldType.FINANCIAL, action=Action.KEEP)}}
    with pytest.raises(ColumnSafetyError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass
    assert exc_info.value.reason is ColumnSafetyReason.CROSS_TYPE_IDENTIFIER_DETECTED
    _assert_no_sentinel_leak(exc_info.value)


def test_scan_process_worksheet_result_repr_does_not_leak_raw_identifier() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Financial"
    ws["A2"] = RAW_INN_SENTINEL
    rule = FieldRule(column_name="Financial", field_type=FieldType.FINANCIAL, action=Action.KEEP)
    result = scan_process_worksheet(ws, 1, {1: rule})
    assert RAW_INN_SENTINEL not in repr(result)


# ---------------------------------------------------------------------------
# E. CoverageError(WORKBOOK_MUTATION_FAILED) — санитизация нижнего слоя
# ---------------------------------------------------------------------------


def test_workbook_mutation_failed_does_not_chain_leaky_underlying_exception(
    source_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Company"
    ws["A2"] = "value"
    wb.save(source_path)
    wb.close()

    def _leaky_save(self, filename):  # noqa: ANN001
        raise RuntimeError(f"leaky failure containing password {PASSWORD_SENTINEL}")

    monkeypatch.setattr(openpyxl.Workbook, "save", _leaky_save)

    rules = {"Data": {1: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)}}
    with pytest.raises(CoverageError) as exc_info:
        with prepare_workbook_for_anonymization(source_path, {"Data": WorksheetPolicy.PROCESS}, rules):
            pass

    exc = exc_info.value
    assert exc.reason is CoverageReason.WORKBOOK_MUTATION_FAILED
    # Ключевая проверка OD-7: __context__ НЕ указывает на перехваченный
    # "утечный" RuntimeError — except-блок был покинут до подъёма нового
    # исключения (тот же приём, что и в app.workspace.workspace).
    assert exc.__context__ is None
    assert exc.__cause__ is None
    _assert_no_sentinel_leak(exc)


# ---------------------------------------------------------------------------
# F. Sentinel-скан по фиксированным сообщениям всех reason (defense-in-depth)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("reason", list(CoverageReason))
def test_coverage_error_fixed_messages_never_contain_sentinels(reason: CoverageReason) -> None:
    exc = CoverageError(reason)
    _assert_no_sentinel_leak(exc)


@pytest.mark.parametrize("reason", list(ColumnSafetyReason))
def test_column_safety_error_fixed_messages_never_contain_sentinels(reason: ColumnSafetyReason) -> None:
    exc = ColumnSafetyError(reason)
    _assert_no_sentinel_leak(exc)
