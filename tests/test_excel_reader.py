"""
Тесты app.excel.reader (list_sheet_names, read_flat_table) и
app.excel.models (CellRecord, FlatTable).

Все Excel-фикстуры создаются исключительно через openpyxl в tmp_path —
никаких реальных пользовательских .xlsx файлов в репозитории.

Некоторые тесты (см. комментарии) построены на ФАКТИЧЕСКИ исследованном
поведении openpyxl (а не на предположениях): например, "круглый" float
(7707083893.0) после реального save/reload сериализуется openpyxl без
дробной части и на чтении сам возвращается как int, а не float — это
поведение openpyxl, а не reader'а. Гарантия Stage 6 — reader не меняет
то, что вернул openpyxl, а НЕ то, что openpyxl всегда сохраняет исходный
Python-тип.
"""

from __future__ import annotations

import dataclasses
import datetime
from pathlib import Path

import openpyxl
import pytest

from app.excel.models import CellRecord, FlatTable
from app.excel.reader import _cell_record, list_sheet_names, read_flat_table


# ---------------------------------------------------------------------------
# 1-6. Sheet selection
# ---------------------------------------------------------------------------


def test_single_worksheet(tmp_path: Path) -> None:
    path = tmp_path / "single.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws["A1"] = "Колонка"
    ws["A2"] = "значение"
    wb.save(path)

    table = read_flat_table(path)
    assert table.sheet_name == "Данные"
    assert table.header_row == (CellRecord(row=1, column=1, value="Колонка"),)
    assert table.rows == ((CellRecord(row=2, column=1, value="значение"),),)


def test_multiple_worksheets(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "f1"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "s1"
    wb.save(path)

    assert list_sheet_names(path) == ("First", "Second")
    assert read_flat_table(path, sheet_name="First").header_row[0].value == "f1"
    assert read_flat_table(path, sheet_name="Second").header_row[0].value == "s1"


def test_list_sheet_names_preserves_order(tmp_path: Path) -> None:
    path = tmp_path / "order.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zebra"
    wb.create_sheet("Alpha")
    wb.create_sheet("Mango")
    wb.save(path)

    assert list_sheet_names(path) == ("Zebra", "Alpha", "Mango")


def test_sheet_name_none_reads_active_worksheet(tmp_path: Path) -> None:
    path = tmp_path / "active.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "not-active"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "is-active"
    wb.active = wb.sheetnames.index("Second")
    wb.save(path)

    table = read_flat_table(path)
    assert table.sheet_name == "Second"
    assert table.header_row[0].value == "is-active"


def test_explicit_sheet_name(tmp_path: Path) -> None:
    path = tmp_path / "explicit.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1["A1"] = "first-value"
    ws2 = wb.create_sheet("Second")
    ws2["A1"] = "second-value"
    wb.save(path)

    table = read_flat_table(path, sheet_name="First")
    assert table.sheet_name == "First"
    assert table.header_row[0].value == "first-value"


def test_unknown_sheet_name_raises(tmp_path: Path) -> None:
    path = tmp_path / "unknown_sheet.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "OnlySheet"
    wb.save(path)

    with pytest.raises(KeyError):
        read_flat_table(path, sheet_name="DoesNotExist")


def test_hidden_worksheet_can_be_read_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "hidden.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Visible"
    hidden_ws = wb.create_sheet("Hidden")
    hidden_ws.sheet_state = "hidden"
    hidden_ws["A1"] = "hidden-data"
    wb.save(path)

    assert "Hidden" in list_sheet_names(path)
    table = read_flat_table(path, sheet_name="Hidden")
    assert table.header_row[0].value == "hidden-data"


# ---------------------------------------------------------------------------
# 7-8. header_row
# ---------------------------------------------------------------------------


def test_custom_header_row(tmp_path: Path) -> None:
    path = tmp_path / "custom_header.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "преамбула"
    ws["A2"] = "ещё преамбула"
    ws["A3"] = "ИНН"
    ws["A4"] = "7707083893"
    ws["A5"] = "500100732259"
    wb.save(path)

    table = read_flat_table(path, header_row=3)
    assert table.header_row == (CellRecord(row=3, column=1, value="ИНН"),)
    assert table.rows == (
        (CellRecord(row=4, column=1, value="7707083893"),),
        (CellRecord(row=5, column=1, value="500100732259"),),
    )


def test_header_row_at_last_row_gives_no_data_rows(tmp_path: Path) -> None:
    path = tmp_path / "header_only.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "только заголовок"
    wb.save(path)

    table = read_flat_table(path, header_row=1)
    assert table.rows == ()


@pytest.mark.parametrize("bad_value", [0, -1, -100])
def test_invalid_header_row_value_raises_value_error(tmp_path: Path, bad_value: int) -> None:
    path = tmp_path / "any.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(ValueError):
        read_flat_table(path, header_row=bad_value)


@pytest.mark.parametrize("bad_type", ["1", 1.0, True, None, [1]])
def test_invalid_header_row_type_raises_type_error(tmp_path: Path, bad_type: object) -> None:
    path = tmp_path / "any2.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(TypeError):
        read_flat_table(path, header_row=bad_type)  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# header_row не должен превышать worksheet.max_row (fail-fast, без
# создания искусственной строки заголовков "из воздуха").
# ---------------------------------------------------------------------------


def _make_five_row_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "five_rows.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 6):
        ws.cell(row=r, column=1, value=f"row{r}")
    wb.save(path)
    assert ws.max_row == 5  # sanity-проверка самой фикстуры
    return path


def test_header_row_one_past_max_row_raises_value_error(tmp_path: Path) -> None:
    path = _make_five_row_workbook(tmp_path)
    with pytest.raises(ValueError, match=r"header_row=6.*max_row=5"):
        read_flat_table(path, header_row=6)


def test_header_row_far_past_max_row_raises_value_error(tmp_path: Path) -> None:
    path = _make_five_row_workbook(tmp_path)
    with pytest.raises(ValueError, match=r"header_row=100.*max_row=5"):
        read_flat_table(path, header_row=100)


def test_header_row_equal_to_max_row_is_allowed(tmp_path: Path) -> None:
    path = _make_five_row_workbook(tmp_path)
    table = read_flat_table(path, header_row=5)
    assert table.header_row == (CellRecord(row=5, column=1, value="row5"),)
    assert table.rows == ()


def test_header_row_one_on_brand_new_empty_worksheet_is_allowed(tmp_path: Path) -> None:
    """
    Пустой новый worksheet openpyxl имеет max_row == 1 — header_row=1
    должен оставаться допустимым и не расцениваться как "превышение".
    """
    path = tmp_path / "brand_new_empty.xlsx"
    wb = openpyxl.Workbook()
    assert wb.active.max_row == 1  # sanity-проверка самой фикстуры
    wb.save(path)

    table = read_flat_table(path, header_row=1)
    assert table.header_row == (CellRecord(row=1, column=1, value=None),)
    assert table.rows == ()


def test_workbook_is_closed_even_when_header_row_exceeds_max_row(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """
    Белый ящик: подтверждает, что workbook.close() вызывается через
    существующий try/finally даже тогда, когда header_row > max_row
    поднимает ValueError внутри try-блока.
    """
    path = _make_five_row_workbook(tmp_path)

    real_load_workbook = openpyxl.load_workbook
    close_calls: list[bool] = []

    def spy_load_workbook(*args: object, **kwargs: object):
        workbook = real_load_workbook(*args, **kwargs)
        real_close = workbook.close

        def spy_close() -> None:
            close_calls.append(True)
            real_close()

        workbook.close = spy_close
        return workbook

    monkeypatch.setattr(openpyxl, "load_workbook", spy_load_workbook)

    with pytest.raises(ValueError):
        read_flat_table(path, header_row=100)

    assert close_calls == [True]


# ---------------------------------------------------------------------------
# 9-12. Обычные строки, пустые ячейки/строки, координаты
# ---------------------------------------------------------------------------


def test_ordinary_rows(tmp_path: Path) -> None:
    path = tmp_path / "ordinary.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Компания", "ИНН"
    ws["A2"], ws["B2"] = "ООО Ромашка", "7707083893"
    ws["A3"], ws["B3"] = "ООО Лютик", "500100732259"
    wb.save(path)

    table = read_flat_table(path)
    assert len(table.rows) == 2
    assert table.rows[0][0].value == "ООО Ромашка"
    assert table.rows[1][1].value == "500100732259"


def test_blank_cells_inside_populated_row(tmp_path: Path) -> None:
    path = tmp_path / "blank_cells.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "A", "B", "C"
    ws["A2"] = "value"
    # B2 намеренно не заполняется
    ws["C2"] = "another"
    wb.save(path)

    table = read_flat_table(path)
    row = table.rows[0]
    assert row[0].value == "value"
    assert row[1].value is None
    assert row[2].value == "another"


def test_completely_empty_row_between_data_is_preserved(tmp_path: Path) -> None:
    path = tmp_path / "empty_row.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Header"
    ws["A2"] = "row2"
    # строка 3 полностью пуста
    ws["A4"] = "row4"
    wb.save(path)

    table = read_flat_table(path)
    assert len(table.rows) == 3  # строки 2, 3, 4 — без сжатия
    assert table.rows[0] == (CellRecord(row=2, column=1, value="row2"),)
    assert table.rows[1] == (CellRecord(row=3, column=1, value=None),)
    assert table.rows[2] == (CellRecord(row=4, column=1, value="row4"),)


def test_cell_coordinates_are_1_based_and_match_source(tmp_path: Path) -> None:
    path = tmp_path / "coords.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="H1")
    ws.cell(row=1, column=2, value="H2")
    ws.cell(row=2, column=1, value="R2C1")
    ws.cell(row=2, column=2, value="R2C2")
    wb.save(path)

    table = read_flat_table(path)
    h1, h2 = table.header_row
    assert (h1.row, h1.column) == (1, 1)
    assert (h2.row, h2.column) == (1, 2)
    r2c1, r2c2 = table.rows[0]
    assert (r2c1.row, r2c1.column) == (2, 1)
    assert (r2c2.row, r2c2.column) == (2, 2)


# ---------------------------------------------------------------------------
# 13-15. Заголовки: дубликаты, пустой, числовой
# ---------------------------------------------------------------------------


def test_duplicate_headers_are_preserved_positionally(tmp_path: Path) -> None:
    path = tmp_path / "dup_headers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "ИНН", "ИНН"
    ws["A2"], ws["B2"] = "value-a", "value-b"
    wb.save(path)

    table = read_flat_table(path)
    assert table.header_row[0].value == "ИНН"
    assert table.header_row[1].value == "ИНН"
    assert table.header_row[0].column == 1
    assert table.header_row[1].column == 2
    # Обе колонки данных сохранены отдельно, несмотря на одинаковый заголовок.
    assert table.rows[0][0].value == "value-a"
    assert table.rows[0][1].value == "value-b"


def test_blank_header_is_preserved_as_none(tmp_path: Path) -> None:
    path = tmp_path / "blank_header.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Есть заголовок"
    # B1 намеренно не заполняется
    ws["C1"] = "Ещё заголовок"
    wb.save(path)

    table = read_flat_table(path)
    assert table.header_row[1].value is None


def test_numeric_header_is_preserved_as_number(tmp_path: Path) -> None:
    path = tmp_path / "numeric_header.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 2024
    wb.save(path)

    table = read_flat_table(path)
    assert table.header_row[0].value == 2024
    assert isinstance(table.header_row[0].value, int)


# ---------------------------------------------------------------------------
# 16-22. Типы значений
# ---------------------------------------------------------------------------


def test_value_type_str(tmp_path: Path) -> None:
    path = tmp_path / "t_str.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A2"] = "текст"
    wb.save(path)
    value = read_flat_table(path).rows[0][0].value
    assert value == "текст"
    assert isinstance(value, str)


def test_value_type_int(tmp_path: Path) -> None:
    path = tmp_path / "t_int.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A2"] = 42
    wb.save(path)
    value = read_flat_table(path).rows[0][0].value
    assert value == 42
    assert type(value) is int


def test_value_type_float(tmp_path: Path) -> None:
    path = tmp_path / "t_float.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A2"] = 1234567.5  # дробная часть — гарантированно остаётся float
    wb.save(path)
    value = read_flat_table(path).rows[0][0].value
    assert value == 1234567.5
    assert type(value) is float


def test_value_type_bool(tmp_path: Path) -> None:
    path = tmp_path / "t_bool.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A2"] = True
    wb.active["A3"] = False
    wb.save(path)
    table = read_flat_table(path)
    assert table.rows[0][0].value is True
    assert table.rows[1][0].value is False


def test_value_type_none(tmp_path: Path) -> None:
    path = tmp_path / "t_none.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "H1", "H2"
    ws["A2"] = None
    # B2 закрепляет worksheet.max_row=2 — иначе строка с одним только None
    # в A2 не расширяет dimension листа (openpyxl не считает её "занятой").
    ws["B2"] = "anchor"
    wb.save(path)
    value = read_flat_table(path).rows[0][0].value
    assert value is None


def test_value_type_datetime(tmp_path: Path) -> None:
    path = tmp_path / "t_datetime.xlsx"
    wb = openpyxl.Workbook()
    dt = datetime.datetime(2024, 5, 1, 13, 30, 0)
    wb.active["A2"] = dt
    wb.save(path)
    value = read_flat_table(path).rows[0][0].value
    assert value == dt
    assert isinstance(value, datetime.datetime)


def test_value_type_time(tmp_path: Path) -> None:
    """
    datetime.time надёжно проходит реальный save/reload как time (в отличие
    от datetime.date — см. test_value_type_date_roundtrips_as_datetime).
    """
    path = tmp_path / "t_time.xlsx"
    wb = openpyxl.Workbook()
    t = datetime.time(13, 30, 0)
    wb.active["A2"] = t
    wb.save(path)
    value = read_flat_table(path).rows[0][0].value
    assert value == t
    assert isinstance(value, datetime.time)


def test_value_type_date_roundtrips_as_datetime(tmp_path: Path) -> None:
    """
    Фактически исследованное поведение openpyxl (не предположение): при
    реальном save/reload значение, записанное как datetime.date, openpyxl
    сам возвращает как datetime.datetime с временем 00:00, а не как
    datetime.date. Это поведение openpyxl, а не reader'а — reader здесь
    лишь обязан не исказить то, что вернул openpyxl (что и проверяется).
    """
    path = tmp_path / "t_date.xlsx"
    wb = openpyxl.Workbook()
    d = datetime.date(2024, 5, 1)
    wb.active["A2"] = d
    wb.save(path)

    raw_wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    raw_value = raw_wb.active["A2"].value
    raw_wb.close()

    value = read_flat_table(path).rows[0][0].value
    assert value == raw_value
    assert type(value) is type(raw_value)


def test_cell_record_helper_preserves_date_value_untouched() -> None:
    """
    Белый ящик: проверяет саму функцию преобразования Cell -> CellRecord
    напрямую на синтетическом объекте с value=datetime.date(...), в обход
    ограничения реального openpyxl round-trip (см. тест выше) — доказывает,
    что _cell_record() не содержит никакой ветки, которая бы что-то делала
    с датой/типом value, если бы openpyxl когда-либо вернул чистый date.
    """

    class _FakeCell:
        row = 5
        column = 2
        value = datetime.date(2030, 1, 1)

    record = _cell_record(_FakeCell())
    assert record == CellRecord(row=5, column=2, value=datetime.date(2030, 1, 1))
    assert isinstance(record.value, datetime.date)


def test_formula_is_preserved_as_formula_string(tmp_path: Path) -> None:
    path = tmp_path / "formula.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A2"], ws["A3"] = 1, 2
    ws["A4"] = "=SUM(A2:A3)"
    wb.save(path)

    value = read_flat_table(path).rows[2][0].value
    assert value == "=SUM(A2:A3)"
    assert isinstance(value, str)


# ---------------------------------------------------------------------------
# 23. Merged cells
# ---------------------------------------------------------------------------


def test_merged_cells_are_not_expanded(tmp_path: Path) -> None:
    path = tmp_path / "merged.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "H1"
    ws["B1"] = "H2"
    ws["A2"] = "merged value"
    ws.merge_cells("A2:B2")
    wb.save(path)

    table = read_flat_table(path)
    row = table.rows[0]
    assert row[0].value == "merged value"  # top-left
    assert row[1].value is None  # остальная часть merge — родное поведение openpyxl


# ---------------------------------------------------------------------------
# 24-27. Числовые идентификаторы: float/int/text/leading zero — без коэрсии
# ---------------------------------------------------------------------------


def test_stage5_safety_contract_reader_never_coerces_numeric_types(tmp_path: Path) -> None:
    """
    Особенно важный regression test (см. отчёт по Stage 6).

    Гарантия: reader НЕ меняет тип/значение, которое вернул openpyxl —
    для каждой ячейки сравниваем результат read_flat_table() с прямым
    независимым чтением того же файла через openpyxl (теми же
    параметрами load_workbook).

    НЕ гарантия: то, что openpyxl всегда сохраняет исходный Python
    numeric type. Фактически исследовано (см. отчёт), что "круглый" float
    7707083893.0 при реальном save/reload сам openpyxl возвращает как int
    — это не подделывается и не выдаётся здесь за требование к openpyxl.
    """
    path = tmp_path / "inn_types.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "ИНН"
    ws["A2"] = "7707083893"
    ws["A2"].number_format = "@"  # явно текстовый формат
    ws["A3"] = 7707083893
    ws["A4"] = 7707083893.0
    wb.save(path)

    raw_wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    raw_ws = raw_wb["Sheet1"]
    raw_values = [raw_ws.cell(row=r, column=1).value for r in (2, 3, 4)]
    raw_wb.close()

    table = read_flat_table(path, sheet_name="Sheet1")
    read_values = [table.rows[i][0].value for i in range(3)]

    for raw, read in zip(raw_values, read_values):
        assert read == raw
        assert type(read) is type(raw)

    # Фиксируем сами фактически наблюдаемые значения (для прозрачности
    # отчёта, не как отдельное требование к openpyxl):
    assert raw_values[0] == "7707083893" and isinstance(raw_values[0], str)
    assert raw_values[1] == 7707083893 and type(raw_values[1]) is int


def test_float_value_is_not_reconstructed_into_str_or_int(tmp_path: Path) -> None:
    path = tmp_path / "float_not_coerced.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A2"] = 7707083893.5  # дробная часть -> гарантированно float
    wb.save(path)

    value = read_flat_table(path).rows[0][0].value
    assert type(value) is float
    assert value != "7707083893"
    assert value != 7707083893


def test_integer_identifier_is_not_reconstructed_into_str(tmp_path: Path) -> None:
    path = tmp_path / "int_not_coerced.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A2"] = 7707083893
    wb.save(path)

    value = read_flat_table(path).rows[0][0].value
    assert type(value) is int
    assert value != "7707083893"


def test_numeric_identifier_stored_as_text_stays_str(tmp_path: Path) -> None:
    path = tmp_path / "text_identifier.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A2"] = "7707083893"
    ws["A2"].number_format = "@"
    wb.save(path)

    value = read_flat_table(path).rows[0][0].value
    assert value == "7707083893"
    assert isinstance(value, str)


def test_leading_zero_identifier_stored_as_text_keeps_zero(tmp_path: Path) -> None:
    path = tmp_path / "leading_zero.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A2"] = "0087654321"
    ws["A2"].number_format = "@"
    wb.save(path)

    value = read_flat_table(path).rows[0][0].value
    assert value == "0087654321"
    assert isinstance(value, str)
    assert value.startswith("00")


# ---------------------------------------------------------------------------
# 28-29. Read-only контракт: исходный файл не меняется
# ---------------------------------------------------------------------------


def test_read_flat_table_does_not_modify_source_file(tmp_path: Path) -> None:
    path = tmp_path / "readonly1.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "data"
    wb.save(path)

    before = path.read_bytes()
    read_flat_table(path)
    after = path.read_bytes()

    assert after == before


def test_list_sheet_names_does_not_modify_source_file(tmp_path: Path) -> None:
    path = tmp_path / "readonly2.xlsx"
    wb = openpyxl.Workbook()
    wb.create_sheet("Second")
    wb.save(path)

    before = path.read_bytes()
    list_sheet_names(path)
    after = path.read_bytes()

    assert after == before


# ---------------------------------------------------------------------------
# Дополнительно: линейность (не строгий бенчмарк, а sanity-проверка того,
# что большая таблица читается целиком и без ошибок за один проход).
# ---------------------------------------------------------------------------


def test_reasonably_large_table_reads_completely(tmp_path: Path) -> None:
    path = tmp_path / "large.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    n_rows, n_cols = 500, 5
    for c in range(1, n_cols + 1):
        ws.cell(row=1, column=c, value=f"col{c}")
    for r in range(2, n_rows + 2):
        for c in range(1, n_cols + 1):
            ws.cell(row=r, column=c, value=r * 100 + c)
    wb.save(path)

    table = read_flat_table(path)
    assert len(table.header_row) == n_cols
    assert len(table.rows) == n_rows
    assert table.rows[-1][-1].value == (n_rows + 1) * 100 + n_cols


# ---------------------------------------------------------------------------
# Модели: базовая валидация координат CellRecord
# ---------------------------------------------------------------------------


def test_cell_record_rejects_non_positive_row() -> None:
    with pytest.raises(ValueError):
        CellRecord(row=0, column=1, value="x")


def test_cell_record_rejects_non_positive_column() -> None:
    with pytest.raises(ValueError):
        CellRecord(row=1, column=0, value="x")


def test_cell_record_rejects_bool_row_or_column() -> None:
    with pytest.raises(TypeError):
        CellRecord(row=True, column=1, value="x")  # type: ignore[arg-type]
    with pytest.raises(TypeError):
        CellRecord(row=1, column=True, value="x")  # type: ignore[arg-type]


def test_flat_table_is_frozen() -> None:
    table = FlatTable(sheet_name="S", header_row=(), rows=())
    with pytest.raises(dataclasses.FrozenInstanceError):
        table.sheet_name = "changed"  # type: ignore[misc]


def test_cell_record_is_frozen() -> None:
    record = CellRecord(row=1, column=1, value="x")
    with pytest.raises(dataclasses.FrozenInstanceError):
        record.value = "changed"  # type: ignore[misc]
