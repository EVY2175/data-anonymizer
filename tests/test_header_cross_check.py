"""
Тесты Stage 10C.1 — header cross-check
(app.coverage.worksheet_policy.scan_process_worksheet, header_mismatches).

Header — всегда row 1 (та же жёсткая конвенция, что и во всей Stage 8
оркестрации). Сравнение точное: без strip/casefold/normalize.
"""

from __future__ import annotations

import openpyxl

from app.coverage.worksheet_policy import scan_process_worksheet
from app.models.rules import Action, FieldRule, FieldType


def _rule(column_name: str) -> FieldRule:
    return FieldRule(column_name=column_name, field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE)


def _worksheet(headers: dict[int, object]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for column, value in headers.items():
        ws.cell(row=1, column=column, value=value)
        ws.cell(row=2, column=column, value="data")
    return ws


def test_exact_match_passes() -> None:
    ws = _worksheet({1: "Company"})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == ()


def test_mismatch_detected() -> None:
    ws = _worksheet({1: "Wrong"})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == (1,)


def test_blank_header_cell_mismatches() -> None:
    ws = _worksheet({1: None})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == (1,)


def test_missing_header_cell_entirely_absent_mismatches() -> None:
    # Ячейка (1, column) никогда не была установлена вообще — worksheet._cells
    # не содержит для неё entry. Header cross-check обязан обнаружить это
    # явным lookup, а не полагаться на то, что ячейка "случайно" попадёт в
    # общий проход по существующим ячейкам.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.cell(row=2, column=1, value="data-only-no-header")
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == (1,)


def test_duplicate_header_text_on_different_columns_both_checked_independently() -> None:
    ws = _worksheet({1: "ИНН", 2: "ИНН"})
    result = scan_process_worksheet(
        ws,
        1,
        {
            1: _rule("ИНН"),
            2: _rule("ИНН"),
        },
    )
    assert result.header_mismatches == ()


def test_extra_column_without_rule_not_checked() -> None:
    ws = _worksheet({1: "Company", 2: "Untouched"})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == ()


def test_reordered_columns_detected_as_mismatch() -> None:
    # Правило колонки 1 ожидает "Email", но физически в колонке 1 теперь
    # "Company" (реструктурирование источника) — биндинг по индексу
    # немедленно выявляет несоответствие.
    ws = _worksheet({1: "Company", 2: "Email"})
    rules = {
        1: FieldRule(column_name="Email", field_type=FieldType.EMAIL, action=Action.REMOVE),
        2: FieldRule(column_name="Company", field_type=FieldType.COMPANY, action=Action.PSEUDONYMIZE),
    }
    result = scan_process_worksheet(ws, 1, rules)
    assert set(result.header_mismatches) == {1, 2}


def test_case_difference_mismatches() -> None:
    ws = _worksheet({1: "company"})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == (1,)


def test_whitespace_difference_mismatches() -> None:
    ws = _worksheet({1: " Company "})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company")})
    assert result.header_mismatches == (1,)


def test_multiple_mismatches_all_reported() -> None:
    ws = _worksheet({1: "Wrong1", 2: "Wrong2"})
    result = scan_process_worksheet(ws, 1, {1: _rule("Company"), 2: _rule("Email")})
    assert result.header_mismatches == (1, 2)
