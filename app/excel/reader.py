"""
Read-only чтение обычной плоской Excel-таблицы через openpyxl.

======================================================================
Границы ответственности Stage 6
======================================================================

Этот модуль делает РОВНО одно преобразование:

    XLSX -> openpyxl -> FlatTable

Он НЕ анонимизирует, НЕ определяет типы полей/значений, НЕ строит
иерархию, НЕ пишет файлы. В частности, этот модуль сознательно НЕ
импортирует app.detectors и не оперирует FieldType/EntityType/Action/
FieldRule — reader только читает Excel, классификация значений/колонок
(если понадобится) — забота отдельного, более позднего слоя.

======================================================================
load_workbook(..., data_only=False, read_only=False)
======================================================================

data_only=False: мы хотим видеть исходные формулы (например,
"=SUM(A2:A5)"), а не зависеть от закэшированного Excel'ем результата
вычисления (который к тому же может отсутствовать, если файл не
пересчитывался/не сохранялся в реальном Excel). Вычисление формул,
получение cached result и вообще какая-либо логика вокруг формул — не
Stage 6.

read_only=False: обычный (не потоковый) режим — используется, чтобы
получить доступ к worksheet.max_row/max_column и обычному API Cell/
merged_cells без ограничений read-only-режима openpyxl. Никакой мутации
воркбука в этом режиме не производится — модуль ни разу не присваивает
cell.value, не вызывает workbook.save() и вообще не пишет на диск.

Workbook загружается РОВНО ОДИН раз за вызов read_flat_table — двойная
загрузка (например, отдельно под data_only=True) на этом этапе не
делается.

======================================================================
Значения ячеек — без нормализации
======================================================================

CellRecord.value содержит ровно то, что вернул openpyxl для данной
ячейки — никакого приведения типов (float -> int/str, int -> str,
None -> "", date/datetime -> str, bool -> int и т.п.). Это необходимо
для сохранения safety-контракта Stage 5: float никогда не должен
"восстанавливаться" в ИНН/ОГРН/ОГРНИП, а бывает это только тогда, когда
чей-то код пытается досрочно привести значение к строке/int. Этот модуль
такого не делает.

======================================================================
Границы таблицы и merged cells
======================================================================

Диапазон чтения — worksheet.max_row/worksheet.max_column "как есть", без
попытки эвристически определить "настоящий" конец таблицы (это отдельная
политика будущих этапов). Merged cells не разворачиваются: top-left
ячейка диапазона хранит значение, остальные позиции естественно дают
None (родное поведение openpyxl) — это тоже сознательно не трогается на
Stage 6.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Union

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from app.excel.models import CellRecord, FlatTable

_PathLike = Union[str, Path]


def list_sheet_names(path: _PathLike) -> tuple[str, ...]:
    """
    Возвращает имена worksheets в исходном порядке книги, включая скрытые
    листы (hidden sheets НЕ фильтруются). Файл только читается.
    """
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def read_flat_table(
    path: _PathLike,
    *,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
) -> FlatTable:
    """
    Читает один worksheet целиком в FlatTable.

    :param path: путь к .xlsx файлу (только читается).
    :param sheet_name: имя листа; None -> активный лист книги (без каких-
        либо попыток угадать "правильный" лист иначе). Если указанного
        листа нет — поднимается KeyError (стандартное поведение
        openpyxl.Workbook.__getitem__, не подменяется молчаливым
        переключением на другой лист).
    :param header_row: 1-based номер строки заголовков; строки данных
        начинаются со следующей строки (header_row + 1). Значение
        строго >= 1 — иначе ValueError/TypeError с понятным сообщением.
        Также header_row не может превышать worksheet.max_row выбранного
        листа — иначе ValueError (см. _read_flat_table): reader не должен
        молча создавать искусственную строку заголовков "из воздуха" для
        строки, которой в листе физически не существует. Единственное
        исключение — пустой новый лист (worksheet.max_row == 1, как у
        openpyxl по умолчанию): header_row=1 для него остаётся допустимым.
        Автоматическое угадывание строки заголовка не производится.
    """
    _validate_header_row(header_row)

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    try:
        worksheet = workbook.active if sheet_name is None else workbook[sheet_name]
        return _read_flat_table(worksheet, header_row)
    finally:
        workbook.close()


def _validate_header_row(header_row: object) -> None:
    if isinstance(header_row, bool) or not isinstance(header_row, int):
        raise TypeError(f"header_row должен быть int, получено: {type(header_row)!r}")
    if header_row < 1:
        raise ValueError(f"header_row должен быть >= 1, получено: {header_row}")


def _read_flat_table(worksheet: Worksheet, header_row: int) -> FlatTable:
    """
    Один линейный проход по worksheet.iter_rows() в пределах фактического
    worksheet.max_row/max_column — O(количество ячеек), без повторных
    сканирований и без поиска по уже прочитанным данным.
    """
    max_row = worksheet.max_row
    max_column = worksheet.max_column

    if header_row > max_row:
        raise ValueError(
            f"header_row={header_row} превышает worksheet.max_row={max_row} "
            f"листа {worksheet.title!r} — reader не создаёт искусственную "
            "строку заголовков за пределами фактического листа"
        )

    header_cells = next(
        worksheet.iter_rows(min_row=header_row, max_row=header_row, min_col=1, max_col=max_column)
    )
    header_record = tuple(_cell_record(cell) for cell in header_cells)

    rows = tuple(
        tuple(_cell_record(cell) for cell in row)
        for row in worksheet.iter_rows(
            min_row=header_row + 1, max_row=max_row, min_col=1, max_col=max_column
        )
    )

    return FlatTable(sheet_name=worksheet.title, header_row=header_record, rows=rows)


def _cell_record(cell: Cell) -> CellRecord:
    """
    Единственное место, где Cell openpyxl превращается в CellRecord.
    cell.value передаётся как есть — без нормализации/coercion (см.
    docstring модуля).
    """
    return CellRecord(row=cell.row, column=cell.column, value=cell.value)
