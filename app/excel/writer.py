"""
Excel Writer (Stage 8.2) — запись анонимизированных значений обратно в
.xlsx и встраивание externally supplied job_id в metadata.

======================================================================
Границы ответственности (frozen Stage 8.1 contract)
======================================================================

Writer НЕ создаёт новый Workbook из FlatTable — FlatTable/CellRecord
(app.excel.models) не содержат ни стилей, ни merged ranges, ни formulas
для нетронутых ячеек, ни hidden-sheet state, ни defined names, ни
workbook/core properties, поэтому построение книги "с нуля" не могло бы
воспроизвести структуру исходной книги. Вместо этого Writer:

    1. загружает исходную .xlsx книгу через openpyxl;
    2. ПОЛНОСТЬЮ проверяет ВСЕ переданные FlatTable (validation-before-
       mutation, см. ниже) — ни одна ячейка не изменяется, пока не
       проверены все таблицы целиком;
    3. только после успешной проверки изменяет значения существующих
       ячеек через cell.value = ... (без пересоздания Cell и без
       копирования/сброса style — font/fill/border/alignment/
       number_format/protection остаются нетронутыми);
    4. заменяет устаревшее/отсутствующее workbook-metadata значение
       job_id ровно одним свежим;
    5. сохраняет отдельную destination-книгу атомарно.

Исходная книга — read-only input на уровне проектного контракта: Writer
никогда не сохраняет результат в source_path, независимо от исхода
операции.

Writer НЕ анонимизирует значения, НЕ генерирует alias/token/job_id, НЕ
знает про MappingStore/IdentifierMappingStore/ProvenanceStore и не
принимает пароль или сырые идентификаторы — он работает исключительно с
уже готовыми FlatTable (результат app.anonymizer), путями и job_id.

======================================================================
job_id metadata — custom document property (frozen Stage 8.1 design)
======================================================================

job_id встраивается как единственная custom document property с именем
DataAnonymizer.JobId (openpyxl.packaging.custom.StringProperty). Если
исходная книга уже содержит property с этим именем (в т.ч. устаревший
job_id от предыдущего запуска анонимизации), ВСЕ такие property
удаляются ПЕРЕД добавлением свежей. CustomPropertyList.append() сам по
себе не допускает дублирования имён (поднимает ValueError) — но это
гарантирует уникальность только для книг, полностью построенных через
openpyxl API; CustomPropertyList.from_tree() (парсинг docProps/
custom.xml при ЗАГРУЗКЕ книги) уникальность НЕ проверяет, поэтому
исходная книга, отредактированная не через openpyxl, может физически
содержать несколько property с этим именем. CustomPropertyList.
__delitem__ удаляет ровно ОДНО совпадение за вызов, поэтому Writer
повторяет удаление до KeyError — гарантированно приводя output к
exactly-one состоянию независимо от числа stale-properties в источнике.
Прочие custom properties (не совпадающие по имени) не затрагиваются.

job_id — opaque non-empty str, полностью generation-method agnostic:
Writer не требует конкретного формата (32 hex, UUID, secrets.token_hex
и т.п.), не обрезает и не нормализует строку — валидная строка
сохраняется EXACT как передана.

======================================================================
Координаты и структурные границы
======================================================================

Запись выполняется исключительно по CellRecord.row/CellRecord.column
(1-based, как у openpyxl) — никогда по позиции в списке/enumerate.

worksheet.cell(row, column) в openpyxl молча РАСШИРЯЕТ worksheet, если
координата выходит за текущий max_row/max_column, — даже при обычном
чтении, без присваивания value (проверено экспериментально на
openpyxl 3.1.5). Поэтому Writer фиксирует max_row/max_column каждого
целевого листа СРАЗУ после загрузки книги, ДО первого обращения к
worksheet.cell() для supplied-координат, и сравнивает с этим снимком —
никогда с текущим (потенциально уже расширенным) состоянием листа.

Эта проверка — только структурная защита от случайного расширения
листа. Она НЕ доказывает, что переданный FlatTable действительно был
получен именно из этой исходной книги (см. Stage 8.1 corrective pass,
раздел про FlatTable ↔ source workbook binding) — такое сопоставление
остаётся ответственностью orchestration layer, вне границ Writer.

======================================================================
Merged cells (frozen Stage 8.1 contract)
======================================================================

Top-left ячейка merged range — обычный Cell, запись в неё выполняется
как в любую другую ячейку.

Non-top-left ячейка merged range — openpyxl.cell.cell.MergedCell,
её .value физически всегда None и доступен только для чтения
(присваивание поднимает AttributeError на уровне openpyxl). Writer НЕ
использует это исключение как механизм валидации — вместо этого:

    - CellRecord.value is None  -> безопасный no-op, merged range не
      трогается;
    - CellRecord.value is not None -> явный WriterValidationError ДО
      какой-либо мутации: non-top-left merged cell не может
      самостоятельно нести значение в исходной книге, поэтому непустое
      значение здесь означает некорректный/синтетический FlatTable.

======================================================================
Что Writer сознательно не проверяет повторно
======================================================================

CellRecord.row/column уже валидируются как строго положительные int
(bool исключён) в CellRecord.__post_init__ (app.excel.models) — Writer
не дублирует эту проверку, полагаясь на то, что любой переданный
CellRecord уже прошёл её на этапе конструирования.
"""

from __future__ import annotations

import os
import tempfile
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Union

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.packaging.custom import StringProperty
from openpyxl.workbook.workbook import Workbook

from app.excel.models import FlatTable

_PathLike = Union[str, Path]

_SUPPORTED_EXTENSION = ".xlsx"
_JOB_ID_PROPERTY_NAME = "DataAnonymizer.JobId"


class WriterValidationError(ValueError):
    """
    Переданные аргументы/FlatTable не удовлетворяют Writer-контракту:
    неверное расширение файла, source_path == destination_path, пустой
    job_id, пустой набор tables, дублирующиеся FlatTable.sheet_name,
    отсутствующий в исходной книге лист, дублирующиеся координаты
    внутри одного FlatTable, координата за структурными границами
    исходного листа, непустое значение для non-top-left merged cell.

    Все перечисленные проверки выполняются ДО какой-либо мутации
    исходной/выходной книги (см. docstring модуля).
    """


def write_anonymized_workbook(
    source_path: _PathLike,
    tables: Iterable[FlatTable],
    destination_path: _PathLike,
    *,
    job_id: str,
) -> None:
    """
    Записывает анонимизированные значения tables поверх копии
    source_path и сохраняет результат в destination_path вместе с
    embedded job_id.

    :param source_path: путь к исходной .xlsx книге; только читается,
        никогда не изменяется.
    :param tables: один или несколько FlatTable (обычно — результат
        app.anonymizer.anonymize_flat_table). Пустая коллекция
        отклоняется как invalid input — в текущей архитектуре Writer
        вызывается только когда есть хотя бы один обработанный
        результат анонимизации; создание "пустой" копии книги
        исключительно ради job_id не имеет текущего use case.
    :param destination_path: путь к выходной .xlsx книге; должен
        отличаться от source_path; родительский каталог должен уже
        существовать (автоматически не создаётся).
    :param job_id: externally supplied opaque non-empty строка;
        Writer её не генерирует и не проверяет происхождение/формат.

    :raises TypeError: source_path/destination_path не str/Path, либо
        tables содержит элемент, не являющийся FlatTable.
    :raises FileNotFoundError: source_path не существует, либо
        родительский каталог destination_path не существует.
    :raises WriterValidationError: любое нарушение Writer-контракта,
        перечисленное в докстринге класса.
    :raises Exception: любая ошибка openpyxl (повреждённый файл,
        ошибка сохранения и т.п.) распространяется как есть.
    """
    source, destination = _validate_paths(source_path, destination_path)
    _validate_job_id(job_id)
    materialized_tables = _validate_tables(tables)

    workbook = openpyxl.load_workbook(
        str(source), read_only=False, data_only=False, keep_vba=False, keep_links=True
    )
    try:
        pending_writes = _validate_and_resolve(workbook, materialized_tables)

        # Mutation phase — начинается ТОЛЬКО после успешной валидации
        # ВСЕХ переданных tables (validation-before-mutation).
        for cell, value in pending_writes:
            cell.value = value

        _replace_job_id_property(workbook, job_id)
        _atomic_save(workbook, destination)
    finally:
        workbook.close()


# ----------------------------------------------------------------------
# Валидация аргументов (до загрузки книги)
# ----------------------------------------------------------------------


def _validate_paths(source_path: object, destination_path: object) -> tuple[Path, Path]:
    if not isinstance(source_path, (str, Path)):
        raise TypeError(f"source_path должен быть str или Path, получено: {type(source_path)!r}")
    if not isinstance(destination_path, (str, Path)):
        raise TypeError(
            f"destination_path должен быть str или Path, получено: {type(destination_path)!r}"
        )

    source = Path(source_path)
    destination = Path(destination_path)

    if source.suffix.lower() != _SUPPORTED_EXTENSION:
        raise WriterValidationError(
            f"source_path должен иметь расширение {_SUPPORTED_EXTENSION}, "
            f"получено: {source.suffix!r} ({source})"
        )
    if destination.suffix.lower() != _SUPPORTED_EXTENSION:
        raise WriterValidationError(
            f"destination_path должен иметь расширение {_SUPPORTED_EXTENSION}, "
            f"получено: {destination.suffix!r} ({destination})"
        )

    if not source.exists():
        raise FileNotFoundError(f"source_path не существует: {source}")

    if _paths_refer_to_same_file(source, destination):
        raise WriterValidationError(
            "source_path и destination_path не должны совпадать — "
            "исходная книга не должна изменяться in-place"
        )

    if not destination.parent.exists():
        raise FileNotFoundError(
            f"Родительский каталог destination_path не существует: {destination.parent}"
        )

    return source, destination


def _paths_refer_to_same_file(a: Path, b: Path) -> bool:
    """
    Минимальная, но надёжная нормализация: resolve() устраняет
    относительность/'..', os.path.normcase учитывает
    регистронезависимость файловой системы Windows. Полноценный
    filesystem-identity framework (inode/dev, символьные ссылки сверх
    resolve()) сознательно не строится — избыточно для MVP-контракта
    "source_path == destination_path -> reject".
    """
    return os.path.normcase(str(a.resolve())) == os.path.normcase(str(b.resolve()))


def _validate_job_id(job_id: object) -> None:
    if not isinstance(job_id, str) or not job_id.strip():
        raise WriterValidationError(f"job_id должен быть непустой строкой, получено: {job_id!r}")


def _validate_tables(tables: Iterable[FlatTable]) -> tuple[FlatTable, ...]:
    materialized = tuple(tables)

    for table in materialized:
        if not isinstance(table, FlatTable):
            raise TypeError(f"tables должен содержать только FlatTable, получено: {type(table)!r}")

    if not materialized:
        raise WriterValidationError("tables не должен быть пустым")

    sheet_name_counts = Counter(table.sheet_name for table in materialized)
    duplicates = sorted(name for name, count in sheet_name_counts.items() if count > 1)
    if duplicates:
        raise WriterValidationError(
            f"Обнаружены дублирующиеся FlatTable.sheet_name: {duplicates}"
        )

    return materialized


# ----------------------------------------------------------------------
# Валидация координат/merged-cells относительно загруженной книги
# ----------------------------------------------------------------------


def _validate_and_resolve(
    workbook: Workbook, tables: tuple[FlatTable, ...]
) -> list[tuple[Cell, object]]:
    """
    Проверяет ВСЕ координаты ВСЕХ tables относительно уже загруженной
    книги и возвращает список (cell, value) для последующей mutation
    phase. Ни одна cell.value не присваивается здесь — только чтение и
    структурная проверка.
    """
    pending_writes: list[tuple[Cell, object]] = []

    for table in tables:
        if table.sheet_name not in workbook.sheetnames:
            raise WriterValidationError(
                f"Лист {table.sheet_name!r} отсутствует в исходной книге"
            )

        worksheet = workbook[table.sheet_name]
        # Снимок ДО первого обращения к worksheet.cell() для supplied-
        # координат — само обращение может расширить max_row/max_column
        # (см. docstring модуля), поэтому сравнивать нужно только с этим
        # зафиксированным снимком, а не с текущим состоянием листа.
        original_max_row = worksheet.max_row
        original_max_column = worksheet.max_column

        all_records = tuple(table.header_row) + tuple(
            cell_record for row in table.rows for cell_record in row
        )

        seen_coordinates: set[tuple[int, int]] = set()
        for record in all_records:
            coordinate = (record.row, record.column)
            if coordinate in seen_coordinates:
                raise WriterValidationError(
                    f"Лист {table.sheet_name!r}: дублирующаяся координата "
                    f"(row={record.row}, column={record.column})"
                )
            seen_coordinates.add(coordinate)

            if record.row > original_max_row or record.column > original_max_column:
                raise WriterValidationError(
                    f"Лист {table.sheet_name!r}: координата (row={record.row}, "
                    f"column={record.column}) выходит за структурные границы "
                    f"исходного листа (max_row={original_max_row}, "
                    f"max_column={original_max_column})"
                )

            cell = worksheet.cell(row=record.row, column=record.column)
            if isinstance(cell, MergedCell):
                if record.value is None:
                    continue
                raise WriterValidationError(
                    f"Лист {table.sheet_name!r}: координата (row={record.row}, "
                    f"column={record.column}) — non-top-left ячейка "
                    "объединённого диапазона и не может получить непустое значение"
                )

            pending_writes.append((cell, record.value))

    return pending_writes


# ----------------------------------------------------------------------
# job_id metadata
# ----------------------------------------------------------------------


def _replace_job_id_property(workbook: Workbook, job_id: str) -> None:
    """
    Удаляет ВСЕ существующие custom properties с именем
    _JOB_ID_PROPERTY_NAME (не только первую) и добавляет ровно одну
    свежую. Одного del недостаточно: CustomPropertyList.append() сам не
    допускает дублирования имён (Stage 8.1), но
    CustomPropertyList.from_tree() (парсинг docProps/custom.xml при
    ЗАГРУЗКЕ книги) уникальность НЕ проверяет — исходная книга,
    отредактированная не через openpyxl, может физически содержать
    несколько property с этим именем. CustomPropertyList.__delitem__
    удаляет ровно ОДНО (первое найденное) совпадение по имени за вызов
    (см. openpyxl.packaging.custom), поэтому цикл повторяет del до
    KeyError — только тогда все совпадения гарантированно удалены.
    Unrelated properties не затрагиваются (сравнение по exact имени).
    """
    while True:
        try:
            del workbook.custom_doc_props[_JOB_ID_PROPERTY_NAME]
        except KeyError:
            break
    workbook.custom_doc_props.append(StringProperty(name=_JOB_ID_PROPERTY_NAME, value=job_id))


# ----------------------------------------------------------------------
# Атомарная запись (тот же паттерн, что и у encrypted-store сериализации:
# temp-файл в целевом каталоге -> save -> fsync -> os.replace)
# ----------------------------------------------------------------------


def _atomic_save(workbook: Workbook, destination: Path) -> None:
    directory = destination.parent
    fd, tmp_name = tempfile.mkstemp(dir=str(directory), prefix=f".{destination.name}.", suffix=".tmp")
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        workbook.save(str(tmp_path))

        # fsync требует дескриптор, открытый для записи (O_RDWR) — на
        # Windows os.fsync на дескрипторе, открытом только для чтения
        # (O_RDONLY), поднимает OSError: Bad file descriptor (проверено
        # экспериментально на этом окружении).
        fsync_fd = os.open(str(tmp_path), os.O_RDWR)
        try:
            os.fsync(fsync_fd)
        finally:
            os.close(fsync_fd)

        os.replace(str(tmp_path), str(destination))
    except Exception:
        try:
            tmp_path.unlink(missing_ok=True)
        except OSError:
            # Best-effort cleanup: ошибка удаления temp-файла не должна
            # маскировать исходную ошибку.
            pass
        raise
