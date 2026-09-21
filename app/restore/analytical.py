"""
Derived Analytical Workbook Restore Core (Stage 9C).

Переиспользуемая, чисто локальная value-driven логика восстановления —
без provenance, без job_id, без владения filesystem/store lifecycle.

Этот модуль НЕ читает/пишет файлы, НЕ создаёт workbook/temp, НЕ работает
с metadata и путями, НЕ мутирует MappingStore/IdentifierMappingStore — он
оперирует уже загруженным openpyxl Workbook и уже открытыми store-
объектами, переданными вызывающей стороной (app.orchestration.
restore_analytical).

======================================================================
Главный инвариант — value-driven, а не coordinate-driven
======================================================================

В отличие от Stage 9B (app.restore.core), здесь НЕТ ProvenanceStore и
НЕТ frozen-координат исходного anonymization job — derived analytical
workbook мог быть произвольно реструктурирован внешним инструментом
(новые листы/строки/колонки, сортировка, агрегаты). Единственный
источник истины — ТЕКУЩЕЕ значение ячейки: exact whole-cell lookup в
MappingStore (entity alias) и, опционально, в IdentifierMappingStore
(identifier token).

======================================================================
Dual lookup — без priority-order
======================================================================

Для каждой eligible строки выполняются ОБА lookup (entity и identifier,
если identifier_store передан), не по очереди с коротким замыканием на
первом hit. Production alias/token генераторы сегодня структурно не
могут создать совпадающую строку (разная длина итогового значения из-за
разных префиксов при общем PRODUCTION_RANDOM_LENGTH — см. Stage 9C
Architecture Review), но это dual-lookup — defense-in-depth против
будущего изменения генераторов, а не полагание на текущий факт.

======================================================================
All-or-nothing
======================================================================

build_replacement_plan только строит RAM-план (entity_replacements,
identifier_replacements) и никогда не мутирует workbook. Мутация — через
apply_replacement_plan, единственную точку, вызываемую вызывающей
стороной только после полного успешного scan (без ambiguity).
"""

from __future__ import annotations

import dataclasses
from typing import Optional

from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.workbook import Workbook

from app.mapping.base import MappingStore
from app.mapping.identifier_base import IdentifierMappingStore


class AmbiguousRestoreError(Exception):
    """
    Одна и та же exact строка одновременно найдена и как entity alias
    (MappingStore.get_by_alias), и как identifier token
    (IdentifierMappingStore.get_by_token) — восстановление неоднозначно,
    весь job завершается ошибкой (all-or-nothing).

    Сообщение содержит только 1-based worksheet index, row, column и
    generic reason — никогда sheet name, cell value, alias, token, real
    company, identifier value, repr каких-либо store-объектов/записей.
    """


@dataclasses.dataclass(frozen=True)
class AnalyticalCellReplacement:
    """Один запланированный, но ещё не применённый к workbook replacement."""

    worksheet_index: int
    row: int
    column: int
    value: str


def build_replacement_plan(
    workbook: Workbook,
    mapping_store: MappingStore,
    identifier_store: Optional[IdentifierMappingStore],
) -> tuple[tuple[AnalyticalCellReplacement, ...], tuple[AnalyticalCellReplacement, ...]]:
    """
    Сканирует ВСЕ worksheets workbook (включая hidden/veryHidden;
    chartsheets не входят в workbook.worksheets) и строит полный RAM-план
    замен, не мутируя workbook. При обнаружении ambiguity (см.
    AmbiguousRestoreError) немедленно поднимает исключение — до
    какой-либо мутации.

    :returns: (entity_replacements, identifier_replacements).
    """
    entity_replacements: list[AnalyticalCellReplacement] = []
    identifier_replacements: list[AnalyticalCellReplacement] = []

    for worksheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        max_row = worksheet.max_row
        max_column = worksheet.max_column

        for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if cell.data_type == "e":
                    continue

                value = cell.value
                if not isinstance(value, str):
                    continue
                if value.startswith("="):
                    continue

                _resolve_cell(
                    value,
                    worksheet_index,
                    cell.row,
                    cell.column,
                    mapping_store,
                    identifier_store,
                    entity_replacements,
                    identifier_replacements,
                )

    return tuple(entity_replacements), tuple(identifier_replacements)


def _resolve_cell(
    value: str,
    worksheet_index: int,
    row: int,
    column: int,
    mapping_store: MappingStore,
    identifier_store: Optional[IdentifierMappingStore],
    entity_replacements: list[AnalyticalCellReplacement],
    identifier_replacements: list[AnalyticalCellReplacement],
) -> None:
    entity_entry = mapping_store.get_by_alias(value)
    identifier_entry = None if identifier_store is None else identifier_store.get_by_token(value)

    if entity_entry is not None and identifier_entry is not None:
        raise AmbiguousRestoreError(
            f"Ambiguous analytical restore value at worksheet {worksheet_index}, "
            f"row {row}, column {column}"
        )

    if entity_entry is not None:
        entity_replacements.append(
            AnalyticalCellReplacement(
                worksheet_index=worksheet_index, row=row, column=column, value=entity_entry.real_value
            )
        )
    elif identifier_entry is not None:
        identifier_replacements.append(
            AnalyticalCellReplacement(
                worksheet_index=worksheet_index,
                row=row,
                column=column,
                value=identifier_entry.identifier_value,
            )
        )
    # Иначе — значение не найдено ни в одном store: unchanged, не ошибка.


def apply_replacement_plan(
    workbook: Workbook,
    entity_replacements: tuple[AnalyticalCellReplacement, ...],
    identifier_replacements: tuple[AnalyticalCellReplacement, ...],
) -> None:
    """
    Единственная точка мутации workbook — применяет уже полностью
    разрешённый (без ambiguity) replacement plan. Координаты плана уже
    получены из существующих ячеек через iter_rows() на этапе scan,
    поэтому повторное обращение через worksheet.cell() безопасно и не
    расширяет worksheet.
    """
    for replacement in (*entity_replacements, *identifier_replacements):
        worksheet = workbook.worksheets[replacement.worksheet_index - 1]
        worksheet.cell(row=replacement.row, column=replacement.column).value = replacement.value
