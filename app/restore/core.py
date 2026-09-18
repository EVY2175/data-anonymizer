"""
Restore Core (Stage 9A) — переиспользуемая, чисто локальная логика
восстановления, без владения filesystem/store lifecycle.

Этот модуль НЕ читает/пишет файлы, НЕ создаёт workbook/provenance sidecar,
НЕ генерирует job_id, НЕ мутирует MappingStore/IdentifierMappingStore/
ProvenanceStore и не выполняет network/batch/Entity Resolution — он
оперирует уже загруженным openpyxl Workbook и уже открытыми store-объектами,
переданными вызывающей стороной (Stage 9B, app.orchestration.
restore_single_file).

======================================================================
Identifier vs entity restore — разные источники истины
======================================================================

Identifier restore целиком управляется provenance: координата + exact
token-equality — единственное основание для lookup в IdentifierMappingStore.
Никакого identifier lookup не выполняется до подтверждения exact token
match (frozen Stage 7C.5/Stage 9 Contract Freeze).

Entity restore управляется значением ячейки: exact whole-cell
MappingStore.get_by_alias(value) lookup, без rules_by_sheet и без
provenance. Если alias не найден — это НЕ ошибка (без rules нельзя
утверждать, что произвольная строка обязана быть alias), ячейка остаётся
unchanged. Provenance-координаты (identifier) исключаются из entity scan
полностью, до фактического сравнения — identifier provenance всегда
authoritative для своих координат.

======================================================================
All-or-nothing
======================================================================

resolve_identifier_replacements/resolve_entity_replacements только строят
RAM-план замен (ReplacementPlanEntry) и никогда не мутируют workbook.
Мутация происходит исключительно через apply_replacement_plan — единую
точку, вызываемую Stage 9B только после полного успеха ОБЕИХ фаз
resolution.
"""

from __future__ import annotations

import dataclasses
from collections.abc import Iterable, Mapping
from typing import Optional, Union

from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.workbook import Workbook

from app.mapping.base import MappingStore
from app.mapping.identifier_base import IdentifierMappingStore
from app.models.provenance import IdentifierCellProvenance, IdentifierRepresentation

_JOB_ID_PROPERTY_NAME = "DataAnonymizer.JobId"


class RestoreValidationError(ValueError):
    """
    Restore-конфигурация не удовлетворяет preflight/binding-контракту:
    неверный тип/расширение пути, отсутствующий source/provenance, уже
    существующий destination, отсутствующий/дублирующийся/пустой workbook
    job_id, job_id mismatch, отсутствующий identifier_store при непустом
    provenance, пустой provenance_password.

    Поднимается ДО начала per-cell resolution — никогда не оборачивает
    UnresolvedRestoreError или leaf-исключения crypto/openpyxl/IO.
    """


class UnresolvedRestoreError(Exception):
    """
    Content-level неразрешимость, обнаруженная в процессе per-cell
    resolution: token mismatch, provenance-координата не резолвится
    (отсутствующий лист/выход за структурные границы/non-top-left merged
    cell), token отсутствует в IdentifierMappingStore, INTEGER
    representation conversion fails.

    Поднимается только после успешного preflight/binding — сигнализирует
    hard failure всего restore job (all-or-nothing).
    """


@dataclasses.dataclass(frozen=True)
class ReplacementPlanEntry:
    """
    Один запланированный, но ещё не применённый к workbook replacement.

    Не хранит expected_current_value: resolution и apply выполняются над
    одним и тем же in-memory workbook-объектом без какого-либо
    промежуточного кода, способного его изменить (единственный поток
    исполнения) — повторная проверка защищала бы от угрозы, которой в
    этой архитектуре не существует.
    """

    sheet_name: str
    row: int
    column: int
    restored_value: Union[str, int]


def read_workbook_job_id(workbook: Workbook) -> str:
    """
    Находит ВСЕ custom properties с именем DataAnonymizer.JobId через
    полную итерацию (openpyxl.packaging.custom.CustomPropertyList.
    from_tree() не гарантирует уникальность имён при парсинге стороннего/
    отредактированного файла) и возвращает EXACT значение единственной
    найденной записи.

    :raises RestoreValidationError: property отсутствует, дублируется,
        значение не str, значение пустое/whitespace-only.
    """
    matches = [prop for prop in workbook.custom_doc_props if prop.name == _JOB_ID_PROPERTY_NAME]

    if not matches:
        raise RestoreValidationError(
            f"В workbook отсутствует обязательная custom property {_JOB_ID_PROPERTY_NAME!r}"
        )
    if len(matches) > 1:
        raise RestoreValidationError(
            f"Custom property {_JOB_ID_PROPERTY_NAME!r} дублируется в workbook"
        )

    value = matches[0].value
    if not isinstance(value, str):
        raise RestoreValidationError(
            f"Custom property {_JOB_ID_PROPERTY_NAME!r} должна быть строкой, "
            f"получено: {type(value)!r}"
        )
    if not value.strip():
        raise RestoreValidationError(
            f"Custom property {_JOB_ID_PROPERTY_NAME!r} не может быть пустой"
        )

    return value


def check_job_id_binding(workbook_job_id: str, provenance_job_id: str) -> None:
    """
    До любого per-cell resolution: workbook job_id должен ТОЧНО совпадать
    с job_id открытого provenance sidecar.

    :raises RestoreValidationError: mismatch.
    """
    if workbook_job_id != provenance_job_id:
        raise RestoreValidationError(
            "job_id workbook не совпадает с job_id provenance sidecar"
        )


def snapshot_structural_bounds(workbook: Workbook) -> dict[str, tuple[int, int]]:
    """
    Снимок (max_row, max_column) для КАЖДОГО worksheet, взятый сразу после
    загрузки книги — до первого worksheet.cell()-вызова для любой
    provenance-координаты. worksheet.cell() с произвольными координатами
    молча расширяет worksheet, даже без присваивания value (проверено
    экспериментально на Stage 8.1) — снимок должен предшествовать любому
    такому вызову и использоваться на протяжении всего job'а.
    """
    return {worksheet.title: (worksheet.max_row, worksheet.max_column) for worksheet in workbook.worksheets}


def resolve_identifier_replacements(
    workbook: Workbook,
    structural_bounds: Mapping[str, tuple[int, int]],
    provenance_entries: Iterable[IdentifierCellProvenance],
    identifier_store: Optional[IdentifierMappingStore],
) -> tuple[ReplacementPlanEntry, ...]:
    """
    Резолвит ВСЕ provenance entries в RAM-план замен, не мутируя workbook.

    Порядок на каждую entry: лист существует -> координата в границах
    structural_bounds -> получить cell -> non-top-left MergedCell
    отклоняется -> cell.value EXACT equal entry.token (identifier lookup
    НЕ выполняется до этого сравнения) -> identifier_store.get_by_token
    -> representation restore. Любой сбой — немедленный
    UnresolvedRestoreError, вся функция возвращает результат только при
    полном успехе всех entries (all-or-nothing).

    Если provenance_entries пуст — identifier_store не используется вовсе
    (может быть None).
    """
    entries = tuple(provenance_entries)
    if not entries:
        return ()

    plan: list[ReplacementPlanEntry] = []
    for entry in entries:
        bounds = structural_bounds.get(entry.sheet_name)
        if bounds is None:
            raise UnresolvedRestoreError(
                f"Provenance-координата ссылается на отсутствующий в workbook лист {entry.sheet_name!r}"
            )

        max_row, max_column = bounds
        if entry.row > max_row or entry.column > max_column:
            raise UnresolvedRestoreError(
                f"Provenance-координата ({entry.sheet_name!r}, row={entry.row}, "
                f"column={entry.column}) выходит за структурные границы листа"
            )

        worksheet = workbook[entry.sheet_name]
        cell = worksheet.cell(row=entry.row, column=entry.column)
        if isinstance(cell, MergedCell):
            raise UnresolvedRestoreError(
                f"Provenance-координата ({entry.sheet_name!r}, row={entry.row}, "
                f"column={entry.column}) указывает на non-top-left ячейку "
                "объединённого диапазона"
            )

        if cell.value != entry.token:
            raise UnresolvedRestoreError(
                f"Значение ячейки ({entry.sheet_name!r}, row={entry.row}, "
                f"column={entry.column}) не совпадает с ожидаемым provenance-token"
            )

        mapping_entry = identifier_store.get_by_token(entry.token)  # type: ignore[union-attr]
        if mapping_entry is None:
            raise UnresolvedRestoreError(
                f"Token для координаты ({entry.sheet_name!r}, row={entry.row}, "
                f"column={entry.column}) отсутствует в IdentifierMappingStore"
            )

        restored_value = _restore_representation(
            mapping_entry.identifier_value, entry.representation, entry.sheet_name, entry.row, entry.column
        )

        plan.append(
            ReplacementPlanEntry(
                sheet_name=entry.sheet_name, row=entry.row, column=entry.column, restored_value=restored_value
            )
        )

    return tuple(plan)


def _restore_representation(
    identifier_value: str, representation: IdentifierRepresentation, sheet_name: str, row: int, column: int
) -> Union[str, int]:
    if representation is IdentifierRepresentation.STRING:
        return identifier_value

    # IdentifierRepresentation содержит ровно два значения (STRING/INTEGER) —
    # единственная оставшаяся ветка.
    try:
        return int(identifier_value)
    except ValueError as exc:
        raise UnresolvedRestoreError(
            f"Не удалось восстановить INTEGER representation для координаты "
            f"({sheet_name!r}, row={row}, column={column})"
        ) from exc


def resolve_entity_replacements(
    workbook: Workbook,
    structural_bounds: Mapping[str, tuple[int, int]],
    mapping_store: MappingStore,
    excluded_coordinates: set[tuple[str, int, int]],
) -> tuple[ReplacementPlanEntry, ...]:
    """
    Резолвит exact entity alias replacements во ВСЕХ worksheets workbook,
    в пределах structural_bounds, не мутируя workbook.

    Пропускаются: координаты из excluded_coordinates (provenance —
    identifier всегда authoritative для своих координат), non-top-left
    MergedCell, non-str значения, формулы ("=..."), пустые/whitespace-only
    строки. Для остальных str-значений — exact
    MappingStore.get_by_alias(value); отсутствие совпадения НЕ является
    ошибкой (без rules нельзя утверждать, что строка обязана быть alias).
    """
    plan: list[ReplacementPlanEntry] = []

    for worksheet in workbook.worksheets:
        max_row, max_column = structural_bounds[worksheet.title]
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
            for cell in row:
                coordinate = (worksheet.title, cell.row, cell.column)
                if coordinate in excluded_coordinates:
                    continue
                if isinstance(cell, MergedCell):
                    continue

                value = cell.value
                if not isinstance(value, str):
                    continue
                if value.startswith("="):
                    continue
                if not value.strip():
                    continue

                entry = mapping_store.get_by_alias(value)
                if entry is not None:
                    plan.append(
                        ReplacementPlanEntry(
                            sheet_name=worksheet.title,
                            row=cell.row,
                            column=cell.column,
                            restored_value=entry.real_value,
                        )
                    )

    return tuple(plan)


def apply_replacement_plan(workbook: Workbook, plan: Iterable[ReplacementPlanEntry]) -> None:
    """
    Единственная точка мутации workbook — применяет уже полностью
    разрешённый (в обеих фазах) replacement plan. Координаты плана уже
    проверены на структурные границы на этапе resolve, повторное
    расширение worksheet здесь невозможно.
    """
    for entry in plan:
        worksheet = workbook[entry.sheet_name]
        worksheet.cell(row=entry.row, column=entry.column).value = entry.restored_value
