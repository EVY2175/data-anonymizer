"""
External AI Safety — inspection primitives (Stage 10A).

Модуль отвечает ТОЛЬКО на узкий вопрос: "какие локально проверяемые
признаки риска присутствуют в этом workbook?". Он НЕ отвечает на вопрос
"можно ли отправлять этот файл во внешний AI" — такого решения здесь
нет и быть не может (нет registry/workspace). Поэтому в API отсутствуют
is_safe/can_upload: результат — SafetyReport, из которого более
высокие слои (Stage 10B/10C) принимают решение сами, fail-closed.

======================================================================
Что проверяется
======================================================================

BLOCKING (известный признак утечки/восстановленных данных):

    - RESTORED_MARKER: любое custom property с exact именем
      DataAnonymizer.AnalyticallyRestored (Stage 9C) либо
      DataAnonymizer.RestoredFromJobId (Stage 9B). Значение не
      интерпретируется; каждое property (в т.ч. дубликат) — отдельный
      finding. DataAnonymizer.JobId сам по себе finding не является.
    - KNOWN_ENTITY_VALUE / KNOWN_IDENTIFIER_VALUE: exact whole-value
      совпадение (ячейка, название листа, core property) с real_value
      из MappingStore либо identifier_value из IdentifierMappingStore.

UNVERIFIED (присутствие содержимого, которое модуль не анализирует):
формулы, комментарии, гиперссылки, непустые core/custom properties,
defined names, колонтитулы, таблицы, скрытые листы и части пакета
(external links, drawings, charts, chartsheets, media, embeddings,
pivot, customXml, VBA, threaded comments).

======================================================================
Явные ограничения (SafetyReport.not_inspected)
======================================================================

Модуль не заявляет полной OOXML-безопасности. Он не проверяет data
validation, conditional formatting, docProps/app.xml (Company/Manager),
подстроки, значения, которых нет в сторах (raw данные, ещё не
попавшие в mapping, распознать нельзя), и ячейки float/bool/дат.

======================================================================
Exact-семантика
======================================================================

Никаких strip/casefold/normalization/regex/substring/fuzzy. str-ячейка
сравнивается с entity- и identifier-множествами как есть. int (не bool)
сравнивается как str(value) ТОЛЬКО с identifier-множеством, без
реконструкции ведущих нулей (тот же принцип, что у Stage 7B). float,
bool, даты и None в lookup не участвуют.

======================================================================
Только существующие ячейки (ws._cells)
======================================================================

Итерация идёт по worksheet._cells (private API openpyxl, версия
запинена: 3.1.5), а не по dense iter_rows: одна далёкая ячейка
(например XFD1048576) заставила бы iter_rows материализовать миллиарды
пустых Cell. Если _cells недоступен/несовместим — inspection падает
безопасной ошибкой; молчаливого fallback на dense-итерацию нет.

======================================================================
Конфиденциальность и read-only
======================================================================

Ни SafetyFinding, ни SafetyReport, ни ошибки не содержат значений
ячеек, alias/token, real/identifier значений, названий листов, текстов
комментариев/формул, URL, значений или имён custom property, паролей.
Только перечисления, счётчики, 1-based индексы и координаты. Модуль
ничего не сохраняет, не изменяет stores и workbook, не пишет логи.
"""

from __future__ import annotations

import dataclasses
import enum
import hashlib
import zipfile
from pathlib import Path
from types import MappingProxyType
from typing import Optional, Union

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.mapping.base import MappingStore
from app.mapping.identifier_base import IdentifierMappingStore

_PathLike = Union[str, Path]

_SUPPORTED_EXTENSION = ".xlsx"
_HASH_CHUNK_SIZE = 1024 * 1024

MAX_FINDINGS_PER_KIND = 100

_RESTORED_MARKER_PROPERTY_NAMES = (
    "DataAnonymizer.AnalyticallyRestored",
    "DataAnonymizer.RestoredFromJobId",
)
_DATA_ANONYMIZER_PROPERTY_PREFIX = "DataAnonymizer."
_BENIGN_CREATOR_VALUE = "openpyxl"

_HEADER_FOOTER_ATTRIBUTES = (
    "oddHeader",
    "oddFooter",
    "evenHeader",
    "evenFooter",
    "firstHeader",
    "firstFooter",
)


# ----------------------------------------------------------------------
# Enums
# ----------------------------------------------------------------------


class Severity(enum.Enum):
    BLOCKING = "blocking"
    UNVERIFIED = "unverified"


class FindingKind(enum.Enum):
    RESTORED_MARKER = "restored_marker"
    KNOWN_ENTITY_VALUE = "known_entity_value"
    KNOWN_IDENTIFIER_VALUE = "known_identifier_value"
    FORMULA = "formula"
    COMMENT = "comment"
    HYPERLINK = "hyperlink"
    CORE_PROPERTY = "core_property"
    CUSTOM_PROPERTY = "custom_property"
    DEFINED_NAME = "defined_name"
    HEADER_FOOTER = "header_footer"
    TABLE = "table"
    HIDDEN_WORKSHEET = "hidden_worksheet"
    EXTERNAL_LINK = "external_link"
    DRAWING = "drawing"
    CHART = "chart"
    CHARTSHEET = "chartsheet"
    MEDIA = "media"
    EMBEDDING = "embedding"
    PIVOT = "pivot"
    CUSTOM_XML = "custom_xml"
    VBA = "vba"
    THREADED_COMMENT = "threaded_comment"


_BLOCKING_KINDS = frozenset(
    {
        FindingKind.RESTORED_MARKER,
        FindingKind.KNOWN_ENTITY_VALUE,
        FindingKind.KNOWN_IDENTIFIER_VALUE,
    }
)

# Централизованная severity: ни одна запись SafetyFinding не хранит severity
# самостоятельно, поэтому противоречивое состояние невозможно.
SEVERITY_BY_KIND = MappingProxyType(
    {
        kind: (Severity.BLOCKING if kind in _BLOCKING_KINDS else Severity.UNVERIFIED)
        for kind in FindingKind
    }
)


class FindingLocation(enum.Enum):
    WORKBOOK = "workbook"
    WORKSHEET = "worksheet"
    SHEET_TITLE = "sheet_title"
    CELL = "cell"
    CORE_PROPERTY = "core_property"
    PACKAGE = "package"


class CoreProperty(enum.Enum):
    """Проверяемые core properties; value — имя атрибута openpyxl."""

    CREATOR = "creator"
    LAST_MODIFIED_BY = "lastModifiedBy"
    TITLE = "title"
    SUBJECT = "subject"
    DESCRIPTION = "description"
    KEYWORDS = "keywords"
    CATEGORY = "category"
    CONTENT_STATUS = "contentStatus"
    IDENTIFIER = "identifier"


class InspectionCheck(enum.Enum):
    """Что реально было проверено (см. SafetyReport.checks_performed)."""

    RESTORED_MARKER = "restored_marker"
    KNOWN_ENTITY_EXACT = "known_entity_exact"
    KNOWN_IDENTIFIER_EXACT = "known_identifier_exact"
    FORMULA_SURVEY = "formula_survey"
    COMMENT_SURVEY = "comment_survey"
    HYPERLINK_SURVEY = "hyperlink_survey"
    WORKSHEET_TITLE_EXACT = "worksheet_title_exact"
    CORE_PROPERTIES = "core_properties"
    CUSTOM_PROPERTIES = "custom_properties"
    WORKSHEET_EXPOSURES = "worksheet_exposures"
    PACKAGE_INVENTORY = "package_inventory"


class NotInspectedArea(enum.Enum):
    """Области, которые Stage 10A сознательно НЕ проверяет."""

    DATA_VALIDATION = "data_validation"
    CONDITIONAL_FORMATTING = "conditional_formatting"
    EXTENDED_PROPERTIES = "extended_properties"
    SUBSTRING_MATCHES = "substring_matches"
    UNKNOWN_VALUES = "unknown_values"
    NON_TEXT_NON_INTEGER_CELLS = "non_text_non_integer_cells"


class InspectionFailureReason(enum.Enum):
    UNREADABLE_WORKBOOK = "unreadable_workbook"
    FILE_CHANGED_DURING_INSPECTION = "file_changed_during_inspection"


# ----------------------------------------------------------------------
# Ошибки
# ----------------------------------------------------------------------


class ExternalAiSafetyError(Exception):
    """Базовое исключение Stage 10A."""


class ExternalAiInspectionInputError(ExternalAiSafetyError, ValueError):
    """
    Нарушен preflight-контракт inspection (расширение, существование,
    is_file). Сообщение может содержать путь/суффикс (как в остальных
    preflight-модулях проекта), но никогда — содержимое workbook/stores.
    """


_INSPECTION_ERROR_MESSAGES = {
    InspectionFailureReason.UNREADABLE_WORKBOOK: (
        "Workbook не удалось безопасно прочитать для inspection"
    ),
    InspectionFailureReason.FILE_CHANGED_DURING_INSPECTION: (
        "Файл изменился во время inspection"
    ),
}


class ExternalAiInspectionError(ExternalAiSafetyError):
    """
    Inspection не удалось выполнить. Сообщение фиксированное и безопасное:
    текст исключений нижних слоёв (openpyxl/zip/IO) сюда не попадает.
    """

    def __init__(self, reason: InspectionFailureReason) -> None:
        super().__init__(_INSPECTION_ERROR_MESSAGES[reason])
        self.reason = reason


# ----------------------------------------------------------------------
# Модели
# ----------------------------------------------------------------------


def _is_position(value: object) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) and value >= 1


@dataclasses.dataclass(frozen=True)
class SafetyFinding:
    """
    Один finding без каких-либо значений. Инварианты по location:

        CELL:                   worksheet_index, row, column; core_property None
        SHEET_TITLE / WORKSHEET: только worksheet_index
        CORE_PROPERTY:          только core_property
        WORKBOOK / PACKAGE:     ни одного из полей
    """

    kind: FindingKind
    location: FindingLocation
    worksheet_index: Optional[int] = None
    row: Optional[int] = None
    column: Optional[int] = None
    core_property: Optional[CoreProperty] = None

    def __post_init__(self) -> None:
        if not isinstance(self.kind, FindingKind):
            raise ValueError("SafetyFinding.kind должен быть FindingKind")
        if not isinstance(self.location, FindingLocation):
            raise ValueError("SafetyFinding.location должен быть FindingLocation")

        location = self.location
        if location is FindingLocation.CELL:
            valid = (
                _is_position(self.worksheet_index)
                and _is_position(self.row)
                and _is_position(self.column)
                and self.core_property is None
            )
        elif location in (FindingLocation.SHEET_TITLE, FindingLocation.WORKSHEET):
            valid = (
                _is_position(self.worksheet_index)
                and self.row is None
                and self.column is None
                and self.core_property is None
            )
        elif location is FindingLocation.CORE_PROPERTY:
            valid = (
                isinstance(self.core_property, CoreProperty)
                and self.worksheet_index is None
                and self.row is None
                and self.column is None
            )
        else:  # WORKBOOK / PACKAGE
            valid = (
                self.worksheet_index is None
                and self.row is None
                and self.column is None
                and self.core_property is None
            )
        if not valid:
            raise ValueError("SafetyFinding нарушает инварианты своего location")

    @property
    def severity(self) -> Severity:
        return SEVERITY_BY_KIND[self.kind]


@dataclasses.dataclass(frozen=True)
class SafetyReport:
    """
    Результат inspection. Не утверждает "файл безопасен": решение — за
    вызывающим слоем, который обязан учитывать exact finding_counts (а не
    len(findings)), indexed_* (пустой индекс = проверка ничего не значит)
    и непустой not_inspected.
    """

    sha256: str
    worksheet_count: int
    indexed_entity_values: int
    indexed_identifier_values: Optional[int]
    checks_performed: tuple[InspectionCheck, ...]
    not_inspected: tuple[NotInspectedArea, ...]
    finding_counts: tuple[tuple[FindingKind, int], ...]
    findings: tuple[SafetyFinding, ...]
    findings_truncated: bool

    def count(self, kind: FindingKind) -> int:
        for counted_kind, count in self.finding_counts:
            if counted_kind is kind:
                return count
        return 0

    @property
    def has_blocking_findings(self) -> bool:
        return any(
            SEVERITY_BY_KIND[kind] is Severity.BLOCKING for kind, count in self.finding_counts if count > 0
        )

    @property
    def has_unverified_exposures(self) -> bool:
        return any(
            SEVERITY_BY_KIND[kind] is Severity.UNVERIFIED for kind, count in self.finding_counts if count > 0
        )


# ----------------------------------------------------------------------
# SHA-256
# ----------------------------------------------------------------------


def sha256_file(path: _PathLike) -> str:
    """
    SHA-256 содержимого файла (lowercase hex). Чтение потоковое, блоками
    по 1 MiB, в бинарном режиме; путь в хэш не входит; расширение не
    проверяется. Ошибки ввода-вывода проходят как есть.

    :raises TypeError: path не str/Path.
    """
    if not isinstance(path, (str, Path)):
        raise TypeError(f"path должен быть str или Path, получено: {type(path)!r}")

    digest = hashlib.sha256()
    with open(path, "rb") as file:
        while True:
            chunk = file.read(_HASH_CHUNK_SIZE)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


# ----------------------------------------------------------------------
# Центральный API
# ----------------------------------------------------------------------


def inspect_workbook_for_external_ai(
    path: _PathLike,
    mapping_store: MappingStore,
    identifier_store: Optional[IdentifierMappingStore],
) -> SafetyReport:
    """
    Read-only inspection одного .xlsx: см. docstring модуля. Порядок:
    preflight -> hash до -> индексы из сторов (entries() ровно один раз
    на store) -> загрузка workbook -> scan -> инвентарь пакета ->
    close (finally) -> hash после -> сравнение -> report.

    :raises TypeError: неверный тип path/mapping_store/identifier_store.
    :raises ExternalAiInspectionInputError: нарушен preflight.
    :raises ExternalAiInspectionError: workbook/zip не удалось прочитать
        либо файл изменился во время inspection.
    :raises Exception: ошибки sha256_file (I/O) и сторов проходят как есть.
    """
    source = _preflight(path, mapping_store, identifier_store)

    hash_before = sha256_file(source)

    entity_values = frozenset(entry.real_value for entry in mapping_store.entries())
    identifier_values: Optional[frozenset[str]] = (
        None
        if identifier_store is None
        else frozenset(entry.identifier_value for entry in identifier_store.entries())
    )

    collector = _FindingCollector()

    workbook = _load_workbook(source)
    try:
        worksheet_count = len(workbook.worksheets)
        _scan_workbook(workbook, entity_values, identifier_values, collector)
        _inventory_package(source, collector)
    finally:
        workbook.close()

    hash_after = sha256_file(source)
    if hash_before != hash_after:
        raise ExternalAiInspectionError(InspectionFailureReason.FILE_CHANGED_DURING_INSPECTION)

    checks = tuple(
        check
        for check in InspectionCheck
        if not (check is InspectionCheck.KNOWN_IDENTIFIER_EXACT and identifier_values is None)
    )

    return SafetyReport(
        sha256=hash_before,
        worksheet_count=worksheet_count,
        indexed_entity_values=len(entity_values),
        indexed_identifier_values=None if identifier_values is None else len(identifier_values),
        checks_performed=checks,
        not_inspected=tuple(NotInspectedArea),
        finding_counts=collector.counts(),
        findings=collector.findings(),
        findings_truncated=collector.truncated,
    )


# ----------------------------------------------------------------------
# Preflight / загрузка
# ----------------------------------------------------------------------


def _preflight(
    path: object, mapping_store: object, identifier_store: object
) -> Path:
    if not isinstance(path, (str, Path)):
        raise TypeError(f"path должен быть str или Path, получено: {type(path)!r}")
    if not isinstance(mapping_store, MappingStore):
        raise TypeError(f"mapping_store должен быть MappingStore, получено: {type(mapping_store)!r}")
    if identifier_store is not None and not isinstance(identifier_store, IdentifierMappingStore):
        raise TypeError(
            "identifier_store должен быть IdentifierMappingStore или None, "
            f"получено: {type(identifier_store)!r}"
        )

    source = Path(path)
    if source.suffix.lower() != _SUPPORTED_EXTENSION:
        raise ExternalAiInspectionInputError(
            f"path должен иметь расширение {_SUPPORTED_EXTENSION}, получено: {source.suffix!r}"
        )
    if not source.exists():
        raise ExternalAiInspectionInputError(f"path не существует: {source}")
    if not source.is_file():
        raise ExternalAiInspectionInputError(f"path должен быть файлом: {source}")
    return source


def _load_workbook(source: Path) -> Workbook:
    try:
        return openpyxl.load_workbook(
            str(source),
            read_only=False,
            data_only=False,
            keep_vba=False,
            keep_links=True,
        )
    except Exception:
        # Текст исключений openpyxl может содержать содержимое файла —
        # поэтому фиксированное сообщение и разрыв цепочки (from None).
        raise ExternalAiInspectionError(InspectionFailureReason.UNREADABLE_WORKBOOK) from None


# ----------------------------------------------------------------------
# Сбор findings
# ----------------------------------------------------------------------


class _FindingCollector:
    """Точные счётчики по kind + не более MAX_FINDINGS_PER_KIND конкретных findings."""

    def __init__(self) -> None:
        self._counts: dict[FindingKind, int] = {kind: 0 for kind in FindingKind}
        self._findings: dict[FindingKind, list[SafetyFinding]] = {kind: [] for kind in FindingKind}
        self.truncated = False

    def add(
        self,
        kind: FindingKind,
        location: FindingLocation,
        *,
        worksheet_index: Optional[int] = None,
        row: Optional[int] = None,
        column: Optional[int] = None,
        core_property: Optional[CoreProperty] = None,
    ) -> None:
        self._counts[kind] += 1
        bucket = self._findings[kind]
        if len(bucket) >= MAX_FINDINGS_PER_KIND:
            self.truncated = True
            return
        bucket.append(
            SafetyFinding(
                kind=kind,
                location=location,
                worksheet_index=worksheet_index,
                row=row,
                column=column,
                core_property=core_property,
            )
        )

    def counts(self) -> tuple[tuple[FindingKind, int], ...]:
        return tuple((kind, self._counts[kind]) for kind in FindingKind if self._counts[kind] > 0)

    def findings(self) -> tuple[SafetyFinding, ...]:
        return tuple(finding for kind in FindingKind for finding in self._findings[kind])


# ----------------------------------------------------------------------
# Scan workbook
# ----------------------------------------------------------------------


def _scan_workbook(
    workbook: Workbook,
    entity_values: frozenset[str],
    identifier_values: Optional[frozenset[str]],
    collector: _FindingCollector,
) -> None:
    _scan_restored_markers_and_custom_properties(workbook, collector)
    _scan_core_properties(workbook, entity_values, identifier_values, collector)

    for _ in workbook.defined_names:
        collector.add(FindingKind.DEFINED_NAME, FindingLocation.WORKBOOK)

    for worksheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        _scan_worksheet(worksheet, worksheet_index, entity_values, identifier_values, collector)


def _scan_restored_markers_and_custom_properties(
    workbook: Workbook, collector: _FindingCollector
) -> None:
    for prop in workbook.custom_doc_props:
        name = prop.name
        if name in _RESTORED_MARKER_PROPERTY_NAMES:
            collector.add(FindingKind.RESTORED_MARKER, FindingLocation.WORKBOOK)
        elif isinstance(name, str) and not name.startswith(_DATA_ANONYMIZER_PROPERTY_PREFIX):
            collector.add(FindingKind.CUSTOM_PROPERTY, FindingLocation.WORKBOOK)


def _scan_core_properties(
    workbook: Workbook,
    entity_values: frozenset[str],
    identifier_values: Optional[frozenset[str]],
    collector: _FindingCollector,
) -> None:
    properties = workbook.properties
    for core_property in CoreProperty:
        value = getattr(properties, core_property.value, None)
        if not isinstance(value, str):
            continue

        if value in entity_values:
            collector.add(
                FindingKind.KNOWN_ENTITY_VALUE,
                FindingLocation.CORE_PROPERTY,
                core_property=core_property,
            )
        if identifier_values is not None and value in identifier_values:
            collector.add(
                FindingKind.KNOWN_IDENTIFIER_VALUE,
                FindingLocation.CORE_PROPERTY,
                core_property=core_property,
            )

        # strip() здесь только распознаёт "пустое" значение, в exact-match
        # выше он не участвует.
        if value.strip() == "":
            continue
        if core_property is CoreProperty.CREATOR and value == _BENIGN_CREATOR_VALUE:
            continue
        collector.add(
            FindingKind.CORE_PROPERTY,
            FindingLocation.CORE_PROPERTY,
            core_property=core_property,
        )


def _scan_worksheet(
    worksheet: Worksheet,
    worksheet_index: int,
    entity_values: frozenset[str],
    identifier_values: Optional[frozenset[str]],
    collector: _FindingCollector,
) -> None:
    title = worksheet.title
    if isinstance(title, str):
        if title in entity_values:
            collector.add(
                FindingKind.KNOWN_ENTITY_VALUE,
                FindingLocation.SHEET_TITLE,
                worksheet_index=worksheet_index,
            )
        if identifier_values is not None and title in identifier_values:
            collector.add(
                FindingKind.KNOWN_IDENTIFIER_VALUE,
                FindingLocation.SHEET_TITLE,
                worksheet_index=worksheet_index,
            )

    if worksheet.sheet_state != "visible":
        collector.add(
            FindingKind.HIDDEN_WORKSHEET, FindingLocation.WORKSHEET, worksheet_index=worksheet_index
        )

    for _ in worksheet.tables:
        collector.add(FindingKind.TABLE, FindingLocation.WORKSHEET, worksheet_index=worksheet_index)

    for _ in worksheet.defined_names:
        collector.add(
            FindingKind.DEFINED_NAME, FindingLocation.WORKSHEET, worksheet_index=worksheet_index
        )

    if _has_header_footer_text(worksheet):
        collector.add(
            FindingKind.HEADER_FOOTER, FindingLocation.WORKSHEET, worksheet_index=worksheet_index
        )

    for (row, column), cell in _existing_cells(worksheet):
        _scan_cell(
            cell, worksheet_index, row, column, entity_values, identifier_values, collector
        )


def _has_header_footer_text(worksheet: Worksheet) -> bool:
    for attribute in _HEADER_FOOTER_ATTRIBUTES:
        header_footer = getattr(worksheet, attribute)
        for part in (header_footer.left, header_footer.center, header_footer.right):
            text = part.text
            if isinstance(text, str) and text.strip() != "":
                return True
    return False


def _existing_cells(worksheet: Worksheet) -> list:
    """
    Только реально существующие ячейки, детерминированно по (row, column).
    Молчаливого fallback на dense iter_rows нет: несовместимый _cells —
    безопасная ошибка (fail-closed).
    """
    cells = getattr(worksheet, "_cells", None)
    if not isinstance(cells, dict):
        raise ExternalAiInspectionError(InspectionFailureReason.UNREADABLE_WORKBOOK)
    for key in cells:
        if not (
            isinstance(key, tuple)
            and len(key) == 2
            and isinstance(key[0], int)
            and isinstance(key[1], int)
        ):
            raise ExternalAiInspectionError(InspectionFailureReason.UNREADABLE_WORKBOOK)
    return sorted(cells.items(), key=lambda item: item[0])


def _scan_cell(
    cell: object,
    worksheet_index: int,
    row: int,
    column: int,
    entity_values: frozenset[str],
    identifier_values: Optional[frozenset[str]],
    collector: _FindingCollector,
) -> None:
    if isinstance(cell, MergedCell):
        return

    def add(kind: FindingKind) -> None:
        collector.add(
            kind,
            FindingLocation.CELL,
            worksheet_index=worksheet_index,
            row=row,
            column=column,
        )

    if cell.comment is not None:
        add(FindingKind.COMMENT)
    if cell.hyperlink is not None:
        add(FindingKind.HYPERLINK)

    if cell.data_type == "f":
        add(FindingKind.FORMULA)
        return
    if cell.data_type == "e":
        return

    value = cell.value
    if isinstance(value, str):
        if value in entity_values:
            add(FindingKind.KNOWN_ENTITY_VALUE)
        if identifier_values is not None and value in identifier_values:
            add(FindingKind.KNOWN_IDENTIFIER_VALUE)
    elif isinstance(value, int) and not isinstance(value, bool):
        if identifier_values is not None and str(value) in identifier_values:
            add(FindingKind.KNOWN_IDENTIFIER_VALUE)


# ----------------------------------------------------------------------
# Инвентарь пакета (только имена частей zip, без чтения содержимого)
# ----------------------------------------------------------------------


def _inventory_package(source: Path, collector: _FindingCollector) -> None:
    try:
        with zipfile.ZipFile(source) as archive:
            names = archive.namelist()
    except Exception:
        raise ExternalAiInspectionError(InspectionFailureReason.UNREADABLE_WORKBOOK) from None

    for name in names:
        kind = _classify_part(name)
        if kind is not None:
            collector.add(kind, FindingLocation.PACKAGE)


def _classify_part(name: str) -> Optional[FindingKind]:
    lowered = name.lower()
    if lowered.endswith("/") or lowered.endswith(".rels") or "/_rels/" in lowered:
        return None

    basename = lowered.rsplit("/", 1)[-1]

    if lowered.startswith("xl/externallinks/") and lowered.endswith(".xml"):
        return FindingKind.EXTERNAL_LINK
    if lowered.startswith("xl/drawings/") and lowered.endswith(".xml"):
        return FindingKind.DRAWING
    if lowered.startswith("xl/charts/") and basename.startswith("chart") and lowered.endswith(".xml"):
        return FindingKind.CHART
    if lowered.startswith("xl/chartsheets/") and lowered.endswith(".xml"):
        return FindingKind.CHARTSHEET
    if lowered.startswith("xl/media/"):
        return FindingKind.MEDIA
    if lowered.startswith("xl/embeddings/"):
        return FindingKind.EMBEDDING
    if (
        lowered.startswith("xl/pivotcache/") or lowered.startswith("xl/pivottables/")
    ) and lowered.endswith(".xml"):
        return FindingKind.PIVOT
    if lowered.startswith("customxml/") and lowered.endswith(".xml"):
        return FindingKind.CUSTOM_XML
    if lowered == "xl/vbaproject.bin":
        return FindingKind.VBA
    if (
        lowered.startswith("xl/threadedcomments/") or lowered.startswith("xl/persons/")
    ) and lowered.endswith(".xml"):
        return FindingKind.THREADED_COMMENT
    return None
